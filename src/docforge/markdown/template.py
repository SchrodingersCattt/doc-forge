"""Template-aware Markdown to DOCX assembly.

This module is project agnostic. A caller supplies Markdown paths, a DOCX
template, optional metadata, and an ordered JSON bibliography.
"""

from __future__ import annotations

import copy
import hashlib
import json
import re
import sys
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from pathlib import PurePosixPath
from typing import Iterable, Mapping, Sequence

from docx import Document
from docx.document import Document as DocumentType
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image

from ..bibliography import BibliographyEntry, CitationResolver, format_entry, load_json
from .blocks import Block
from ..output import validate_output_path
from .launcher import (
    accept_docx_revisions,
    add_inline,
    convert_unicode_scripts_in_docx,
    normalize_document_typography,
    parse_markdown,
    remap_image_relationships,
    remove_docx_comments,
    render_blocks_to_doc,
    verify_clean_review_state,
    verify_image_relationships,
    verify_navigation_headings,
)

__all__ = [
    "AssemblyResult",
    "ManuscriptMetadata",
    "assemble_markdown_template",
    "discover_template_styles",
    "_format_bibliography_record",
    "load_bibliography",
    "parse_metadata",
    "sha256_file",
    "verify_template_output",
    "write_assembly_sidecars",
    "CitationResolver",
]

CITATION_RE = re.compile(r"\\citep?\{([^{}]+)\}")
HEADING_RE = re.compile(r"^#\s+(.+?)\s*$")
FORBIDDEN_TOKENS = (
    "[placeholder text]",
    "[title]",
    "[authors]",
    "[abstract]",
    "[section]",
    "[subsection]",
    "[subsubsection]",
    "[citation field]",
    "[acknowledgment]",
)
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


@dataclass(frozen=True)
class ManuscriptMetadata:
    title: str = ""
    authors: str = ""
    affiliations: str = ""
    contacts: tuple[str, ...] = ()
    acknowledgement: str = ""
    author_contributions: str = ""
    code_availability: str = ""
    data_software_availability: str = ""


@dataclass(frozen=True)
class AssemblyResult:
    output: Path
    title: str
    citation_map: Mapping[str, str | int]
    used_citations: tuple[str, ...]
    reference_keys: tuple[str, ...]
    style_map: Mapping[str, str]
    template_sections_before: int
    template_sections_after: int
    section_geometry: tuple[tuple[int, int, int, int, int, int], ...]
    body_columns_before: int
    body_columns_after: int
    column_layout: str
    figure_span: str
    verification: Mapping[str, object]
    font_family: str | None = None
    east_asia_font: str | None = None
    style_profile: str = "template"
    line_numbers: str = "template"
    contacts: tuple[str, ...] = ()
    include_title: bool = True
    strip_level_one_headings: bool = False
    heading_before: float | None = None
    section_sources: tuple[int, ...] = ()
    figures: tuple[Mapping[str, object], ...] = ()
    tables: tuple[Mapping[str, object], ...] = ()
    numbering_prefix: str = ""
    bibliography_scope: str = "auto"
    citation_numbering: str = "first-citation"
    bibliography_profile: str = "markdown"
    include_metadata_back_matter: bool = True
    native_toc: bool = False
    restart_heading_numbering: bool = False
    body_first_line_chars: float | None = None
    page_break_before_h1: bool = False
    body_font_size: float | None = None
    abstract_font_size: float | None = None
    caption_font_size: float | None = None
    reference_font_size: float | None = None


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _metadata_key(label: str) -> str | None:
    label = re.sub(r"[\s:_-]+", " ", label.strip().lower()).strip()
    return {
        "title": "title",
        "author": "authors",
        "authors": "authors",
        "affiliation": "affiliations",
        "affiliations": "affiliations",
        "acknowledgement": "acknowledgement",
        "acknowledgment": "acknowledgement",
        "acknowledgements": "acknowledgement",
        "acknowledgments": "acknowledgement",
        "author contribution": "author_contributions",
        "author contributions": "author_contributions",
        "author contributions statement": "author_contributions",
        "email": "contacts",
        "emails": "contacts",
        "e mail": "contacts",
        "corresponding email": "contacts",
        "corresponding emails": "contacts",
        "code availability": "code_availability",
        "data and software availability": "data_software_availability",
    }.get(label)


def parse_metadata(path: Path) -> ManuscriptMetadata:
    sections: dict[str, list[str]] = {}
    current: str | None = None
    for line in path.read_text(encoding="utf-8-sig").replace("\r\n", "\n").splitlines():
        match = HEADING_RE.match(line.strip())
        if match:
            current = _metadata_key(match.group(1))
            continue
        if current is not None:
            sections.setdefault(current, []).append(line.strip())
    raw_values = {
        key: [part for part in parts if part]
        for key, parts in sections.items()
    }
    author_lines = raw_values.get("authors", [])
    contact_lines = list(raw_values.get("contacts", []))
    retained_authors = []
    for line in author_lines:
        if "@" in line and line.lstrip("*").strip():
            contact_lines.append(line)
        else:
            retained_authors.append(line)
    if len(contact_lines) > 2:
        raise ValueError("Metadata supports at most two contact email lines")
    return ManuscriptMetadata(
        title=" ".join(raw_values.get("title", [])).strip(),
        authors=" ".join(retained_authors).strip(),
        affiliations=" ".join(raw_values.get("affiliations", [])).strip(),
        contacts=tuple(contact_lines),
        acknowledgement=" ".join(raw_values.get("acknowledgement", [])).strip(),
        author_contributions=" ".join(raw_values.get("author_contributions", [])).strip(),
        code_availability=" ".join(raw_values.get("code_availability", [])).strip(),
        data_software_availability=" ".join(raw_values.get("data_software_availability", [])).strip(),
    )


def _is_filled_metadata(value: str) -> bool:
    """Treat unresolved optional metadata markers as absent during assembly."""
    return value.strip().upper() not in {"", "TODO", "TBD", "TBA"}


def _format_bibliography_record(record: Mapping[str, object], key: str) -> str:
    entry = BibliographyEntry.from_mapping(key, record, source_format="json")
    return format_entry(entry, profile="markdown", label=key)


def load_bibliography(path: Path) -> dict[str, str]:
    return {key: format_entry(entry, profile="markdown", label=key) for key, entry in load_json(path).items()}


def _block_texts(block: Block) -> Iterable[str]:
    if block.text:
        yield block.text
    for row in block.rows:
        yield from row


def _citation_keys(blocks: Iterable[Block]) -> tuple[str, ...]:
    found: list[str] = []
    seen: set[str] = set()
    for block in blocks:
        for text in _block_texts(block):
            for match in CITATION_RE.finditer(text):
                for raw_key in match.group(1).split(","):
                    key = raw_key.strip()
                    if key and key not in seen:
                        seen.add(key)
                        found.append(key)
    return tuple(found)


def _load_citation_base(path: Path) -> dict[str, int]:
    payload = json.loads(path.read_text(encoding="utf-8-sig"))
    if not isinstance(payload, dict) or not isinstance(payload.get("citation_map"), dict):
        raise ValueError(f"Citation base must be a docforge manifest with citation_map: {path}")
    result: dict[str, int] = {}
    for key, value in payload["citation_map"].items():
        if not isinstance(key, str) or not isinstance(value, int) or value < 1:
            raise ValueError(f"Citation base contains an invalid mapping: {key!r} -> {value!r}")
        result[key] = value
    return result


def _citation_plan(
    cited: Sequence[str], bibliography: Mapping[str, str | BibliographyEntry], base: Mapping[str, int] | None,
    *, numbering: str = "first-citation"
) -> tuple[tuple[str, ...], dict[str, str | int]]:
    entries = {
        key: value if isinstance(value, BibliographyEntry) else BibliographyEntry.from_mapping(key, value)
        for key, value in bibliography.items()
    }
    resolver = CitationResolver(entries, numbering=numbering, strict=True)
    return resolver.plan(cited, inherited_numbers=base, numbering=numbering)


def _format_citation_labels(labels: Iterable[str | int]) -> str:
    """Sort and collapse consecutive numeric or SI-prefixed citation labels."""
    grouped: dict[str, set[int]] = {"": set(), "S": set()}
    for label in labels:
        value = str(label)
        prefix = "S" if value.startswith("S") else ""
        grouped[prefix].add(int(value.removeprefix("S")))

    ranges: list[str] = []
    for prefix in ("", "S"):
        values = sorted(grouped[prefix])
        start = end = None
        for value in values + [None]:
            if start is None:
                start = end = value
            elif value is not None and value == end + 1:
                end = value
            else:
                ranges.append(f"{prefix}{start}" if start == end else f"{prefix}{start}–{prefix}{end}")
                start = end = value
    return ",".join(ranges)


def _replace_citations(text: str, mapping: Mapping[str, str | int], *, superscript: bool = False) -> str:
    def replace(match: re.Match[str]) -> str:
        keys = [key.strip() for key in match.group(1).split(",") if key.strip()]
        missing = [key for key in keys if key not in mapping]
        if missing:
            raise ValueError(f"Unknown citation key(s): {', '.join(missing)}")
        numbers = _format_citation_labels(mapping[key] for key in keys)
        return ("\ue000" + numbers + "\ue001") if superscript else "[" + numbers + "]"

    pattern = r"\s*" + CITATION_RE.pattern if superscript else CITATION_RE.pattern
    return re.sub(pattern, replace, text)


def _replace_block_citations(
    blocks: Iterable[Block], mapping: Mapping[str, str | int], *, superscript: bool = False
) -> tuple[Block, ...]:
    return tuple(
        Block(
            kind=block.kind,
            text=_replace_citations(block.text, mapping, superscript=superscript),
            level=block.level,
            rows=tuple(
                tuple(_replace_citations(value, mapping, superscript=superscript) for value in row)
                for row in block.rows
            ),
            language=block.language,
            path=block.path,
            options=block.options,
        )
        for block in blocks
    )


def _extract_abstract(blocks: Iterable[Block]) -> tuple[str, tuple[Block, ...]]:
    abstract: list[str] = []
    body: list[Block] = []
    collecting = False
    found = False
    for block in blocks:
        is_abstract = (
            block.kind == "heading"
            and block.level == 1
            and block.text.strip().lower() == "abstract"
        )
        if is_abstract and not found:
            collecting = True
            found = True
            continue
        if collecting and block.kind == "heading" and block.level == 1:
            collecting = False
        if collecting:
            if block.text:
                abstract.append(block.text)
            elif block.rows:
                abstract.extend(" ".join(row) for row in block.rows)
        else:
            body.append(block)
    return " ".join(part.strip() for part in abstract if part.strip()), tuple(body)


def _norm_style(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", value.lower())


def _find_style(document: DocumentType, candidates: Sequence[str]) -> str:
    styles = list(document.styles)
    by_id = {style.style_id: style for style in styles}
    by_name = {style.name: style for style in styles}
    by_norm = {_norm_style(style.name): style for style in styles}
    for candidate in candidates:
        if candidate in by_id:
            return by_id[candidate].style_id
        if candidate in by_name:
            return by_name[candidate].style_id
        style = by_norm.get(_norm_style(candidate))
        if style is not None:
            return style.style_id
    return by_name["Normal"].style_id


def discover_template_styles(document: DocumentType) -> dict[str, str]:
    candidates = {
        "title": ("BBAuthorName", "BB_Author_Name", "BATitle", "BA_Title", "Title"),
        "authors": ("BBAuthorName", "BB_Author_Name", "Author", "Normal"),
        "affiliations": (
            "FACorrespondingAuthorFootnote",
            "FA_Corresponding_Author_Footnote",
            "BCAuthorAddress",
            "BC_Author_Address",
            "Normal",
        ),
        "abstract_title": ("BDAbstractTitle", "BD_Abstract_Title", "Heading 1", "1"),
        "abstract": ("BDAbstract", "BD_Abstract", "Normal"),
        "body": ("TAMainText1", "TAMainText", "TA_Main_Text", "Normal"),
        "heading_1": ("Heading 1", "1"),
        "heading_2": ("Heading 2", "2"),
        "heading_3": ("Heading 3", "3"),
        "caption": ("a4", "Caption", "VA_Figure_Caption"),
        "references_heading": (
            "TFReferencesSection",
            "TF_References_Section",
            "EndNoteBibliographyTitle",
            "Heading 1",
        ),
        "reference": ("EndNoteBibliography", "EndNote Bibliography", "Normal"),
    }
    return {role: _find_style(document, options) for role, options in candidates.items()}


def _resolve_style_profile(
    document: DocumentType, styles: Mapping[str, str], profile: str
) -> dict[str, str]:
    """Resolve an optional JSON semantic-style override against the template."""
    if profile == "template":
        return dict(styles)
    path = Path(profile)
    if not path.is_file():
        raise FileNotFoundError(f"Style profile JSON does not exist: {path}")
    payload = json.loads(path.read_text(encoding="utf-8-sig"))
    if not isinstance(payload, dict):
        raise ValueError(f"Style profile must be a JSON object: {path}")
    result = dict(styles)
    known = set(result)
    unknown = sorted(str(key) for key in payload if key not in known)
    if unknown:
        raise ValueError(f"Style profile contains unknown roles: {', '.join(unknown)}")
    available = {style.style_id for style in document.styles} | {style.name for style in document.styles}
    for role, value in payload.items():
        if not isinstance(value, str) or value not in available:
            raise ValueError(f"Style profile names an unavailable style: {role}={value!r}")
        result[role] = document.styles[value].style_id if value in {style.name for style in document.styles} else value
    return result


def _style(document: DocumentType, style_id: str):
    for style in document.styles:
        if style.style_id == style_id:
            return style
    raise ValueError(f"Template style is unavailable: {style_id}")


def _strip_fonts(element) -> None:
    for properties in element.iter(qn("w:rPr")):
        for tag in (qn("w:rFonts"), qn("w:sz"), qn("w:szCs"), qn("w:kern")):
            for child in list(properties.findall(tag)):
                properties.remove(child)


@dataclass(frozen=True)
class _TemplateRegion:
    index: int
    nodes: tuple
    section: object
    figure: bool = False
    image_node: object | None = None
    caption_node: object | None = None


@dataclass(frozen=True)
class _TemplatePrototypes:
    paragraphs: Mapping[str, object]
    figure_regions: tuple[_TemplateRegion, ...]


def _text_of(element) -> str:
    return "".join(node.text or "" for node in element.iter(qn("w:t")))


def _is_section_break(element) -> bool:
    return element.tag == qn("w:p") and element.find(".//" + qn("w:sectPr")) is not None


def _clean_paragraph_prototype(element):
    """Keep a prototype's paragraph properties but remove example content."""
    result = copy.deepcopy(element)
    properties = result.find(qn("w:pPr"))
    if properties is not None:
        # A paragraph-level rPr controls the mark, not the generated text, and
        # in the reference file contains example-only highlighting.
        mark = properties.find(qn("w:rPr"))
        if mark is not None:
            properties.remove(mark)
        for child in list(properties):
            if child.tag == qn("w:sectPr"):
                properties.remove(child)
    for child in list(result):
        if child.tag != qn("w:pPr"):
            result.remove(child)
    return result


def _run_prototype(element, *, prefer_long: bool = False):
    runs = [run for run in element.findall(".//" + qn("w:r")) if _text_of(run).strip()]
    if not runs:
        runs = list(element.findall(".//" + qn("w:r")))
    if not runs:
        return None
    if prefer_long:
        return copy.deepcopy(max(runs, key=lambda run: len(_text_of(run))))
    return copy.deepcopy(runs[0])


def _sanitize_run_properties(run_properties, *, keep_bold: bool | None = None):
    if run_properties is None:
        return None
    result = copy.deepcopy(run_properties)
    for tag in (qn("w:highlight"), qn("w:shd"), qn("w:rStyle"), qn("w:rPrChange")):
        for child in list(result.findall(tag)):
            result.remove(child)
    if keep_bold is False:
        for child in list(result.findall(qn("w:b"))) + list(result.findall(qn("w:bCs"))):
            result.remove(child)
    return result


def _copy_positive_markers(source_rpr, target_rpr):
    """Carry Markdown emphasis while retaining template font metrics."""
    if source_rpr is None:
        return target_rpr
    if target_rpr is None:
        target_rpr = OxmlElement("w:rPr")
    for tag in ("b", "bCs", "i", "iCs", "u", "strike", "vertAlign"):
        for marker in source_rpr.findall(qn("w:" + tag)):
            value = marker.get(qn("w:val"))
            if value is None or value.lower() not in {"0", "false", "off", "none"}:
                existing = target_rpr.find(qn("w:" + tag))
                if existing is not None:
                    target_rpr.remove(existing)
                target_rpr.append(copy.deepcopy(marker))
    return target_rpr


def _override_run_fonts(element, font_family: str | None, east_asia_font: str | None) -> None:
    """Override generated text fonts only when the caller explicitly requests it."""
    if not font_family and not east_asia_font:
        return
    latin = font_family or east_asia_font
    east_asia = east_asia_font or font_family
    for run in element.iter(qn("w:r")):
        properties = run.find(qn("w:rPr"))
        if properties is None:
            properties = OxmlElement("w:rPr")
            run.insert(0, properties)
        fonts = properties.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            properties.insert(0, fonts)
        for slot in ("ascii", "hAnsi", "cs"):
            if latin:
                fonts.set(qn(f"w:{slot}"), latin)
        if east_asia:
            fonts.set(qn("w:eastAsia"), east_asia)
        for slot in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn(f"w:{slot}"), None)


def _override_run_size(element, points: float) -> None:
    value = str(round(points * 2))
    for run in element.iter(qn("w:r")):
        properties = run.find(qn("w:rPr"))
        if properties is None:
            properties = OxmlElement("w:rPr")
            run.insert(0, properties)
        for tag in ("sz", "szCs"):
            node = properties.find(qn("w:" + tag))
            if node is None:
                node = OxmlElement("w:" + tag)
                properties.append(node)
            node.set(qn("w:val"), value)


def _remove_run_emphasis(element, *, preserve_math: bool = False) -> None:
    for properties in element.iter(qn("w:rPr")):
        tags = ("b", "bCs") if preserve_math else ("b", "bCs", "i", "iCs")
        for tag in tags:
            for node in list(properties.findall(qn("w:" + tag))):
                properties.remove(node)


def _format_caption_runs(element, points: float = 10) -> None:
    """Apply 10 pt roman caption text and bold only the leading label run."""
    _remove_run_emphasis(element, preserve_math=True)
    _override_run_size(element, points)
    first = next(
        (
            run for run in element.iter(qn("w:r"))
            if "".join(node.text or "" for node in run.iter(qn("w:t"))).strip()
        ),
        None,
    )
    if first is None:
        return
    properties = first.find(qn("w:rPr"))
    if properties is None:
        properties = OxmlElement("w:rPr")
        first.insert(0, properties)
    for tag in ("b", "bCs"):
        marker = OxmlElement("w:" + tag)
        properties.append(marker)


def _override_heading_before(element, style_ids: set[str], points: float | None) -> None:
    """Set explicit spacing before generated headings when requested."""
    if points is None:
        return
    before = str(max(0, round(points * 20)))
    for paragraph in element.iter(qn("w:p")):
        properties = paragraph.find(qn("w:pPr"))
        style = properties.find(qn("w:pStyle")) if properties is not None else None
        if style is None or style.get(qn("w:val")) not in style_ids:
            continue
        spacing = properties.find(qn("w:spacing"))
        if spacing is None:
            spacing = OxmlElement("w:spacing")
            properties.append(spacing)
        spacing.set(qn("w:before"), before)


def _new_paragraph(
    document: DocumentType,
    style_id: str,
    text: str,
    *,
    prototype=None,
    bold_default: bool = False,
    uppercase: bool = False,
    font_family: str | None = None,
    east_asia_font: str | None = None,
):
    """Render text with a template paragraph and run-format prototype."""
    temporary = Document()
    paragraph = temporary.add_paragraph()
    add_inline(paragraph, text.upper() if uppercase else text, bold_default=bold_default)
    element = copy.deepcopy(paragraph._p)
    if prototype is not None:
        original_properties = prototype.find(qn("w:pPr"))
        properties = element.find(qn("w:pPr"))
        if properties is not None:
            element.remove(properties)
        if original_properties is not None:
            element.insert(0, _clean_paragraph_prototype(prototype).find(qn("w:pPr")))
    properties = element.find(qn("w:pPr"))
    if properties is None:
        properties = OxmlElement("w:pPr")
        element.insert(0, properties)
    style = properties.find(qn("w:pStyle"))
    if style is None:
        style = OxmlElement("w:pStyle")
        properties.insert(0, style)
    style.set(qn("w:val"), style_id)
    # A generated heading must expose the same outline level as its prototype.
    heading_match = re.fullmatch(r"(?:Heading ?)?([1-9])", style_id)
    if heading_match is not None:
        heading_level = int(heading_match.group(1))
        outline = properties.find(qn("w:outlineLvl"))
        if outline is None:
            outline = OxmlElement("w:outlineLvl")
            properties.append(outline)
        outline.set(qn("w:val"), str(heading_level - 1))

    base_run = _run_prototype(prototype, prefer_long=True) if prototype is not None else None
    base_rpr = _sanitize_run_properties(
        base_run.find(qn("w:rPr")) if base_run is not None else None,
        keep_bold=None,
    )
    for run in element.findall(".//" + qn("w:r")):
        source_rpr = run.find(qn("w:rPr"))
        if source_rpr is None and base_rpr is None:
            continue
        if run.find(qn("w:rPr")) is not None:
            run.remove(run.find(qn("w:rPr")))
        if base_rpr is not None:
            run.insert(0, copy.deepcopy(base_rpr))
        target_rpr = run.find(qn("w:rPr"))
        merged_rpr = _copy_positive_markers(source_rpr, target_rpr)
        if target_rpr is None and merged_rpr is not None and len(merged_rpr):
            run.insert(0, merged_rpr)
    _override_run_fonts(element, font_family, east_asia_font)
    return element


def _new_metadata_paragraph(document: DocumentType, style_id: str, text: str, *, prototype=None):
    """Render metadata literally so correspondence asterisks cannot become emphasis."""
    temporary = Document()
    paragraph = temporary.add_paragraph()
    paragraph.add_run(text)
    element = copy.deepcopy(paragraph._p)
    if prototype is not None:
        properties = element.find(qn("w:pPr"))
        if properties is not None:
            element.remove(properties)
        original_properties = _clean_paragraph_prototype(prototype).find(qn("w:pPr"))
        if original_properties is not None:
            element.insert(0, original_properties)
    properties = element.find(qn("w:pPr"))
    if properties is None:
        properties = OxmlElement("w:pPr")
        element.insert(0, properties)
    style = properties.find(qn("w:pStyle"))
    if prototype is None:
        if style is None:
            style = OxmlElement("w:pStyle")
            properties.insert(0, style)
        style.set(qn("w:val"), style_id)
    base_run = _run_prototype(prototype, prefer_long=True) if prototype is not None else None
    base_rpr = _sanitize_run_properties(base_run.find(qn("w:rPr")) if base_run is not None else None)
    for run in element.findall(".//" + qn("w:r")):
        existing = run.find(qn("w:rPr"))
        if existing is not None:
            run.remove(existing)
        if base_rpr is not None:
            run.insert(0, copy.deepcopy(base_rpr))
    return element


def _native_toc_nodes(document: DocumentType, heading_style_id: str) -> list:
    """Create a native Word TOC field covering Heading 1 through Heading 3."""
    toc_style_id = _find_style(document, ("TOC Heading", "TOCHeading"))
    heading = _new_paragraph(document, toc_style_id or heading_style_id, "CONTENTS", uppercase=False)
    _flush_left_heading(heading)
    paragraph = OxmlElement("w:p")
    run_begin = OxmlElement("w:r")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    run_begin.append(begin)
    paragraph.append(run_begin)
    run_instruction = OxmlElement("w:r")
    instruction = OxmlElement("w:instrText")
    instruction.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    instruction.text = ' TOC \\o "1-3" \\h \\z \\u '
    run_instruction.append(instruction)
    paragraph.append(run_instruction)
    run_separate = OxmlElement("w:r")
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run_separate.append(separate)
    paragraph.append(run_separate)
    placeholder = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "Right-click and update field to populate this table of contents."
    placeholder.append(text)
    paragraph.append(placeholder)
    run_end = OxmlElement("w:r")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run_end.append(end)
    paragraph.append(run_end)
    page_break = OxmlElement("w:p")
    page_run = OxmlElement("w:r")
    break_node = OxmlElement("w:br")
    break_node.set(qn("w:type"), "page")
    page_run.append(break_node)
    page_break.append(page_run)
    return [heading, paragraph, page_break]


def _flush_left_heading(element) -> None:
    properties = element.find(qn("w:pPr"))
    if properties is None:
        properties = OxmlElement("w:pPr")
        element.insert(0, properties)
    for node in list(properties.findall(qn("w:ind"))) + list(properties.findall(qn("w:tabs"))):
        properties.remove(node)
    indentation = OxmlElement("w:ind")
    indentation.set(qn("w:left"), "0")
    indentation.set(qn("w:firstLine"), "0")
    indentation.set(qn("w:hanging"), "0")
    properties.append(indentation)


def _set_heading_outline(element, level: int) -> None:
    properties = element.find(qn("w:pPr"))
    if properties is None:
        properties = OxmlElement("w:pPr")
        element.insert(0, properties)
    outline = properties.find(qn("w:outlineLvl"))
    if outline is None:
        outline = OxmlElement("w:outlineLvl")
        properties.append(outline)
    outline.set(qn("w:val"), str(max(0, level - 1)))


def _set_page_break_before(element) -> None:
    properties = element.find(qn("w:pPr"))
    if properties is None:
        properties = OxmlElement("w:pPr")
        element.insert(0, properties)
    node = properties.find(qn("w:pageBreakBefore"))
    if node is None:
        node = OxmlElement("w:pageBreakBefore")
        properties.append(node)
    node.set(qn("w:val"), "true")


def _remove_paragraph_numbering(element) -> None:
    """Remove inherited list numbering without changing document line numbers."""
    properties = element.find(qn("w:pPr"))
    if properties is None:
        return
    numbering = properties.find(qn("w:numPr"))
    if numbering is not None:
        properties.remove(numbering)


def _terminal_heading(
    document: DocumentType,
    style_id: str,
    text: str,
    *,
    prototype=None,
    page_break_before: bool = False,
):
    """Create one visible, unnumbered terminal Heading 1 paragraph."""
    heading = _new_paragraph(
        document,
        style_id,
        text,
        prototype=prototype,
        uppercase=False,
    )
    _flush_left_heading(heading)
    _set_heading_outline(heading, 1)
    _remove_paragraph_numbering(heading)
    if page_break_before:
        _set_page_break_before(heading)
    return heading


def _number_heading(element, number: int) -> None:
    paragraph = next(iter(element.iter(qn("w:p"))), None)
    if paragraph is None:
        return
    _flush_left_heading(paragraph)
    first_run = next(iter(paragraph.findall(qn("w:r"))), None)
    if first_run is None:
        first_run = OxmlElement("w:r")
        paragraph.append(first_run)
    text_node = first_run.find(qn("w:t"))
    if text_node is None:
        text_node = OxmlElement("w:t")
        first_run.append(text_node)
    text_node.text = f"{number}. " + (text_node.text or "")


def _set_body_first_line_indent(element, characters: float | None) -> None:
    if characters is None:
        return
    properties = element.find(qn("w:pPr"))
    if properties is None:
        properties = OxmlElement("w:pPr")
        element.insert(0, properties)
    indentation = properties.find(qn("w:ind"))
    if indentation is None:
        indentation = OxmlElement("w:ind")
        properties.append(indentation)
    indentation.attrib.pop(qn("w:firstLine"), None)
    indentation.attrib.pop(qn("w:hanging"), None)
    indentation.set(qn("w:firstLineChars"), str(max(0, round(characters * 100))))


def _is_post_equation_continuation(block: Block, previous: Block | None) -> bool:
    """Return whether prose grammatically continues the preceding equation."""
    return (
        previous is not None
        and previous.kind == "equation"
        and re.match(r"^\s*(?:where|with)\b", block.text, re.IGNORECASE) is not None
    )


def _remap_styles(element, source: DocumentType, target: DocumentType) -> None:
    source_styles = {style.style_id: style for style in source.styles}
    target_styles = {style.name: style for style in target.styles}
    for tag in (qn("w:pStyle"), qn("w:rStyle"), qn("w:tblStyle")):
        for node in element.iter(tag):
            source_style = source_styles.get(node.get(qn("w:val"), ""))
            target_style = target_styles.get(source_style.name) if source_style else None
            if target_style is not None:
                node.set(qn("w:val"), target_style.style_id)


def _add_body_style(element, body_style_id: str) -> None:
    for paragraph in element.iter(qn("w:p")):
        properties = paragraph.find(qn("w:pPr"))
        if properties is None:
            properties = OxmlElement("w:pPr")
            paragraph.insert(0, properties)
        if properties.find(qn("w:pStyle")) is None:
            style = OxmlElement("w:pStyle")
            style.set(qn("w:val"), body_style_id)
            properties.insert(0, style)


def _body_column_width_twips(document: DocumentType) -> int:
    """Return the usable width of the section receiving generated body nodes."""
    section = document.sections[1] if len(document.sections) > 1 else document.sections[0]
    page = int(section.page_width / 635)
    margins = int((section.left_margin + section.right_margin) / 635)
    columns = section._sectPr.find(qn("w:cols"))
    count = int(columns.get(qn("w:num"), "1")) if columns is not None else 1
    gap = int(columns.get(qn("w:space"), "0")) if columns is not None else 0
    return max(1, (page - margins - gap * max(count - 1, 0)) // max(count, 1))


def _fit_tables(element, width_twips: int) -> None:
    """Scale generated table grids to the receiving section's column width."""
    for table in element.iter(qn("w:tbl")):
        grid = table.find(qn("w:tblGrid"))
        if grid is None:
            continue
        columns = grid.findall(qn("w:gridCol"))
        old = [max(1, int(column.get(qn("w:w"), "1"))) for column in columns]
        if not old:
            continue
        total = sum(old)
        scaled = [max(240, round(width_twips * value / total)) for value in old]
        scaled[-1] += width_twips - sum(scaled)
        for column, value in zip(columns, scaled):
            column.set(qn("w:w"), str(value))

        properties = table.find(qn("w:tblPr"))
        if properties is None:
            properties = OxmlElement("w:tblPr")
            table.insert(0, properties)
        table_width = properties.find(qn("w:tblW"))
        if table_width is None:
            table_width = OxmlElement("w:tblW")
            properties.insert(0, table_width)
        table_width.set(qn("w:w"), str(width_twips))
        table_width.set(qn("w:type"), "dxa")
        layout = properties.find(qn("w:tblLayout"))
        if layout is None:
            layout = OxmlElement("w:tblLayout")
            properties.append(layout)
        layout.set(qn("w:type"), "fixed")
        rows = table.findall(qn("w:tr"))
        for row_index, row in enumerate(rows):
            row_properties = row.find(qn("w:trPr"))
            if row_properties is None:
                row_properties = OxmlElement("w:trPr")
                row.insert(0, row_properties)
            cant_split = row_properties.find(qn("w:cantSplit"))
            if cant_split is None:
                cant_split = OxmlElement("w:cantSplit")
                row_properties.append(cant_split)
            cells = row.findall(qn("w:tc"))
            for cell, value in zip(cells, scaled):
                cell_properties = cell.find(qn("w:tcPr"))
                if cell_properties is None:
                    cell_properties = OxmlElement("w:tcPr")
                    cell.insert(0, cell_properties)
                cell_width = cell_properties.find(qn("w:tcW"))
                if cell_width is None:
                    cell_width = OxmlElement("w:tcW")
                    cell_properties.insert(0, cell_width)
                cell_width.set(qn("w:w"), str(value))
                cell_width.set(qn("w:type"), "dxa")
                if row_index == 0:
                    for paragraph in cell.findall(qn("w:p")):
                        ppr = paragraph.find(qn("w:pPr"))
                        if ppr is None:
                            ppr = OxmlElement("w:pPr")
                            paragraph.insert(0, ppr)
                        if ppr.find(qn("w:keepNext")) is None:
                            ppr.append(OxmlElement("w:keepNext"))
        # Small audit tables are easier to read as a unit.  Keeping their
        # rows together prevents a continuation page without its header;
        # larger tables are deliberately allowed to flow across pages.
        if len(rows) <= 7:
            for row in rows[:-1]:
                for paragraph in row.iter(qn("w:p")):
                    ppr = paragraph.find(qn("w:pPr"))
                    if ppr is None:
                        ppr = OxmlElement("w:pPr")
                        paragraph.insert(0, ppr)
                    if ppr.find(qn("w:keepNext")) is None:
                        ppr.append(OxmlElement("w:keepNext"))


def _clone_body(blocks: Iterable[Block], target: DocumentType, styles: Mapping[str, str]) -> list:
    source = render_blocks_to_doc(blocks)
    result: list = []
    body_width = _body_column_width_twips(target)
    for child in source._element.body.iterchildren():
        if child.tag == qn("w:sectPr"):
            continue
        clone = copy.deepcopy(child)
        remap_image_relationships(clone, source, target)
        _remap_styles(clone, source, target)
        _add_body_style(clone, styles["body"])
        _fit_tables(clone, body_width)
        result.append(clone)
    return result


def _template_regions(document: DocumentType) -> tuple[_TemplateRegion, ...]:
    body = document._element.body
    regions: list[_TemplateRegion] = []
    pending: list = []
    index = 0
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            regions.append(_TemplateRegion(index, tuple(pending), copy.deepcopy(child)))
            pending = []
            index += 1
            continue
        if _is_section_break(child):
            properties = child.find(qn("w:pPr"))
            section = properties.find(qn("w:sectPr")) if properties is not None else None
            if section is None:
                pending.append(copy.deepcopy(child))
                continue
            nodes = tuple(pending)
            regions.append(_TemplateRegion(index, nodes, copy.deepcopy(section)))
            pending = []
            index += 1
            continue
        pending.append(copy.deepcopy(child))
    if pending:
        terminal = body.sectPr
        if terminal is None:
            raise ValueError("Template has no section properties")
        regions.append(_TemplateRegion(index, tuple(pending), copy.deepcopy(terminal)))
    if not regions:
        raise ValueError("Template has no section properties")

    marked: list[_TemplateRegion] = []
    for region in regions:
        drawing = next((node for node in region.nodes if next(node.iter(qn("w:drawing")), None) is not None), None)
        caption = next(
            (
                node
                for node in region.nodes
                if node.tag == qn("w:p")
                and _is_figure_caption_text(_text_of(node))
            ),
            None,
        )
        non_caption_text = [
            _text_of(node).strip()
            for node in region.nodes
            if node.tag == qn("w:p") and _text_of(node).strip() and node is not caption
        ]
        is_figure = drawing is not None and caption is not None and not non_caption_text
        marked.append(
            _TemplateRegion(region.index, region.nodes, region.section, is_figure, drawing, caption)
        )
    return tuple(marked)


def _template_inline_figure_slots(regions) -> tuple[_TemplateRegion, ...]:
    """Find every image/caption pair, including pairs inside mixed regions.

    Some SI templates keep sample prose, figures, and captions in one section.
    A section-level figure test would miss those slots because the prose is not
    part of the figure itself.  Pairing the drawing paragraph with the next
    non-empty caption paragraph keeps the template's own section geometry while
    making every explicit figure slot reusable.
    """
    slots: list[_TemplateRegion] = []
    for region in regions:
        for index, node in enumerate(region.nodes):
            if node.tag != qn("w:p") or next(node.iter(qn("w:drawing")), None) is None:
                continue
            for following in region.nodes[index + 1:]:
                if following.tag != qn("w:p"):
                    continue
                text = _text_of(following).strip()
                if not text:
                    continue
                if _is_figure_caption_text(text):
                    slots.append(
                        _TemplateRegion(
                            region.index,
                            (node, following),
                            region.section,
                            True,
                            node,
                            following,
                        )
                    )
                break
    return tuple(slots)


def _template_prototypes(document: DocumentType, styles: Mapping[str, str], regions):
    paragraphs = list(document.paragraphs)
    by_style: dict[str, list] = {}
    for paragraph in paragraphs:
        by_style.setdefault(paragraph.style.style_id, []).append(paragraph._p)

    def choose(role: str, *, long: bool = False, text_match: str | None = None):
        values = by_style.get(styles[role], [])
        if text_match:
            values = [node for node in values if text_match in _text_of(node).lower()] or values
        values = [node for node in values if _text_of(node).strip()] or values
        if not values:
            return None
        return copy.deepcopy(max(values, key=lambda node: len(_text_of(node)))) if long else copy.deepcopy(values[0])

    # The reference file uses the same heading style IDs as the numbering
    # definitions; selecting an actual paragraph preserves its numPr/spacing.
    heading_1 = choose("heading_1")
    heading_2 = choose("heading_2")
    heading_3 = choose("heading_3")
    body_candidates = [
        (index, node)
        for index, node in enumerate(paragraphs)
        if node.style.style_id in {styles["body"], "a", "Normal"}
        and _text_of(node._p).strip()
        and "<w:drawing" not in node._p.xml
    ]
    heading_style_ids = {
        styles["heading_1"],
        styles["heading_2"],
        styles["heading_3"],
    }
    first_heading = next(
        (
            index
            for index, node in enumerate(paragraphs)
            if node.style.style_id in heading_style_ids and _text_of(node._p).strip()
        ),
        -1,
    )
    post_front = [
        node
        for index, node in body_candidates
        if index > first_heading
    ]
    body = copy.deepcopy(
        max(post_front or [node for _, node in body_candidates], key=lambda node: len(_text_of(node._p)))._p
    ) if (post_front or body_candidates) else choose("body")
    caption = choose("caption", long=True)
    references = choose("reference", long=True)
    references_heading = next((copy.deepcopy(paragraph._p) for paragraph in paragraphs if paragraph.text.strip().lower() in {"references", "bibliography"}), None)
    supporting_information = next(
        (copy.deepcopy(paragraph._p) for paragraph in paragraphs if paragraph.text.strip().lower() == "supporting information"),
        None,
    )
    si_title = None
    si_authors = None
    si_affiliations = None
    si_contacts = None
    if supporting_information is not None:
        front = [paragraph for paragraph in paragraphs[:first_heading] if paragraph.text.strip()]
        if len(front) >= 5:
            _, si_title, si_authors, si_affiliations, si_contacts = [copy.deepcopy(paragraph._p) for paragraph in front[:5]]
    authors = si_authors if si_authors is not None else choose("authors")
    if styles["authors"] == styles["title"]:
        same_style = by_style.get(styles["authors"], [])
        authors = copy.deepcopy(same_style[1]) if len(same_style) > 1 else None
    paragraphs_map = {
        "supporting_information": supporting_information,
        "title": si_title if si_title is not None else choose("title"),
        "authors": authors,
        "affiliations": si_affiliations if si_affiliations is not None else choose("affiliations"),
        "contacts": si_contacts if si_contacts is not None else choose("affiliations"),
        "abstract": choose("abstract", long=True),
        "body": body,
        "heading_1": heading_1,
        "heading_2": heading_2,
        "heading_3": heading_3,
        "caption": caption,
        "reference": references,
        "references_heading": references_heading if references_heading is not None else choose("references_heading", text_match="references"),
    }
    figure_regions = _template_inline_figure_slots(regions)
    if not figure_regions:
        figure_regions = tuple(region for region in regions if region.figure)
    return _TemplatePrototypes(paragraphs_map, figure_regions)


def _template_uses_superscript_citations(document: DocumentType) -> bool:
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            value = run.text.strip()
            marker = run._r.find(".//" + qn("w:vertAlign"))
            if marker is not None and marker.get(qn("w:val")) == "superscript" and re.fullmatch(r"[0-9]+(?:[,–-][0-9]+)*", value):
                return True
    return False


def _materialize_superscript_citations(document: DocumentType) -> int:
    """Replace private citation sentinels with template-style superscripts."""
    converted = 0
    pattern = re.compile(r"(\ue000(?:S?[0-9]+)(?:[–,](?:S?[0-9]+))*\ue001)")
    # Materialize the run list before replacing nodes; mutating a live lxml
    # iterator otherwise skips sibling paragraphs after the first citation.
    for parent in list(document._element.body.iter(qn("w:r"))):
        text_nodes = [child for child in parent if child.tag == qn("w:t") and child.text and pattern.search(child.text)]
        if not text_nodes:
            continue
        grandparent = parent.getparent()
        if grandparent is None:
            continue
        base_properties = parent.find(qn("w:rPr"))
        replacements = []
        for node in text_nodes:
            for fragment in pattern.split(node.text):
                if not fragment:
                    continue
                run = OxmlElement("w:r")
                if base_properties is not None:
                    run.append(copy.deepcopy(base_properties))
                text = OxmlElement("w:t")
                text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                if fragment.startswith("\ue000"):
                    text.text = fragment[1:-1]
                    properties = run.find(qn("w:rPr"))
                    if properties is None:
                        properties = OxmlElement("w:rPr")
                        run.insert(0, properties)
                    for existing in list(properties.findall(qn("w:vertAlign"))):
                        properties.remove(existing)
                    marker = OxmlElement("w:vertAlign")
                    marker.set(qn("w:val"), "superscript")
                    properties.append(marker)
                    converted += 1
                else:
                    text.text = fragment
                run.append(text)
                replacements.append(run)
        position = list(grandparent).index(parent)
        grandparent.remove(parent)
        for offset, run in enumerate(replacements):
            grandparent.insert(position + offset, run)
    return converted


FIGURE_CAPTION_RE = re.compile(r"\s*(?:Figure|Scheme|Chart)\s+[A-Za-z0-9]+[.:|]?", re.IGNORECASE)
FIGURE_CAPTION_SLOT_RE = re.compile(
    r"\s*\[(?:Figure|Scheme|Chart)\s+Caption\]", re.IGNORECASE
)
CAPTION_LABEL_RE = re.compile(
    r"^\s*(Figure|Scheme|Chart)\s+[A-Za-z]*\d+\s*[.:]?\s*",
    re.IGNORECASE,
)


def _numbered_caption(caption: str, number: int, prefix: str) -> tuple[str, str]:
    match = CAPTION_LABEL_RE.match(caption)
    kind = match.group(1).capitalize() if match else "Figure"
    body = caption[match.end():].strip() if match else caption.strip()
    label = f"{kind} {prefix}{number}"
    return label, f"{label}. {body}".rstrip()


def _is_figure_caption_text(text: str) -> bool:
    """Recognize rendered captions and ACS-style placeholder caption slots."""
    return bool(FIGURE_CAPTION_RE.match(text) or FIGURE_CAPTION_SLOT_RE.match(text))


def _figure_groups(blocks: Sequence[Block]) -> tuple[tuple[tuple[Block, ...], Block | None, str], ...]:
    """Split Markdown into text runs and image/caption pairs.

    A caption immediately following an image is consumed as the caption.  An
    image alt text remains a usable fallback, so ordinary Markdown image
    syntax stays sufficient for callers.
    """
    groups: list[tuple[tuple[Block, ...], Block | None, str]] = []
    text: list[Block] = []
    index = 0
    while index < len(blocks):
        block = blocks[index]
        if block.kind != "image":
            text.append(block)
            index += 1
            continue
        groups.append((tuple(text), block, block.text.strip()))
        text = []
        index += 1
        if index < len(blocks) and blocks[index].kind == "paragraph" and FIGURE_CAPTION_RE.match(blocks[index].text):
            groups[-1] = (groups[-1][0], block, blocks[index].text.strip())
            index += 1
    groups.append((tuple(text), None, ""))
    return tuple(groups)


def _section_break_node(section) -> object:
    paragraph = OxmlElement("w:p")
    properties = OxmlElement("w:pPr")
    properties.append(copy.deepcopy(section))
    paragraph.append(properties)
    return paragraph


def _page_break_node() -> object:
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    node = OxmlElement("w:br")
    node.set(qn("w:type"), "page")
    run.append(node)
    paragraph.append(run)
    return paragraph


def _section_column_count(section) -> int:
    columns = section.find(qn("w:cols"))
    if columns is None:
        return 1
    return max(1, int(columns.get(qn("w:num"), "1")))


def _section_width_twips(section) -> int:
    page = section.find(qn("w:pgSz"))
    margins = section.find(qn("w:pgMar"))
    if page is None or margins is None:
        return 7200
    width = int(page.get(qn("w:w"), "12240"))
    left = int(margins.get(qn("w:left"), "0"))
    right = int(margins.get(qn("w:right"), "0"))
    columns = section.find(qn("w:cols"))
    count = _section_column_count(section)
    gap = int(columns.get(qn("w:space"), "0")) if columns is not None else 0
    return max(1, (width - left - right - gap * max(count - 1, 0)) // count)


def _with_column_layout(section, layout: str):
    if layout not in {"one", "two"}:
        raise ValueError("image columns must be one or two")
    section = copy.deepcopy(section)
    columns = section.find(qn("w:cols"))
    if columns is None:
        columns = OxmlElement("w:cols")
        section.append(columns)
    if layout == "one":
        columns.attrib.pop(qn("w:num"), None)
        columns.attrib.pop(qn("w:space"), None)
    else:
        columns.set(qn("w:num"), "2")
        columns.set(qn("w:space"), columns.get(qn("w:space"), "475"))
    return section


def _with_page_orientation(section, orientation: str):
    if orientation not in {"portrait", "landscape"}:
        raise ValueError("image orientation must be portrait or landscape")
    section = copy.deepcopy(section)
    page = section.find(qn("w:pgSz"))
    if page is None:
        page = OxmlElement("w:pgSz")
        section.insert(0, page)
    width = int(page.get(qn("w:w"), "12240"))
    height = int(page.get(qn("w:h"), "15840"))
    if orientation == "portrait" and width > height:
        width, height = height, width
    elif orientation == "landscape" and width < height:
        width, height = height, width
    page.set(qn("w:w"), str(width))
    page.set(qn("w:h"), str(height))
    if orientation == "landscape":
        page.set(qn("w:orient"), "landscape")
    else:
        page.attrib.pop(qn("w:orient"), None)
    return section


def _with_continuous_section(section):
    """Return a section whose transition is explicitly continuous."""
    section = copy.deepcopy(section)
    section_type = section.find(qn("w:type"))
    if section_type is None:
        section_type = OxmlElement("w:type")
        section.insert(0, section_type)
    section_type.set(qn("w:val"), "continuous")
    return section


def _apply_template_run_formatting(element, prototype) -> None:
    if prototype is None:
        return
    base = _run_prototype(prototype, prefer_long=True)
    base_rpr = _sanitize_run_properties(base.find(qn("w:rPr")) if base is not None else None)
    for run in element.findall(".//" + qn("w:r")):
        source = run.find(qn("w:rPr"))
        if source is not None:
            run.remove(source)
        if base_rpr is not None:
            run.insert(0, copy.deepcopy(base_rpr))
        target_rpr = run.find(qn("w:rPr"))
        merged_rpr = _copy_positive_markers(source, target_rpr)
        if target_rpr is None and merged_rpr is not None and len(merged_rpr):
            run.insert(0, merged_rpr)


def _clone_rendered_block(
    block: Block,
    target: DocumentType,
    styles: Mapping[str, str],
    prototypes: _TemplatePrototypes,
    *,
    equation_number: int | None = None,
    table_number: int | None = None,
    number_prefix: str = "",
    body_font_size: float | None = None,
    caption_font_size: float | None = None,
) -> list:
    """Render one block and graft the template's paragraph prototype."""
    p = prototypes.paragraphs
    if block.kind == "table_caption":
        label = f"Table {number_prefix}{table_number}. " if table_number is not None else "Table. "
        paragraph = _new_paragraph(
            target,
            styles["caption"],
            f"**{label.strip()}** {block.text}",
            prototype=p.get("caption"),
        )
        _format_caption_runs(paragraph, caption_font_size or 10)
        return [paragraph]
    if block.kind in {"paragraph", "reference"}:
        paragraph = _new_paragraph(target, styles["body"], block.text, prototype=p.get("body"))
        if block.kind == "paragraph":
            _override_run_size(paragraph, body_font_size or 11)
        return [paragraph]
    if block.kind == "heading":
        level = min(max(block.level, 1), 3)
        role = f"heading_{level}"
        return [_new_paragraph(target, styles[role], block.text, prototype=p.get(role), bold_default=False, uppercase=level == 1)]
    if block.kind in {"paragraph", "reference", "quote", "ordered", "bullet", "code", "equation", "table", "separator"}:
        generated = render_blocks_to_doc([block], equation_start=equation_number or 1, number_prefix=number_prefix)
        result: list = []
        width = _body_column_width_twips(target)
        for child in generated._element.body.iterchildren():
            if child.tag == qn("w:sectPr"):
                continue
            clone = copy.deepcopy(child)
            remap_image_relationships(clone, generated, target)
            _remap_styles(clone, generated, target)
            _fit_tables(clone, width)
            if block.kind == "equation":
                _add_body_style(clone, styles["body"])
            if block.kind in {"paragraph", "reference", "quote", "ordered", "bullet"}:
                properties = clone.find(qn("w:pPr"))
                if properties is None:
                    properties = OxmlElement("w:pPr")
                    clone.insert(0, properties)
                old_properties = properties
                source_properties = _clean_paragraph_prototype(p.get("body")) if p.get("body") is not None else None
                # Keep list indentation and quote offsets from the generated
                # paragraph while importing the template's font/spacing.
                keep = {qn("w:ind"), qn("w:numPr"), qn("w:tabs"), qn("w:jc")}
                for child_prop in list(old_properties):
                    if child_prop.tag not in keep and child_prop.tag != qn("w:pStyle"):
                        old_properties.remove(child_prop)
                if source_properties is not None:
                    source_ppr = source_properties.find(qn("w:pPr"))
                    if source_ppr is not None:
                        for child_prop in source_ppr:
                            if child_prop.tag not in {qn("w:pStyle"), qn("w:ind"), qn("w:numPr")} and old_properties.find(child_prop.tag) is None:
                                old_properties.append(copy.deepcopy(child_prop))
                style = old_properties.find(qn("w:pStyle"))
                if style is None:
                    style = OxmlElement("w:pStyle")
                    old_properties.insert(0, style)
                style.set(qn("w:val"), styles["body"])
                _apply_template_run_formatting(clone, p.get("body"))
            result.append(clone)
        return result
    return []


def _clone_figure(
    target: DocumentType,
    image: Block,
    caption: str,
    prototype: _TemplateRegion | None,
    caption_prototype,
    caption_style_id: str,
    figure_number: int,
    section_width_twips: int | None = None,
    number_prefix: str = "",
    caption_font_size: float | None = None,
) -> list:
    image_path = Path(image.path)
    if not image_path.is_file():
        raise FileNotFoundError(f"Markdown image does not exist: {image_path}")
    image_bytes = _word_compatible_image_bytes(image_path)
    if prototype is not None and prototype.image_node is not None:
        paragraph = copy.deepcopy(prototype.image_node)
        properties = paragraph.find(qn("w:pPr"))
        if properties is None:
            properties = OxmlElement("w:pPr")
            paragraph.insert(0, properties)
        for child in list(paragraph):
            if child.tag != qn("w:pPr"):
                paragraph.remove(child)
        old_run = next(iter(prototype.image_node.findall(".//" + qn("w:r"))), None)
        run = OxmlElement("w:r")
        if old_run is not None and old_run.find(qn("w:rPr")) is not None:
            run.append(copy.deepcopy(old_run.find(qn("w:rPr"))))
        drawing = next(iter(prototype.image_node.findall(".//" + qn("w:drawing"))), None)
        if drawing is None:
            raise ValueError("Figure prototype has no drawing")
        drawing = copy.deepcopy(drawing)
        relation_id, _ = target.part.get_or_add_image(BytesIO(image_bytes))
        for blip in drawing.iter(qn("a:blip")):
            blip.set(qn("r:embed"), relation_id)
        # Keep the template drawing style, but replace template-specific crop
        # metadata and size the new asset to the available section width.
        for source_rect in list(drawing.iter(qn("a:srcRect"))):
            source_rect.attrib.clear()
        extent = next(iter(drawing.iter(qn("wp:extent"))), None)
        source_extent = next(iter(prototype.image_node.iter(qn("wp:extent"))), None)
        if extent is not None and source_extent is not None:
            width = (
                section_width_twips * 635
                if section_width_twips is not None
                else int(source_extent.get("cx", "1"))
            )
            with Image.open(image_path) as image_file:
                ratio = image_file.height / max(1, image_file.width)
            height = max(1, round(width * ratio))
            extent.set("cx", str(width))
            extent.set("cy", str(height))
            for xfrm in drawing.iter(qn("a:xfrm")):
                ext = next(iter(xfrm.iter(qn("a:ext"))), None)
                if ext is not None:
                    ext.set("cx", str(width))
                    ext.set("cy", str(height))
        for doc_pr in drawing.iter(qn("wp:docPr")):
            doc_pr.set("id", str(1000 + figure_number))
            doc_pr.set("name", image_path.stem)
            doc_pr.set("descr", caption[:255])
        run.append(drawing)
        paragraph.append(run)
    else:
        # Templates without explicit figure slots still receive a real inline
        # drawing, constrained to the current section column.
        temporary = Document()
        picture = temporary.add_paragraph()
        with Image.open(image_path) as image_file:
            available = section_width_twips * 635 if section_width_twips is not None else _body_column_width_twips(target) * 635
            width = min(914400 * 6.5, max(914400, available))
            picture.add_run().add_picture(str(image_path), width=width)
        paragraph = copy.deepcopy(picture._p)
        relation_id, _ = target.part.get_or_add_image(BytesIO(image_bytes))
        for blip in paragraph.iter(qn("a:blip")):
            blip.set(qn("r:embed"), relation_id)
    _, caption_text = _numbered_caption(caption, figure_number, number_prefix)
    caption_text = re.sub(r"^((?:Figure|Scheme|Chart)\s+[A-Za-z0-9]+[.:])", r"**\1**", caption_text)
    cap = _new_paragraph(target, caption_style_id, caption_text, prototype=caption_prototype)
    _format_caption_runs(cap, caption_font_size or 10)
    return [paragraph, cap]


def _word_compatible_image_bytes(image_path: Path, *, max_dimension: int = 4096) -> bytes:
    """Return a conservative RGB PNG payload for reliable Word rendering."""
    with Image.open(image_path) as source:
        image = source.convert("RGB")
        if max(image.size) > max_dimension:
            scale = max_dimension / max(image.size)
            image = image.resize(
                (max(1, round(image.width * scale)), max(1, round(image.height * scale))),
                Image.Resampling.LANCZOS,
            )
        payload = BytesIO()
        image.save(payload, format="PNG", optimize=False)
        return payload.getvalue()


def _section_columns(document: DocumentType, index: int = 1) -> int:
    section = document.sections[min(index, len(document.sections) - 1)]
    columns = section._sectPr.find(qn("w:cols"))
    if columns is None:
        return 1
    return max(1, int(columns.get(qn("w:num"), "1")))


def _set_body_columns(
    section_breaks: Sequence,
    layout: str,
    document: DocumentType,
) -> None:
    """Apply an explicit one/two-column choice to the generated body section."""
    if layout == "template":
        return
    if layout not in {"one", "two"}:
        raise ValueError("columns must be one of: template, one, two")
    if len(section_breaks) >= 2:
        properties = section_breaks[1].find(qn("w:pPr"))
        section = properties.find(qn("w:sectPr"))
    else:
        section = document._element.body.sectPr
    columns = section.find(qn("w:cols"))
    if columns is None:
        columns = OxmlElement("w:cols")
        section.append(columns)
    if layout == "one":
        columns.attrib.pop(qn("w:num"), None)
        columns.attrib.pop(qn("w:space"), None)
    else:
        columns.set(qn("w:num"), "2")
        columns.set(qn("w:space"), columns.get(qn("w:space"), "475"))


def _insert_before(body, anchor, nodes: Iterable) -> None:
    index = list(body).index(anchor)
    for node in nodes:
        body.insert(index, node)
        index += 1


def _geometry(document: DocumentType) -> tuple[tuple[int, int, int, int, int, int], ...]:
    return tuple(
        (
            section.page_width,
            section.page_height,
            section.top_margin,
            section.bottom_margin,
            section.left_margin,
            section.right_margin,
        )
        for section in document.sections
    )


def _update_fields(document: DocumentType) -> None:
    """Disable automatic field updates that trigger Word's external-field prompt."""
    settings = document.settings.element
    node = settings.find(qn("w:updateFields"))
    if node is not None:
        settings.remove(node)


def _source_part(rels_name: str) -> str | None:
    path = PurePosixPath(rels_name)
    if path.parent.name != "_rels" or not path.name.endswith(".rels"):
        return None
    return str(path.parent.parent / path.name[:-5])


def _target(source: str, value: str) -> str:
    if value.startswith("/"):
        return value[1:]
    return str(PurePosixPath(source).parent / value)


def _used_rel_ids(payload: bytes) -> set[str]:
    root = etree.fromstring(payload)
    return {
        value
        for element in root.iter()
        for attribute, value in element.attrib.items()
        if attribute.startswith("{" + REL_NS + "}") and value
    }


def _prune_images(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    targets: set[str] = set()
    for rels_name, payload in list(files.items()):
        source = _source_part(rels_name)
        if source is None or source not in files:
            continue
        root = etree.fromstring(payload)
        used = _used_rel_ids(files[source])
        changed = False
        for relationship in list(root):
            if not relationship.get("Type", "").endswith("/image"):
                continue
            relation_id = relationship.get("Id", "")
            if relation_id not in used:
                root.remove(relationship)
                changed = True
            else:
                targets.add(_target(source, relationship.get("Target", "")))
        if changed:
            files[rels_name] = etree.tostring(
                root, xml_declaration=True, encoding="UTF-8", standalone=True
            )
    for name in list(files):
        if name.startswith("word/media/") and name not in targets:
            del files[name]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, payload in files.items():
            archive.writestr(name, payload)


def _verify_images(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    targets: set[str] = set()
    for rels_name, payload in files.items():
        source = _source_part(rels_name)
        if source is None or source not in files:
            continue
        root = etree.fromstring(payload)
        used = _used_rel_ids(files[source])
        for relationship in root:
            if not relationship.get("Type", "").endswith("/image"):
                continue
            relation_id = relationship.get("Id", "")
            if relation_id not in used:
                raise RuntimeError(f"Unused image relationship: {rels_name}#{relation_id}")
            target = _target(source, relationship.get("Target", ""))
            if target not in files:
                raise RuntimeError(f"Missing image target: {target}")
            targets.add(target)
    leftovers = [name for name in files if name.startswith("word/media/") and name not in targets]
    if leftovers:
        raise RuntimeError(f"Unreferenced image parts remain: {leftovers[:10]}")


def _package_text(path: Path) -> str:
    fragments: list[str] = []
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if not name.startswith("word/") or not name.endswith(".xml"):
                continue
            try:
                root = etree.fromstring(archive.read(name))
            except etree.XMLSyntaxError:
                continue
            fragments.extend(node.text or "" for node in root.iter(qn("w:t")))
            fragments.extend(node.text or "" for node in root.iter(qn("w:instrText")))
    return "\n".join(fragments)


def verify_template_output(
    path: Path,
    *,
    expected_sections: int | None = None,
    expected_geometry: Sequence[tuple[int, int, int, int, int, int]] | None = None,
    expected_body_columns: int | None = None,
    expected_reference_labels: Sequence[str] | None = None,
) -> dict[str, object]:
    verify_navigation_headings(path)
    verify_image_relationships(path)
    verify_clean_review_state(path)
    _verify_images(path)
    with zipfile.ZipFile(path) as archive:
        settings = etree.fromstring(archive.read("word/settings.xml"))
        update_fields = settings.find(qn("w:updateFields"))
        if update_fields is not None and update_fields.get(qn("w:val"), "true").lower() == "true":
            raise RuntimeError("Generated DOCX enables automatic field updates on open")
        unsafe = re.compile(r"\b(?:INCLUDETEXT|INCLUDEPICTURE|DDE(?:AUTO)?|LINK)\b", re.IGNORECASE)
        for name in archive.namelist():
            if not name.startswith("word/") or not name.endswith(".xml"):
                continue
            root = etree.fromstring(archive.read(name))
            field_codes = " ".join(node.text or "" for node in root.iter(qn("w:instrText")))
            if unsafe.search(field_codes):
                raise RuntimeError(f"Unsafe external-file field remains in generated DOCX: {name}")
    document = Document(path)
    geometry = _geometry(document)
    if expected_sections is not None and len(document.sections) != expected_sections:
        raise RuntimeError(
            f"Template section count changed: expected {expected_sections}, found {len(document.sections)}"
        )
    if expected_geometry is not None and tuple(expected_geometry) != geometry:
        raise RuntimeError("Template page geometry changed during assembly")
    body_columns = _section_columns(document)
    if expected_body_columns is not None and body_columns != expected_body_columns:
        raise RuntimeError(
            f"Template body column count changed: expected {expected_body_columns}, found {body_columns}"
        )
    text = _package_text(path)
    lower = text.lower()
    retained = [token for token in FORBIDDEN_TOKENS if token in lower]
    if retained:
        raise RuntimeError(f"Template placeholder text remains: {retained}")
    if CITATION_RE.search(text):
        raise RuntimeError("Raw Markdown citation token remains in generated DOCX")
    if "\ue000" in text or "\ue001" in text:
        raise RuntimeError("Unmaterialized citation marker remains in generated DOCX")
    if expected_reference_labels is not None:
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        if expected_reference_labels and "REFERENCES" not in paragraphs:
            raise RuntimeError("Generated DOCX is missing the REFERENCES heading")
        for label in expected_reference_labels:
            prefix = f"{label}.\t"
            if sum(paragraph.startswith(prefix) for paragraph in paragraphs) != 1:
                raise RuntimeError(
                    f"Generated DOCX must contain exactly one bibliography entry labeled {label}"
                )
    return {
        "clean_review_state": True,
        "image_relationships": "resolved",
        "placeholder_tokens": [],
        "raw_citation_tokens": 0,
        "automatic_field_updates": False,
        "sections": len(document.sections),
        "body_columns": body_columns,
        "section_geometry": [list(value) for value in geometry],
    }


def assemble_markdown_template(
    inputs: Sequence[Path],
    *,
    template_path: Path,
    output: Path,
    metadata_path: Path | None = None,
    bibliography_path: Path | None = None,
    citation_base_path: Path | None = None,
    citation_format: str = "template",
    title: str = "",
    keep_comments: bool = False,
    skip_images: bool = False,
    columns: str = "template",
    figure_span: str = "column",
    font_family: str | None = None,
    east_asia_font: str | None = None,
    style_profile: str = "template",
    line_numbers: str = "template",
    include_title: bool = True,
    strip_level_one_headings: bool = False,
    heading_before: float | None = None,
    numbering_prefix: str = "",
    bibliography_scope: str = "auto",
    citation_numbering: str = "first-citation",
    bibliography_profile: str = "markdown",
    include_metadata_back_matter: bool = True,
    native_toc: bool = False,
    restart_heading_numbering: bool = False,
    body_first_line_chars: float | None = None,
    page_break_before_h1: bool = False,
    body_font_size: float | None = None,
    abstract_font_size: float | None = None,
    caption_font_size: float | None = None,
    reference_font_size: float | None = None,
    force: bool = False,
) -> AssemblyResult:
    validate_output_path(output)
    if output.exists() and not force:
        raise FileExistsError(f"Output exists; pass --force to overwrite: {output}")
    if not inputs:
        raise ValueError("Provide at least one Markdown input")
    if not template_path.is_file():
        raise FileNotFoundError(f"Template DOCX does not exist: {template_path}")
    if metadata_path is not None and not metadata_path.is_file():
        raise FileNotFoundError(f"Metadata Markdown does not exist: {metadata_path}")
    if bibliography_path is not None and not bibliography_path.is_file():
        raise FileNotFoundError(f"Bibliography JSON does not exist: {bibliography_path}")
    if citation_base_path is not None and not citation_base_path.is_file():
        raise FileNotFoundError(f"Citation base manifest does not exist: {citation_base_path}")
    if line_numbers not in {"template", "on", "off"}:
        raise ValueError("line_numbers must be one of: template, on, off")
    if columns not in {"template", "one", "two"}:
        raise ValueError("columns must be one of: template, one, two")
    if figure_span not in {"column", "page"}:
        raise ValueError("figure_span must be one of: column, page")
    if heading_before is not None and heading_before < 0:
        raise ValueError("heading_before must be nonnegative")
    if body_first_line_chars is not None and body_first_line_chars < 0:
        raise ValueError("body_first_line_chars must be nonnegative")
    for name, value in {
        "body_font_size": body_font_size,
        "abstract_font_size": abstract_font_size,
        "caption_font_size": caption_font_size,
        "reference_font_size": reference_font_size,
    }.items():
        if value is not None and value <= 0:
            raise ValueError(f"{name} must be positive")
    if not re.fullmatch(r"[A-Za-z]*", numbering_prefix):
        raise ValueError("numbering_prefix must contain only ASCII letters")
    if bibliography_scope not in {"auto", "all", "new-only"}:
        raise ValueError("bibliography_scope must be one of: auto, all, new-only")
    if citation_numbering not in {"first-citation", "source-order"}:
        raise ValueError("citation_numbering must be one of: first-citation, source-order")
    if bibliography_profile not in {"markdown", "plain"}:
        raise ValueError("bibliography_profile must be one of: markdown, plain")
    resolved_bibliography_scope = (
        "new-only" if bibliography_scope == "auto" and citation_base_path else
        "all" if bibliography_scope == "auto" else
        bibliography_scope
    )
    if resolved_bibliography_scope == "new-only" and citation_base_path is None:
        raise ValueError("bibliography_scope='new-only' requires citation_base_path")

    if citation_format not in {"template", "superscript", "bracketed"}:
        raise ValueError("citation_format must be one of: template, superscript, bracketed")
    template = Document(template_path)
    citation_superscript = citation_format == "superscript" or (
        citation_format == "template" and _template_uses_superscript_citations(template)
    )
    blocks = tuple(
        block
        for path in inputs
        for block in parse_markdown(path, strip_comments=not keep_comments)
    )
    if skip_images:
        blocks = tuple(block for block in blocks if block.kind != "image")
    bibliography_entries = load_json(bibliography_path) if bibliography_path else {}
    bibliography = {
        key: format_entry(entry, profile=bibliography_profile, label=key)
        for key, entry in bibliography_entries.items()
    }
    cited = _citation_keys(blocks)
    missing = sorted(set(cited) - set(bibliography))
    if missing:
        raise ValueError(f"Unknown citation key(s): {', '.join(missing)}")
    citation_base = _load_citation_base(citation_base_path) if citation_base_path else None
    used, mapping = _citation_plan(
        cited, bibliography_entries, citation_base, numbering=citation_numbering
    )
    rendered = _replace_block_citations(blocks, mapping, superscript=citation_superscript)
    abstract, body_blocks = _extract_abstract(rendered)
    if strip_level_one_headings:
        body_blocks = tuple(
            block for block in body_blocks
            if not (block.kind == "heading" and block.level == 1)
        )

    metadata = parse_metadata(metadata_path) if metadata_path else ManuscriptMetadata()
    resolved_title = title.strip() or metadata.title
    if not resolved_title:
        raise ValueError("Template assembly requires --title or a TITLE metadata value")

    sections_before = len(template.sections)
    geometry = _geometry(template)
    styles = _resolve_style_profile(template, discover_template_styles(template), style_profile)
    regions = _template_regions(template)
    prototypes = _template_prototypes(template, styles, regions)
    # Use the actual paragraph style carried by the reference body prototype;
    # A supplied template may carry body paragraphs as ``Normal`` while its
    # semantic discovery role resolves to a custom style.
    body_prototype = prototypes.paragraphs.get("body")
    if body_prototype is not None:
        body_ppr = body_prototype.find(qn("w:pPr"))
        body_style = body_ppr.find(qn("w:pStyle")) if body_ppr is not None else None
        if body_style is not None:
            styles = dict(styles)
            styles["body"] = body_style.get(qn("w:val"), styles["body"])
    body_regions = [region for region in regions if not region.figure]
    if len(body_regions) > 1 and body_regions[0].index == 0:
        front_region = body_regions.pop(0)
    else:
        front_region = None
    if not body_regions:
        body_regions = [front_region] if front_region is not None else [regions[0]]
    body_columns_before = _section_column_count(body_regions[0].section)

    def adjusted(section, *, apply_columns: bool = True):
        section = copy.deepcopy(section)
        line_numbers_node = section.find(qn("w:lnNumType"))
        if line_numbers == "off":
            if line_numbers_node is not None:
                section.remove(line_numbers_node)
        elif line_numbers == "on":
            if line_numbers_node is None:
                line_numbers_node = OxmlElement("w:lnNumType")
                section.append(line_numbers_node)
            line_numbers_node.set(qn("w:countBy"), "1")
            line_numbers_node.set(qn("w:distance"), "120")
            line_numbers_node.set(qn("w:restart"), "newPage")
        elif line_numbers_node is not None:
            # The reference omits a distance value.  LibreOffice then lets
            # three-digit labels touch the column text; the explicit gutter
            # keeps the inherited line-number contract without changing page
            # dimensions or margins.
            line_numbers_node.set(qn("w:distance"), "120")
        if apply_columns and columns in {"one", "two"}:
            col = section.find(qn("w:cols"))
            if col is None:
                col = OxmlElement("w:cols")
                section.append(col)
            if columns == "one":
                col.attrib.pop(qn("w:num"), None)
                col.attrib.pop(qn("w:space"), None)
            else:
                col.set(qn("w:num"), "2")
                col.set(qn("w:space"), col.get(qn("w:space"), "475"))
        return section

    body_columns_after = _section_column_count(adjusted(body_regions[0].section))
    front = []
    if prototypes.paragraphs.get("supporting_information") is not None:
        front.append(_new_metadata_paragraph(template, styles["body"], "Supporting Information", prototype=prototypes.paragraphs.get("supporting_information")))
    if include_title:
        front.append(_new_metadata_paragraph(template, styles["title"], resolved_title, prototype=prototypes.paragraphs.get("title")))
    if metadata.authors:
        front.append(_new_metadata_paragraph(template, styles["authors"], metadata.authors, prototype=prototypes.paragraphs.get("authors")))
    if metadata.affiliations:
        front.append(_new_metadata_paragraph(template, styles["affiliations"], metadata.affiliations, prototype=prototypes.paragraphs.get("affiliations")))
    for contact in metadata.contacts:
        front.append(_new_metadata_paragraph(template, styles["affiliations"], contact, prototype=prototypes.paragraphs.get("contacts")))
    if abstract:
        abstract_prototype = prototypes.paragraphs.get("abstract")
        # Some templates carry the label and the abstract in one styled paragraph.
        abstract_paragraph = _new_paragraph(template, styles["abstract"], f"**ABSTRACT:** {abstract}", prototype=abstract_prototype)
        if abstract_font_size is not None:
            _override_run_size(abstract_paragraph, abstract_font_size)
        front.append(abstract_paragraph)
    if native_toc:
        front.extend(_native_toc_nodes(template, styles["heading_1"]))

    # Build a fresh body while reusing the template's section properties.  The
    # sample text and sample images are discarded; their section geometry,
    # numbering, headers/footers and line-number settings are retained.
    output_nodes: list = []
    section_sources: list[int] = []
    if front_region is not None:
        output_nodes.extend(front)
        output_nodes.append(_section_break_node(adjusted(front_region.section)))
        section_sources.append(front_region.index)
    else:
        output_nodes.extend(front)

    groups = _figure_groups(body_blocks)
    figures: list[dict[str, object]] = []
    tables: list[dict[str, object]] = []
    body_index = 0
    figure_index = 0
    equation_index = 1
    table_index = 1
    heading_counters = {2: 0, 3: 0}
    for text_blocks, image, caption in groups:
        selected_body = body_regions[min(body_index, len(body_regions) - 1)]
        section = adjusted(selected_body.section)
        # _clone_rendered_block computes table widths from the document's
        # active section; the selected section is normally the same geometry,
        # and the explicit fit below handles differing templates.
        previous_block: Block | None = None
        for block in text_blocks:
            if block.kind == "heading" and block.level == 1:
                heading_counters = {2: 0, 3: 0}
            nodes = _clone_rendered_block(
                block,
                template,
                styles,
                prototypes,
                equation_number=equation_index if block.kind == "equation" else None,
                table_number=table_index if block.kind == "table_caption" else None,
                number_prefix=numbering_prefix,
                body_font_size=body_font_size,
                caption_font_size=caption_font_size,
            )
            if block.kind == "heading":
                for node in nodes:
                    _flush_left_heading(node)
                    _set_heading_outline(node, min(max(block.level, 1), 3))
                    if page_break_before_h1 and block.level == 1:
                        _set_page_break_before(node)
                if restart_heading_numbering and block.level == 2:
                    heading_counters[block.level] += 1
                    heading_counters[3] = 0
                    for node in nodes:
                        _number_heading(node, heading_counters[block.level])
            elif block.kind == "paragraph":
                indent = 0 if _is_post_equation_continuation(block, previous_block) else body_first_line_chars
                for node in nodes:
                    _set_body_first_line_indent(node, indent)
            for node in nodes:
                _fit_tables(node, _section_width_twips(section))
            output_nodes.extend(nodes)
            if block.kind == "equation":
                equation_index += 1
            if block.kind == "table_caption":
                label = f"Table {numbering_prefix}{table_index}" if numbering_prefix else f"Table {table_index}"
                tables.append({"number": table_index, "label": label, "caption": block.text})
                table_index += 1
            previous_block = block
        if image is None:
            break
        image_options = dict(image.options)
        requested_orientation = image_options.get("orientation", "portrait")
        if requested_orientation not in {"portrait", "landscape"}:
            raise ValueError("Image orientation marker must be portrait or landscape")
        requested_span = image_options.get("span")
        requested_columns = image_options.get("columns")
        if requested_span is not None and requested_columns is not None:
            raise ValueError("Image span and columns markers cannot be combined")
        if requested_span is not None:
            if requested_span not in {"column", "page"}:
                raise ValueError("Image span marker must be column or page")
            resolved_span = requested_span
        elif requested_columns is not None:
            if requested_columns in {"single", "one"}:
                resolved_span = "page"
            elif requested_columns in {"double", "two"}:
                resolved_span = "column"
            else:
                raise ValueError("Image columns marker must be one, two, single, or double")
        else:
            resolved_span = figure_span
        if requested_orientation == "landscape" and resolved_span != "page":
            raise ValueError("Landscape images require span=page")

        figure_region = None
        if prototypes.figure_regions:
            portrait_regions = [
                region for region in prototypes.figure_regions
                if int(region.section.find(qn("w:pgSz")).get(qn("w:w"), "0"))
                <= int(region.section.find(qn("w:pgSz")).get(qn("w:h"), "0"))
            ]
            landscape_regions = [region for region in prototypes.figure_regions if region not in portrait_regions]
            candidates = landscape_regions if requested_orientation == "landscape" else portrait_regions
            figure_region = candidates[min(figure_index, len(candidates) - 1)] if candidates else prototypes.figure_regions[0]

        if resolved_span == "page":
            base_figure_section = (
                figure_region.section
                if requested_orientation == "landscape" and figure_region is not None
                else section
            )
            figure_section = _with_column_layout(base_figure_section, "one")
            figure_section = _with_page_orientation(figure_section, requested_orientation)
            figure_section = adjusted(figure_section, apply_columns=False)
            if requested_orientation == "portrait":
                figure_section = _with_continuous_section(figure_section)
                body_break_section = _with_continuous_section(section)
            else:
                body_break_section = section
            output_nodes.append(_section_break_node(body_break_section))
            section_sources.append(selected_body.index)
            render_section = figure_section
        else:
            render_section = section

        output_nodes.extend(
            _clone_figure(
                template,
                image,
                caption or image.text or f"Figure {figure_index + 1}.",
                figure_region,
                prototypes.paragraphs.get("caption"),
                styles["caption"],
                figure_index + 1,
                section_width_twips=_section_width_twips(render_section),
                number_prefix=numbering_prefix,
                caption_font_size=caption_font_size,
            )
        )
        if resolved_span == "page":
            output_nodes.append(_section_break_node(figure_section))
            section_sources.append(
                figure_region.index
                if requested_orientation == "landscape" and figure_region is not None
                else selected_body.index
            )
        figures.append(
            {
                "number": figure_index + 1,
                "label": _numbered_caption(caption or image.text, figure_index + 1, numbering_prefix)[0],
                "path": str(Path(image.path).resolve()),
                "sha256": sha256_file(Path(image.path)),
                "caption": caption or image.text,
                "template_region": figure_region.index if figure_region is not None else None,
                "placement_after_body_region": selected_body.index,
                "span": resolved_span,
                "columns": _section_column_count(render_section),
                "orientation": requested_orientation,
            }
        )
        figure_index += 1
        body_index += 1

    reference_keys = tuple(
        key for key in used
        if resolved_bibliography_scope == "all" or key not in citation_base
    ) if citation_base is not None else tuple(used)
    if include_metadata_back_matter and _is_filled_metadata(metadata.acknowledgement):
        output_nodes.append(_terminal_heading(template, styles["heading_1"], "ACKNOWLEDGMENTS", prototype=prototypes.paragraphs.get("heading_1")))
        output_nodes.append(_new_paragraph(template, styles["body"], metadata.acknowledgement, prototype=prototypes.paragraphs.get("body")))
    if include_metadata_back_matter and _is_filled_metadata(metadata.author_contributions):
        output_nodes.append(_terminal_heading(template, styles["heading_1"], "AUTHOR CONTRIBUTIONS", prototype=prototypes.paragraphs.get("heading_1")))
        output_nodes.append(_new_paragraph(template, styles["body"], metadata.author_contributions, prototype=prototypes.paragraphs.get("body")))
    if include_metadata_back_matter and _is_filled_metadata(metadata.code_availability):
        output_nodes.append(_terminal_heading(template, styles["heading_1"], "CODE AVAILABILITY", prototype=prototypes.paragraphs.get("heading_1")))
        output_nodes.append(_new_paragraph(template, styles["body"], metadata.code_availability, prototype=prototypes.paragraphs.get("body")))
    if include_metadata_back_matter and _is_filled_metadata(metadata.data_software_availability):
        output_nodes.append(_terminal_heading(template, styles["heading_1"], "DATA AND SOFTWARE AVAILABILITY", prototype=prototypes.paragraphs.get("heading_1")))
        output_nodes.append(_new_paragraph(template, styles["body"], metadata.data_software_availability, prototype=prototypes.paragraphs.get("body")))
    if reference_keys:
        reference_heading = _terminal_heading(
            template,
            styles["heading_1"],
            "REFERENCES",
            prototype=prototypes.paragraphs.get("heading_1"),
            page_break_before=page_break_before_h1,
        )
        output_nodes.append(reference_heading)
        for key in reference_keys:
            reference = _new_paragraph(template, styles["reference"], f"{mapping[key]}.\t{bibliography[key]}", prototype=prototypes.paragraphs.get("reference"))
            if reference_font_size is not None:
                _override_run_size(reference, reference_font_size)
            output_nodes.append(reference)
    final_region = body_regions[min(body_index, len(body_regions) - 1)]
    # The final section belongs to the document body.  A body-level sectPr is
    # required for Word/LibreOffice to balance the last two-column region.
    output_nodes.append(adjusted(final_region.section))
    section_sources.append(final_region.index)

    heading_style_ids = {
        styles["heading_1"],
        styles["heading_2"],
        styles["heading_3"],
        styles["references_heading"],
    }
    for node in output_nodes:
        _override_heading_before(node, heading_style_ids, heading_before)
        _override_run_fonts(node, font_family, east_asia_font)
    body_element = template._element.body
    for child in list(body_element):
        body_element.remove(child)
    for node in output_nodes:
        body_element.append(node)
    template.core_properties.title = resolved_title
    template.core_properties.author = metadata.authors
    _update_fields(template)
    citation_converted = _materialize_superscript_citations(template) if citation_superscript else 0
    normalize_document_typography(template)
    output.parent.mkdir(parents=True, exist_ok=True)
    template.save(output)
    accept_docx_revisions(output)
    remove_docx_comments(output)
    _prune_images(output)
    convert_unicode_scripts_in_docx(output)
    selected_geometry = tuple(geometry[index] for index in section_sources if index < len(geometry))
    verification = verify_template_output(
        output,
        expected_sections=len(section_sources),
        expected_geometry=selected_geometry,
        expected_body_columns=body_columns_after,
        expected_reference_labels=tuple(str(mapping[key]) for key in reference_keys),
    )
    verification["citation_format"] = "superscript" if citation_superscript else "bracketed"
    verification["citation_count"] = citation_converted
    return AssemblyResult(
        output=output,
        title=resolved_title,
        citation_map=mapping,
        used_citations=used,
        reference_keys=reference_keys,
        style_map=styles,
        template_sections_before=sections_before,
        template_sections_after=len(Document(output).sections),
        section_geometry=selected_geometry,
        body_columns_before=body_columns_before,
        body_columns_after=body_columns_after,
        column_layout=columns,
        figure_span=figure_span,
        verification=verification,
        font_family=font_family,
        east_asia_font=east_asia_font,
        style_profile=style_profile,
        line_numbers=line_numbers,
        contacts=metadata.contacts,
        include_title=include_title,
        strip_level_one_headings=strip_level_one_headings,
        heading_before=heading_before,
        section_sources=tuple(section_sources),
        figures=tuple(figures),
        tables=tuple(tables),
        numbering_prefix=numbering_prefix,
        bibliography_scope=resolved_bibliography_scope,
        citation_numbering=citation_numbering,
        bibliography_profile=bibliography_profile,
        include_metadata_back_matter=include_metadata_back_matter,
        native_toc=native_toc,
        restart_heading_numbering=restart_heading_numbering,
        body_first_line_chars=body_first_line_chars,
        page_break_before_h1=page_break_before_h1,
        body_font_size=body_font_size,
        abstract_font_size=abstract_font_size,
        caption_font_size=caption_font_size,
        reference_font_size=reference_font_size,
    )


def _descriptor(path: Path) -> dict[str, object]:
    return {"path": str(path.resolve()), "sha256": sha256_file(path), "size_bytes": path.stat().st_size}


def write_assembly_sidecars(
    result: AssemblyResult,
    *,
    inputs: Sequence[Path],
    template_path: Path,
    metadata_path: Path | None,
    bibliography_path: Path | None,
    command: Sequence[str],
    citation_base_path: Path | None = None,
) -> tuple[Path, Path]:
    import docx
    import lxml
    import PIL

    descriptors: dict[str, object] = {
        "markdown_inputs": [_descriptor(path) for path in inputs],
        "template": _descriptor(template_path),
    }
    if metadata_path is not None:
        descriptors["metadata"] = _descriptor(metadata_path)
    if bibliography_path is not None:
        descriptors["bibliography"] = _descriptor(bibliography_path)
    if citation_base_path is not None:
        descriptors["citation_base"] = _descriptor(citation_base_path)
    output = _descriptor(result.output)
    manifest = {
        "schema": "docforge.md2docx.template-manifest.v1",
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "command": list(command),
        "inputs": descriptors,
        "output": output,
        "title": result.title,
        "metadata_fields": {
            "contacts": list(result.contacts),
            "include_back_matter": result.include_metadata_back_matter,
        },
        "rendering": {
            "font_family": result.font_family,
            "east_asia_font": result.east_asia_font,
            "style_profile": result.style_profile,
            "line_numbers": result.line_numbers,
            "include_title": result.include_title,
            "strip_level_one_headings": result.strip_level_one_headings,
            "heading_before": result.heading_before,
            "native_toc": result.native_toc,
            "restart_heading_numbering": result.restart_heading_numbering,
            "body_first_line_chars": result.body_first_line_chars,
            "page_break_before_h1": result.page_break_before_h1,
            "figure_span": result.figure_span,
            "body_font_size": result.body_font_size,
            "abstract_font_size": result.abstract_font_size,
            "caption_font_size": result.caption_font_size,
            "reference_font_size": result.reference_font_size,
        },
        "citation_map": dict(result.citation_map),
        "used_citations": list(result.used_citations),
        "reference_keys": list(result.reference_keys),
        "bibliography_scope": result.bibliography_scope,
        "citation_numbering": result.citation_numbering,
        "bibliography_profile": result.bibliography_profile,
        "style_map": dict(result.style_map),
        "figures": [dict(figure) for figure in result.figures],
        "tables": [dict(table) for table in result.tables],
        "numbering_prefix": result.numbering_prefix,
        "template_sections": {
            "before": result.template_sections_before,
            "after": result.template_sections_after,
            "geometry": [list(value) for value in result.section_geometry],
            "body_columns_before": result.body_columns_before,
            "body_columns_after": result.body_columns_after,
            "column_layout": result.column_layout,
            "source_indices": list(result.section_sources),
        },
        "verification": dict(result.verification),
        "environment": {
            "python": sys.version.split()[0],
            "python_docx": docx.__version__,
            "lxml": lxml.__version__,
            "pillow": PIL.__version__,
        },
    }
    manifest_path = result.output.with_suffix(".manifest.json")
    checksum_path = result.output.with_suffix(".sha256")
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    checksum_path.write_text(f"{output['sha256']}  {result.output.name}\n", encoding="utf-8")
    return manifest_path, checksum_path
