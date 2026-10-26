"""DOCX semantic projection, package audit, and exact-reuse metadata.

Rendering a DOCX from scratch stays with the legacy Markdown and TeX
converters. This module reads a package into SemanticDocument and checks
relationships without rebuilding the file.
"""

from __future__ import annotations

import hashlib
import zipfile
from pathlib import Path

from lxml import etree

from ..pipeline.models import Diagnostic, ValidationReport
from ..semantic import BlockNode, DocumentMetadata, InlineNode, Provenance, ResourceDescriptor, SemanticDocument, text_node
from .protocol import FormatCapabilities, ReadOptions, ReadResult, SourceContext, finding


W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"w": W, "r": R, "pr": REL}

_CORE_PARTS = {
    "[Content_Types].xml",
    "_rels/.rels",
    "word/document.xml",
    "word/styles.xml",
    "word/settings.xml",
    "word/fontTable.xml",
    "word/numbering.xml",
    "word/_rels/document.xml.rels",
}

_CAPABILITIES = FormatCapabilities(
    format_id="docx",
    node_kinds=("paragraph", "heading", "table", "figure"),
    citations=False,
    media=True,
    annotations=False,
    preservation=("EXACT_REUSE", "SOURCE_PATCH", "REBUILD"),
    notes=("Semantic DOCX reading drops unsupported native constructs into diagnostics and SourceContext.",),
)


class DocxReader:
    format_id = "docx"

    def probe(self, source: Path) -> FormatCapabilities:
        from ..output import detect_format

        detect_format(source)
        return _CAPABILITIES

    def read(self, source: Path, *, options: ReadOptions) -> ReadResult:
        from ..output import detect_format

        source = Path(source)
        detect_format(source)
        payload = source.read_bytes()
        digest = hashlib.sha256(payload).hexdigest()
        document, diagnostics, resources, features, opaque = project_docx(source)
        return ReadResult(
            document=document,
            source=SourceContext(
                "docx", digest, source.name,
                opaque_parts=tuple(opaque),
                retained_features=tuple(features),
            ),
            diagnostics=tuple(diagnostics),
            capabilities=_CAPABILITIES,
            origin_format="docx",
            origin_name=source.name,
            content_hash=document.semantic_hash(),
        )


def project_docx(path: Path) -> tuple[SemanticDocument, list[Diagnostic], list[ResourceDescriptor], list[str], list[str]]:
    from docx import Document

    diagnostics: list[Diagnostic] = []
    resources: list[ResourceDescriptor] = []
    features: list[str] = []
    opaque: list[str] = []
    with zipfile.ZipFile(path) as package:
        names = set(package.namelist())
        for name in sorted(names):
            if name.startswith("word/media/"):
                resources.append(ResourceDescriptor(
                    "media-" + Path(name).name,
                    "application/octet-stream",
                    hashlib.sha256(package.read(name)).hexdigest(),
                    Path(name).name,
                ))
            elif name not in _CORE_PARTS and not name.startswith("word/media/") and not name.startswith("word/theme/"):
                opaque.append(name)
        document_xml = package.read("word/document.xml")
        if b"w:comment" in document_xml or "word/comments.xml" in names:
            features.append("comments")
            diagnostics.append(finding("info", "docx.comments", "comments are retained on the source package"))
        if b"w:ins" in document_xml or b"w:del" in document_xml:
            features.append("revisions")
            diagnostics.append(finding("info", "docx.revisions", "tracked revisions are retained on the source package"))
        if b"w:txbx" in document_xml:
            features.append("textboxes")
            diagnostics.append(finding("warning", "docx.textboxes", "text boxes are opaque in the semantic projection"))

    doc = Document(str(path))
    counter = 0

    def block_id() -> str:
        nonlocal counter
        counter += 1
        return f"b{counter}"

    blocks: list[BlockNode] = []
    title = ""
    for paragraph in doc.paragraphs:
        style = paragraph.style.name if paragraph.style is not None else ""
        inlines = tuple(_run_inline(run) for run in paragraph.runs if run.text)
        if not inlines and not paragraph.text:
            continue
        if style == "Title" and not title:
            title = paragraph.text
        if style.startswith("Heading"):
            digits = "".join(ch for ch in style if ch.isdigit()) or "1"
            blocks.append(BlockNode(block_id(), "heading", inlines=inlines or (text_node(paragraph.text),), attrs=(("level", digits),), provenance=Provenance(path.name)))
        else:
            blocks.append(BlockNode(block_id(), "paragraph", inlines=inlines or (text_node(paragraph.text),), attrs=(("style", style),), provenance=Provenance(path.name)))
    for table in doc.tables:
        rows = []
        for row in table.rows:
            cells = []
            for cell in row.cells:
                cells.append(BlockNode(block_id(), "table_cell", inlines=(text_node(cell.text),), provenance=Provenance(path.name)))
            rows.append(BlockNode(block_id(), "table_row", children=tuple(cells), provenance=Provenance(path.name)))
        blocks.append(BlockNode(block_id(), "table", children=tuple(rows), provenance=Provenance(path.name)))
    if resources:
        features.append("media")
    document = SemanticDocument(metadata=DocumentMetadata(title=title), blocks=tuple(blocks), resources=tuple(resources))
    return document, diagnostics, resources, features, opaque


def _run_inline(run) -> InlineNode:
    node = text_node(run.text)
    if run.font.subscript:
        return InlineNode("subscript", children=(node,))
    if run.font.superscript:
        return InlineNode("superscript", children=(node,))
    if run.bold and run.italic:
        return InlineNode("strong", children=(InlineNode("emph", children=(node,)),))
    if run.bold:
        return InlineNode("strong", children=(node,))
    if run.italic:
        return InlineNode("emph", children=(node,))
    if run.underline:
        return InlineNode("underline", children=(node,))
    return node


def audit_docx(path: Path) -> ValidationReport:
    """Check package structure and relationship targets. Does not render the file."""

    from ..output import detect_format

    path = Path(path)
    try:
        detect_format(path)
    except Exception as exc:  # format errors are already user-facing
        return ValidationReport("fail", (finding("error", "docx.invalid", str(exc)),))
    diagnostics: list[Diagnostic] = []
    with zipfile.ZipFile(path) as package:
        corrupt = package.testzip()
        if corrupt is not None:
            return ValidationReport("fail", (finding("error", "docx.corrupt", "DOCX member failed CRC"),))
        names = set(package.namelist())
        if "word/document.xml" not in names:
            return ValidationReport("fail", (finding("error", "docx.invalid", "DOCX is missing word/document.xml"),))
        rels_name = "word/_rels/document.xml.rels"
        if rels_name not in names:
            diagnostics.append(finding("error", "docx.relationships", "document relationships part is missing"))
            return ValidationReport("fail", tuple(diagnostics))
        rels, external = _relationship_map(package.read(rels_name))
        diagnostics.extend(_dangling(package.read("word/document.xml"), rels, external, names, "word/document.xml"))
        root = etree.fromstring(package.read("word/document.xml"))
        for reference in root.xpath(".//w:headerReference | .//w:footerReference", namespaces=NS):
            rel_id = reference.get(f"{{{R}}}id")
            if rel_id in external:
                continue
            target = rels.get(rel_id or "")
            if target is None:
                diagnostics.append(finding("error", "docx.header_footer", "section references a missing header or footer"))
                continue
            part = _resolve("word", target)
            if part not in names:
                diagnostics.append(finding("error", "docx.header_footer", "header or footer target is missing"))
                continue
            part_rels = str(Path(part).parent / "_rels" / (Path(part).name + ".rels")).replace("\\", "/")
            if part_rels in names:
                part_map, part_external = _relationship_map(package.read(part_rels))
                diagnostics.extend(_dangling(package.read(part), part_map, part_external, names, part))
    status = "fail" if any(item.severity == "error" for item in diagnostics) else "pass"
    return ValidationReport(status, tuple(diagnostics), {"relationships": "checked"})


def _relationship_map(payload: bytes) -> tuple[dict[str, str], set[str]]:
    root = etree.fromstring(payload)
    internal: dict[str, str] = {}
    external: set[str] = set()
    for item in root.findall(f"{{{REL}}}Relationship"):
        rel_id = item.get("Id")
        if not rel_id:
            continue
        if item.get("TargetMode") == "External":
            external.add(rel_id)
        elif item.get("Target"):
            internal[rel_id] = item.get("Target") or ""
    return internal, external


def _dangling(payload: bytes, rels: dict[str, str], external: set[str], names: set[str], owner: str) -> list[Diagnostic]:
    root = etree.fromstring(payload)
    findings: list[Diagnostic] = []
    parent = str(Path(owner).parent).replace("\\", "/")
    for node in root.xpath(".//*[@r:id]", namespaces=NS):
        rel_id = node.get(f"{{{R}}}id")
        if not rel_id or rel_id in external:
            continue
        if rel_id not in rels:
            findings.append(finding("error", "docx.dangling_relationship", "element references a missing relationship", owner=Path(owner).name))
            continue
        target = rels[rel_id]
        if target.startswith("#") or "://" in target:
            continue
        part = _resolve(parent, target)
        if part not in names:
            findings.append(finding("error", "docx.missing_target", "relationship target is missing from the package", owner=Path(owner).name))
    return findings


def _resolve(parent: str, target: str) -> str:
    return str(Path(parent) / target).replace("\\", "/")


def paragraph_text(path: Path) -> tuple[str, ...]:
    root = etree.fromstring(zipfile.ZipFile(path).read("word/document.xml"))
    texts = []
    for paragraph in root.xpath(".//w:p", namespaces=NS):
        value = "".join(paragraph.xpath(".//w:t/text()", namespaces=NS))
        texts.append(value)
    return tuple(texts)


__all__ = ["DocxReader", "audit_docx", "paragraph_text", "project_docx"]
