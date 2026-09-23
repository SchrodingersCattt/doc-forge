from __future__ import annotations

import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml import OxmlElement
from PIL import Image

from docforge.tex import build_validated_docx, convert_files, parse_bib, validate_docx_package
from docforge.tex.converter import ConversionContext, _ACTIVE_CONTEXT, add_display_math


class TexWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.figure_dir = self.root / "figures"
        self.figure_dir.mkdir()
        Image.new("RGB", (16, 16), "white").save(self.figure_dir / "one.png")
        self.main = self.root / "main.tex"
        self.si = self.root / "SI.tex"
        self.bib = self.root / "refs.bib"
        self.main.write_text(
            r"""\title{A {nested} title}
\begin{abstract}Abstract with \textbf{bold}, $x_i^2$, and \cite{one}.\end{abstract}
\begin{document}
\section{Introduction}
See Figure~\ref{fig:one} and \cite{one}.
\begin{figure}
\includegraphics{figures/one.png}
\caption{A {nested} caption}
\label{fig:one}
\end{figure}
\begin{table}
\caption{Values}
\begin{tabular}{ll}
Group & Value \\
A & 1 \\
\end{tabular}
\end{table}
\begin{algorithm}
\caption{Example}
\begin{algorithmic}
\State \textbf{Return} $x_i$
\end{algorithmic}
\end{algorithm}
\bibliography{refs}
\end{document}
""",
            encoding="utf-8",
        )
        self.si.write_text(
            r"""\title{Supplementary Information \\ Main title}
\begin{document}
\section{Supplementary Note 1}
See main Figure~\ref{fig:one} and \mainref{one}.
\begin{figure}
\includegraphics{figures/one.png}
\caption{SI figure}
\label{fig:si}
\end{figure}
\end{document}
""",
            encoding="utf-8",
        )
        self.bib.write_text(
            "@article{one,\n  author = {Alice Author and Bob Writer},\n  title = {A title},\n  journal = {Journal of Tests},\n  year = {2025},\n  volume = {1},\n  pages = {1--3}\n}\n",
            encoding="utf-8",
        )

    def tearDown(self):
        self.temporary.cleanup()

    def test_convert_files_preserves_document_features_and_validates(self):
        output = self.root / "main.docx"
        convert_files(self.main, self.si, self.bib, output=output)
        summary = validate_docx_package(output, self.main)
        self.assertEqual(1, summary["figures"])
        self.assertEqual(1, summary["tables"])
        text = "\n".join(p.text for p in Document(output).paragraphs)
        self.assertIn("A nested title", text)
        self.assertIn("Figure 1", text)
        self.assertIn("References", text)

    def test_si_cross_reference_and_inherited_citation_number(self):
        output = self.root / "si.docx"
        convert_files(
            self.main, self.si, self.bib, output=output, target="si",
            inherited_citations={"one": 1}, citation_prefix="S", float_prefix="S",
        )
        result = validate_docx_package(output, self.si, self.main)
        self.assertEqual(1, result["figures"])
        visible = "\n".join(p.text for p in Document(output).paragraphs)
        self.assertIn("Figure 1", visible)

    def test_display_math_uses_docforge_math_backend(self):
        document = Document()
        token = _ACTIVE_CONTEXT.set(ConversionContext(self.figure_dir, self.root))
        try:
            with patch("docforge.tex.converter.latex_to_omml", return_value=OxmlElement("m:oMathPara")) as backend:
                add_display_math(document, r"x_i = 1", None)
            backend.assert_called_once_with("x_i = 1")
            self.assertIn("(1)", document.paragraphs[0].text)
        finally:
            _ACTIVE_CONTEXT.reset(token)

    def test_corrupt_build_does_not_replace_previous_output(self):
        output = self.root / "latest.docx"
        output.write_bytes(b"previous valid file")
        def broken(path: Path) -> None:
            path.write_bytes(b"invalid package")
        with self.assertRaises(zipfile.BadZipFile):
            build_validated_docx(output, self.main, broken)
        self.assertEqual(b"previous valid file", output.read_bytes())
        self.assertEqual([], list(self.root.glob(".latest-*.docx")))

    def test_latex_leak_is_rejected(self):
        output = self.root / "leak.docx"
        document = Document()
        document.add_paragraph("A nested title")
        document.add_paragraph(r"Leaked \cite{one}")
        document.save(output)
        with self.assertRaisesRegex(ValueError, "LaTeX command leaked"):
            validate_docx_package(output, self.main)

    def test_missing_figure_is_rejected(self):
        source = self.root / "missing.tex"
        source.write_text(
            r"\begin{document}\begin{figure}\includegraphics{missing.png}\end{figure}\end{document}",
            encoding="utf-8",
        )
        output = self.root / "missing.docx"
        convert_files(source, output=output)
        with self.assertRaisesRegex(ValueError, "missing-figure placeholder"):
            validate_docx_package(output, source)


if __name__ == "__main__":
    unittest.main()
