"""Regression: redline must not emit sectPr header/footer refs without r:id.

Also: deleted paragraphs must not retain live w:sectPr (Word still applies
those for column layout while revisions are shown).
"""

from __future__ import annotations

import zipfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from docforge.docxdiff import create_tracked_docx
from docforge.docxdiff.package import Package
from docforge.docxdiff.redline import (
    R,
    W,
    Context,
    _base_relationships,
    _mark_paragraph,
    _strip_incomplete_header_footer_refs,
)

NS = {"w": W}


def test_redline_omits_baseline_header_refs_without_rid() -> None:
    """Deleted baseline sectPr must not emit header/footer refs without r:id.

    AML baselines may cite word/header*.xml. Those non-media relationships are
    skipped when carrying deleted paragraphs; the reference elements themselves
    must be removed, or Word refuses to open the tracked package.
    """
    header_type = (
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/header"
    )
    paragraph = etree.fromstring(
        f"""
        <w:p xmlns:w="{W}" xmlns:r="{R}">
          <w:pPr>
            <w:sectPr>
              <w:headerReference w:type="even" r:id="rId99"/>
              <w:footerReference w:type="default" r:id="rId99"/>
            </w:sectPr>
          </w:pPr>
          <w:r><w:t>baseline-only section break</w:t></w:r>
        </w:p>
        """
    )
    base_rels = etree.fromstring(
        f"""
        <Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
          <Relationship Id="rId99" Type="{header_type}" Target="header2.xml"/>
        </Relationships>
        """
    )
    current_rels = etree.fromstring(
        """
        <Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>
        """
    )
    base = Package(
        {
            "word/_rels/document.xml.rels": etree.tostring(base_rels),
            "word/header2.xml": b'<w:hdr xmlns:w="%s"/>' % W.encode(),
        }
    )
    current = Package({"word/_rels/document.xml.rels": etree.tostring(current_rels)})

    carry, _save = _base_relationships(base, current)
    carried = carry(paragraph)

    refs = carried.findall(".//w:headerReference", {"w": W}) + carried.findall(
        ".//w:footerReference", {"w": W}
    )
    assert refs == [], f"expected no header/footer refs, got {refs!r}"
    assert (
        current.xml("word/_rels/document.xml.rels").findall(
            "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship"
        )
        == []
    )


def test_strip_incomplete_header_footer_refs_removes_missing_rid() -> None:
    root = etree.fromstring(
        f"""
        <w:document xmlns:w="{W}" xmlns:r="{R}">
          <w:body>
            <w:sectPr>
              <w:headerReference w:type="even"/>
              <w:headerReference w:type="default" r:id="rId1"/>
              <w:footerReference w:type="primary"/>
            </w:sectPr>
          </w:body>
        </w:document>
        """
    )
    _strip_incomplete_header_footer_refs(root)
    headers = root.findall(".//w:headerReference", {"w": W})
    footers = root.findall(".//w:footerReference", {"w": W})
    assert len(headers) == 1
    assert headers[0].get(f"{{{R}}}id") == "rId1"
    assert footers == []


def test_mark_paragraph_deletion_strips_sectpr() -> None:
    """Paragraph-mark deletions must not keep live w:sectPr for Word layout."""
    paragraph = etree.fromstring(
        f"""
        <w:p xmlns:w="{W}">
          <w:pPr>
            <w:sectPr>
              <w:type w:val="continuous"/>
              <w:cols w:num="2" w:space="720"/>
            </w:sectPr>
          </w:pPr>
          <w:r><w:t>baseline-only section break</w:t></w:r>
        </w:p>
        """
    )
    marked = _mark_paragraph(paragraph, "del", Context("tester", 1))
    assert marked.find("./w:pPr/w:rPr/w:del", NS) is not None
    assert marked.findall("./w:pPr/w:sectPr", NS) == []


def _attach_continuous_sectpr(paragraph, cols_num: int | None) -> None:
    """Attach a continuous section break to a python-docx paragraph."""
    ppr = paragraph._p.get_or_add_pPr()
    sectpr = OxmlElement("w:sectPr")
    sect_type = OxmlElement("w:type")
    sect_type.set(qn("w:val"), "continuous")
    sectpr.append(sect_type)
    cols = OxmlElement("w:cols")
    if cols_num is not None:
        cols.set(qn("w:num"), str(cols_num))
    cols.set(qn("w:space"), "720")
    sectpr.append(cols)
    ppr.append(sectpr)


def _mid_doc_cols_sequence(root: etree._Element) -> list[str | None]:
    """Return w:cols/@w:num for each mid-document paragraph sectPr + body."""
    body = root.find("./w:body", NS)
    assert body is not None
    out: list[str | None] = []
    for child in list(body):
        if child.tag == f"{{{W}}}sectPr":
            cols = child.find("./w:cols", NS)
            out.append(cols.get(f"{{{W}}}num") if cols is not None else None)
            continue
        if child.tag != f"{{{W}}}p":
            continue
        for sectpr in child.findall("./w:pPr/w:sectPr", NS):
            cols = sectpr.find("./w:cols", NS)
            out.append(cols.get(f"{{{W}}}num") if cols is not None else None)
    return out


def _write_column_fixture(path: Path, paragraphs: list[tuple[str, int | None]]) -> None:
    """Write a DOCX whose paragraphs carry continuous sectPr with optional cols."""
    document = Document()
    # Reuse the default empty paragraph for the first fixture entry when present.
    first = True
    for text, cols_num in paragraphs:
        if first and document.paragraphs:
            para = document.paragraphs[0]
            para.text = text
            first = False
        else:
            para = document.add_paragraph(text)
        _attach_continuous_sectpr(para, cols_num)
    # Final body section: two-column (typical JPCL letter body).
    body = document.element.body
    body_sect = body.find(qn("w:sectPr"))
    if body_sect is None:
        body_sect = OxmlElement("w:sectPr")
        body.append(body_sect)
    cols = body_sect.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        body_sect.append(cols)
    cols.set(qn("w:num"), "2")
    cols.set(qn("w:space"), "720")
    document.save(path)


def test_deleted_paragraph_sectpr_does_not_skew_column_sequence(tmp_path: Path) -> None:
    """Tracked output must not keep sectPr on deleted paras; cols match current.

    Mid-document continuous section breaks alternate single-column (no w:cols
    num) and two-column around figures. A baseline-only paragraph that carried
    a section break must not leave that sectPr live under w:pPr/w:rPr/w:del.
    """
    base = tmp_path / "base.docx"
    current = tmp_path / "current.docx"
    tracked = tmp_path / "tracked.docx"

    # Base has an extra mid-doc paragraph with a two-column sectPr that current
    # drops. Current keeps the single/two/single pattern around live content.
    _write_column_fixture(
        base,
        [
            ("title abstract single", None),
            ("old figure only in baseline", 2),
            ("body after figure single", None),
            ("second figure two-col", 2),
            ("closing single", None),
        ],
    )
    _write_column_fixture(
        current,
        [
            ("title abstract single", None),
            ("body after figure single", None),
            ("second figure two-col", 2),
            ("closing single", None),
        ],
    )

    summary = create_tracked_docx(base, current, tracked, overwrite=True)
    assert summary["deleted"] >= 1

    root = etree.fromstring(zipfile.ZipFile(tracked).read("word/document.xml"))
    current_root = etree.fromstring(zipfile.ZipFile(current).read("word/document.xml"))

    deleted_with_sectpr = root.xpath(
        ".//w:p[w:pPr/w:rPr/w:del]/w:pPr/w:sectPr",
        namespaces=NS,
    )
    assert deleted_with_sectpr == [], (
        "deleted paragraphs must not retain w:sectPr "
        f"(found {len(deleted_with_sectpr)})"
    )

    assert _mid_doc_cols_sequence(root) == _mid_doc_cols_sequence(current_root)