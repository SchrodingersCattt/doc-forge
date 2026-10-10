from __future__ import annotations

import json
import hashlib
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from lxml import etree

from docforge.cli import build_parser
from docforge.markdown import deliver
import docforge.markdown.delivery as delivery_module


def _revision_docx(path: Path, paragraphs: list[tuple[str, str, str]]) -> None:
    """Create a tiny reviewed DOCX with stable para IDs and real revisions."""
    document = Document()
    for _old, _new, _identifier in paragraphs:
        document.add_paragraph("placeholder")
    document.save(path)
    with zipfile.ZipFile(path, "r") as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    namespace = {
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        "w14": "http://schemas.microsoft.com/office/word/2010/wordml",
    }
    root = etree.fromstring(files["word/document.xml"])
    body_paragraphs = root.xpath(".//w:body/w:p", namespaces=namespace)
    for paragraph, (old, new, identifier) in zip(body_paragraphs, paragraphs):
        paragraph.set("{%s}paraId" % namespace["w14"], identifier)
        for child in list(paragraph):
            if child.tag != "{%s}pPr" % namespace["w"]:
                paragraph.remove(child)
        deleted = etree.SubElement(paragraph, "{%s}del" % namespace["w"], {"{%s}id" % namespace["w"]: "1"})
        deleted_run = etree.SubElement(deleted, "{%s}r" % namespace["w"])
        etree.SubElement(deleted_run, "{%s}delText" % namespace["w"]).text = old
        inserted = etree.SubElement(paragraph, "{%s}ins" % namespace["w"], {"{%s}id" % namespace["w"]: "2"})
        inserted_run = etree.SubElement(inserted, "{%s}r" % namespace["w"])
        etree.SubElement(inserted_run, "{%s}t" % namespace["w"]).text = new
    files["word/document.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, payload in files.items():
            archive.writestr(name, payload)


def test_redline_accept_baseline_named_arguments_parse() -> None:
    args = build_parser().parse_args(
        [
            "redline",
            "--baseline",
            "reviewed.docx",
            "--current",
            "current.docx",
            "--output",
            "tracked.docx",
            "--accept-baseline",
            "--revision-author",
            "Reviewer",
            "--report",
            "report.json",
        ]
    )
    assert args.baseline == Path("reviewed.docx")
    assert args.current_option == Path("current.docx")
    assert args.accept_baseline is True
    assert args.revision_author == "Reviewer"
    assert args.report_path == Path("report.json")


def test_deliver_accept_revisions_fails_closed_without_review_input(tmp_path: Path) -> None:
    profile = {
        "delivery_dir": "delivery",
        "article": {"inputs": ["body.md"], "template": "template.docx"},
    }
    (tmp_path / "body.md").write_text("Body.\n", encoding="utf-8")
    document = Document()
    document.add_paragraph("Template")
    document.save(tmp_path / "template.docx")
    path = tmp_path / "profile.json"
    path.write_text(json.dumps(profile), encoding="utf-8")
    with pytest.raises(ValueError, match="reviewed_docx"):
        deliver(path, accept_revisions=True)


def test_deliver_default_rejects_unapplied_revisions_before_output(tmp_path: Path) -> None:
    reviewed = tmp_path / "reviewed.docx"
    document = Document()
    document.add_paragraph("Reviewed")
    document.save(reviewed)
    with zipfile.ZipFile(reviewed, "r") as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    xml = files["word/document.xml"]
    marker = b'<w:ins w:id="7" w:author="Reviewer">'
    xml = xml.replace(b"<w:r>", marker + b"<w:r>", 1).replace(b"</w:r>", b"</w:r></w:ins>", 1)
    files["word/document.xml"] = xml
    with zipfile.ZipFile(reviewed, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, payload in files.items():
            archive.writestr(name, payload)
    (tmp_path / "body.md").write_text("Body.\n", encoding="utf-8")
    template = Document()
    template.add_paragraph("Template")
    template.save(tmp_path / "template.docx")
    profile = {
        "delivery_dir": "delivery",
        "article": {
            "inputs": ["body.md"], "template": "template.docx",
            "reviewed_docx": "reviewed.docx", "section_map": "map.json", "source_dir": ".",
        },
    }
    (tmp_path / "map.json").write_text(json.dumps({"sections": [{"file": "body.md", "start_line": 1, "end_line": 1, "hash": "bad"}]}), encoding="utf-8")
    profile_path = tmp_path / "profile.json"
    profile_path.write_text(json.dumps(profile), encoding="utf-8")
    with pytest.raises(ValueError, match="unapplied revisions"):
        deliver(profile_path)
    assert not (tmp_path / "delivery").exists()


def test_delivery_output_pattern_is_confined_and_expands_role_timestamp() -> None:
    from docforge.markdown.delivery import _output_name

    profile = {"output_pattern": "{role}_{timestamp}.docx"}
    path = _output_name({}, "article", profile, timestamp="20261010T000000Z")
    assert path == Path("article_20261010T000000Z.docx")
    with pytest.raises(ValueError, match="confined"):
        _output_name({}, "article", {"output_pattern": "../{role}.docx"}, timestamp="now")


def test_delivery_rejects_colliding_artifact_destinations(tmp_path: Path) -> None:
    (tmp_path / "body.md").write_text("Body.\n", encoding="utf-8")
    template = Document()
    template.add_paragraph("Template")
    template.save(tmp_path / "template.docx")
    profile = {
        "delivery_dir": "delivery",
        "artifacts": [
            {"id": "article", "inputs": ["body.md"], "template": "template.docx", "output": "same.docx"},
            {"id": "supplement", "inputs": ["body.md"], "template": "template.docx", "output": "same.docx"},
        ],
    }
    path = tmp_path / "profile.json"
    path.write_text(json.dumps(profile), encoding="utf-8")
    with pytest.raises(ValueError, match="collide"):
        deliver(path)
    assert not (tmp_path / "delivery").exists()


def test_review_source_ranges_are_staged_and_hash_checked(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source = tmp_path / "body.md"
    source.write_text("old\nkeep\n", encoding="utf-8")
    template = Document()
    template.add_paragraph("Template")
    template.save(tmp_path / "template.docx")
    reviewed = tmp_path / "reviewed.docx"
    _revision_docx(reviewed, [("old", "new", "p1")])
    old_hash = delivery_module._sha256_text("old\n")
    reviewed_hash = delivery_module._sha256_text("new\n")
    section_map = {
        "sections": [{
            "file": "body.md", "start_line": 1, "end_line": 1,
            "reviewed_start_line": 1, "reviewed_end_line": 1, "hash": old_hash,
            "reviewed_paragraph": {"id": "p1", "hash": reviewed_hash},
        }],
    }
    (tmp_path / "map.json").write_text(json.dumps(section_map), encoding="utf-8")
    profile = {
        "delivery_dir": "delivery",
        "article": {
            "inputs": ["body.md"], "template": "template.docx",
            "reviewed_docx": "reviewed.docx", "section_map": "map.json",
            "source_paragraph_map": "map.json", "source_dir": ".",
        },
    }
    profile_path = tmp_path / "profile.json"
    profile_path.write_text(json.dumps(profile), encoding="utf-8")

    def fake_docx_to_markdown(_input: Path, *, output: Path, **_kwargs):
        output.write_text("new\n", encoding="utf-8")

    def fake_assemble(_inputs, *, output: Path, **_kwargs):
        document = Document()
        document.add_paragraph("new")
        document.save(output)
        return SimpleNamespace(output=output)

    def fake_sidecars(result, **_kwargs):
        manifest = result.output.with_suffix(".manifest.json")
        checksum = result.output.with_suffix(".sha256")
        manifest.write_text("{}\n", encoding="utf-8")
        checksum.write_text("hash\n", encoding="utf-8")
        return manifest, checksum

    monkeypatch.setattr(delivery_module, "docx_to_markdown", fake_docx_to_markdown)
    monkeypatch.setattr(delivery_module, "assemble_markdown_template", fake_assemble)
    monkeypatch.setattr(delivery_module, "write_assembly_sidecars", fake_sidecars)
    deliver(profile_path, accept_revisions=True, force=True)
    assert source.read_text(encoding="utf-8") == "new\nkeep\n"
    manifest = json.loads((tmp_path / "delivery" / "delivery.manifest.json").read_text(encoding="utf-8"))
    exported = manifest["sections"][0]
    assert exported["file"] == "body.md"
    assert exported["reviewed_paragraph"]["id"] == "p1"
    assert exported["reviewed_start_line"] == 1
    assert exported["docx_paragraph_id"] == hashlib.sha1(b"p1").hexdigest()[:8].upper()
    assert exported["docx_paragraph_range"] == {"start": 1, "end": 1}
    source.write_text("changed\nkeep\n", encoding="utf-8")
    with pytest.raises(RuntimeError, match="hash mismatch"):
        deliver(profile_path, accept_revisions=True, force=True)
    assert source.read_text(encoding="utf-8") == "changed\nkeep\n"


def test_review_ranges_use_immutable_coordinates_for_one_source(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source_dir = tmp_path / "sources"
    source_dir.mkdir()
    source = source_dir / "body.md"
    source.write_text("a\nb\nc\nd\n", encoding="utf-8")
    section_map = {
        "sections": [
            {
                "file": "body.md", "start_line": 1, "end_line": 1,
                "reviewed_start_line": 1, "reviewed_end_line": 1,
                "hash": delivery_module._sha256_text("a\n"),
                "reviewed_paragraph": {"id": "p1", "hash": delivery_module._sha256_text("A\n")},
            },
            {
                "file": "body.md", "start_line": 3, "end_line": 3,
                "reviewed_start_line": 2, "reviewed_end_line": 2,
                "hash": delivery_module._sha256_text("c\n"),
                "reviewed_paragraph": {"id": "p2", "hash": delivery_module._sha256_text("C\n")},
            },
        ],
    }
    (tmp_path / "map.json").write_text(json.dumps(section_map), encoding="utf-8")
    reviewed = tmp_path / "reviewed.docx"
    _revision_docx(reviewed, [("a", "A", "p1"), ("c", "C", "p2")])
    monkeypatch.setattr(delivery_module, "docx_to_markdown", lambda _input, *, output, **_kwargs: output.write_text("A\nC\n", encoding="utf-8"))
    prepared = delivery_module._prepare_review(
        {
            "reviewed_docx": "reviewed.docx", "section_map": "map.json",
            "source_paragraph_map": "map.json", "source_dir": "sources",
        },
        tmp_path,
        tmp_path / "stage",
    )
    assert prepared.inputs[0].read_text(encoding="utf-8") == "A\nb\nC\nd\n"


def test_review_rejects_unmatched_reviewed_paragraph(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """An unrelated DOCX paragraph must never become a source-file overwrite."""
    source = tmp_path / "body.md"
    source.write_text("original\n", encoding="utf-8")
    template = Document()
    template.add_paragraph("Template")
    template.save(tmp_path / "template.docx")
    _revision_docx(tmp_path / "reviewed.docx", [("original", "original", "p-original")])
    section_map = {
        "sections": [{
            "file": "body.md", "start_line": 1, "end_line": 1,
            "reviewed_start_line": 1, "reviewed_end_line": 1,
            "hash": delivery_module._sha256_text("original\n"),
            "reviewed_paragraph": {"id": "p-original", "hash": delivery_module._sha256_text("original\n")},
        }],
    }
    (tmp_path / "map.json").write_text(json.dumps(section_map), encoding="utf-8")
    profile = {
        "delivery_dir": "delivery",
        "article": {
            "inputs": ["body.md"], "template": "template.docx",
            "reviewed_docx": "reviewed.docx", "section_map": "map.json",
            "source_paragraph_map": "map.json", "source_dir": ".",
        },
    }
    profile_path = tmp_path / "profile.json"
    profile_path.write_text(json.dumps(profile), encoding="utf-8")

    monkeypatch.setattr(
        delivery_module,
        "docx_to_markdown",
        lambda _input, *, output, **_kwargs: output.write_text("original\nunrelated insertion\n", encoding="utf-8"),
    )
    def fake_assemble(_inputs, *, output: Path, **_kwargs):
        output.write_bytes(b"docx")
        return SimpleNamespace(output=output)

    def fake_sidecars(result, **_kwargs):
        manifest = result.output.with_suffix(".manifest.json")
        checksum = result.output.with_suffix(".sha256")
        manifest.write_text("{}\n", encoding="utf-8")
        checksum.write_text("hash\n", encoding="utf-8")
        return manifest, checksum

    monkeypatch.setattr(delivery_module, "assemble_markdown_template", fake_assemble)
    monkeypatch.setattr(delivery_module, "write_assembly_sidecars", fake_sidecars)
    with pytest.raises(ValueError, match="unmatched paragraphs"):
        deliver(profile_path, accept_revisions=True, force=True)
    assert source.read_text(encoding="utf-8") == "original\n"
    assert not (tmp_path / "delivery").exists()


def test_delivery_rejects_sidecar_and_aggregate_collisions(tmp_path: Path) -> None:
    (tmp_path / "body.md").write_text("Body.\n", encoding="utf-8")
    template = Document()
    template.add_paragraph("Template")
    template.save(tmp_path / "template.docx")
    profile = {
        "delivery_dir": "delivery",
        "manifest": "article.manifest.json",
        "article": {"inputs": ["body.md"], "template": "template.docx", "output": "article.docx"},
    }
    with pytest.raises(ValueError, match="collide"):
        deliver(profile)
    assert not (tmp_path / "delivery").exists()


def test_review_stages_only_mapped_sources_and_extracted_media(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source_dir = tmp_path / "sources"
    source_dir.mkdir()
    source = source_dir / "body.md"
    source.write_text("old\n", encoding="utf-8")
    (source_dir / "unrelated.txt").write_text("private\n", encoding="utf-8")
    template = Document()
    template.add_paragraph("Template")
    template.save(tmp_path / "template.docx")
    _revision_docx(tmp_path / "reviewed.docx", [("old", "Figure", "p1")])
    section_map = {"sections": [{
        "file": "body.md", "start_line": 1, "end_line": 1,
        "reviewed_start_line": 1, "reviewed_end_line": 1,
        "hash": delivery_module._sha256_text("old\n"),
        "reviewed_paragraph": {"id": "p1", "hash": delivery_module._sha256_text("![Figure](media/image.png)\n")},
    }]}
    (tmp_path / "map.json").write_text(json.dumps(section_map), encoding="utf-8")

    def fake_docx_to_markdown(_input: Path, *, output: Path, **_kwargs):
        media = output.parent / "media"
        media.mkdir()
        (media / "image.png").write_bytes(b"embedded image")
        output.write_text("![Figure](media/image.png)\n", encoding="utf-8")

    monkeypatch.setattr(delivery_module, "docx_to_markdown", fake_docx_to_markdown)
    prepared = delivery_module._prepare_review(
        {
            "reviewed_docx": "reviewed.docx", "section_map": "map.json",
            "source_paragraph_map": "map.json", "source_dir": "sources",
            "inputs": ["sources/body.md"],
        },
        tmp_path,
        tmp_path / "stage",
    )
    staged = prepared.inputs[0]
    assert staged.read_text(encoding="utf-8") == "![Figure](media/image.png)\n"
    assert (staged.parent / "media/image.png").read_bytes() == b"embedded image"
    assert not (staged.parent / "unrelated.txt").exists()


def test_review_rejects_malformed_docx_before_delivery(tmp_path: Path) -> None:
    (tmp_path / "body.md").write_text("Body.\n", encoding="utf-8")
    template = Document()
    template.add_paragraph("Template")
    template.save(tmp_path / "template.docx")
    reviewed = tmp_path / "reviewed.docx"
    template.save(reviewed)
    with zipfile.ZipFile(reviewed, "r") as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    files["word/document.xml"] = b"<w:document>"
    with zipfile.ZipFile(reviewed, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, payload in files.items():
            archive.writestr(name, payload)
    profile = {
        "delivery_dir": "delivery",
        "article": {"inputs": ["body.md"], "template": "template.docx", "reviewed_docx": "reviewed.docx"},
    }
    profile_path = tmp_path / "profile.json"
    profile_path.write_text(json.dumps(profile), encoding="utf-8")
    with pytest.raises(ValueError, match="invalid XML.*document.xml"):
        deliver(profile_path)
    assert not (tmp_path / "delivery").exists()


def test_review_rejects_mapped_source_symlink_escape(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source_dir = tmp_path / "sources"
    source_dir.mkdir()
    outside = tmp_path / "outside.md"
    outside.write_text("outside\n", encoding="utf-8")
    (source_dir / "link.md").symlink_to(outside)
    _revision_docx(tmp_path / "reviewed.docx", [("outside", "new", "p1")])
    section_map = {"sections": [{
        "file": "link.md", "start_line": 1, "end_line": 1,
        "reviewed_start_line": 1, "reviewed_end_line": 1,
        "hash": delivery_module._sha256_text("outside\n"),
        "reviewed_paragraph": {"id": "p1", "hash": delivery_module._sha256_text("new\n")},
    }]}
    (tmp_path / "map.json").write_text(json.dumps(section_map), encoding="utf-8")
    monkeypatch.setattr(delivery_module, "docx_to_markdown", lambda _input, *, output, **_kwargs: output.write_text("new\n", encoding="utf-8"))
    with pytest.raises(ValueError, match="escapes its configured directory"):
        delivery_module._prepare_review(
            {"reviewed_docx": "reviewed.docx", "section_map": "map.json", "source_paragraph_map": "map.json", "source_dir": "sources"},
            tmp_path,
            tmp_path / "stage",
        )
