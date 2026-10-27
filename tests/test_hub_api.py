"""Hub API, legacy parity, and fidelity behavior on synthetic documents."""

from __future__ import annotations

import hashlib
import shutil
import subprocess
import zipfile
from dataclasses import replace
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import parse_xml
from lxml import etree
from docx.shared import Pt
from PIL import Image

from docforge import audit, convert, read, restyle, write
from docforge.adapters.protocol import ConversionError, LossReport
from docforge.cli import main as cli_main
from docforge.semantic import BlockNode, InlineNode, SemanticDocument, text_node
from docforge.style import load_profile, resolved_paragraph_styles, restyle_docx


def test_public_commands_remain_available() -> None:
    from docforge.cli import build_parser

    choices = set(build_parser()._subparsers._group_actions[0].choices)
    required = {
        "md2docx", "tex2docx", "docx2md", "docx2tex", "md-export", "md-build",
        "roundtrip", "redline", "contract", "deliver", "convert", "restyle", "audit",
        "export", "build",
    }
    assert required <= choices


def test_md2docx_api_matches_cli(tmp_path: Path) -> None:
    source = tmp_path / "sample.md"
    source.write_text(
        "# Sample\n\nA paragraph with **bold** and *italic* text.\n\n- alpha\n- beta\n\n"
        "1. first\n2. second\n",
        encoding="utf-8",
    )
    via_cli = tmp_path / "cli.docx"
    via_api = tmp_path / "api.docx"
    assert cli_main(["md2docx", str(source), "-o", str(via_cli), "--force"]) == 0
    result = convert(source, via_api, force=True)
    assert result.execution_mode == "REBUILD"
    assert _docx_view(via_cli) == _docx_view(via_api)
    assert hashlib.sha256(source.read_bytes()).hexdigest() == hashlib.sha256(source.read_bytes()).hexdigest()


def test_tex2docx_api_matches_cli(tmp_path: Path) -> None:
    source = tmp_path / "sample.tex"
    source.write_text(
        "\\documentclass{article}\n\\begin{document}\n\\section{Introduction}\n"
        "A short paragraph with \\textbf{bold} text.\n\\end{document}\n",
        encoding="utf-8",
    )
    via_cli = tmp_path / "cli.docx"
    via_api = tmp_path / "api.docx"
    assert cli_main(["tex2docx", str(source), "-o", str(via_cli), "--force"]) == 0
    convert(source, via_api, force=True)
    assert _docx_view(via_cli) == _docx_view(via_api)


def test_generic_final_filename_is_accepted(tmp_path: Path) -> None:
    source = tmp_path / "sample.md"
    source.write_text("Hello\n", encoding="utf-8")
    output = tmp_path / "paper_final.docx"
    assert cli_main(["md2docx", str(source), "-o", str(output), "--force"]) == 0
    assert output.is_file()


def test_markdown_latex_round_trip_projection(tmp_path: Path) -> None:
    source = tmp_path / "sample.md"
    source.write_text(
        "# Methods\n\nWater is **bold** and $x_i$ is inline.\n\n- one\n- two\n\n"
        "$$\nE = mc^2\n$$\n\nSee \\cite{smith2020}.\n",
        encoding="utf-8",
    )
    tex = tmp_path / "sample.tex"
    back = tmp_path / "back.md"
    convert(source, tex, force=True)
    convert(tex, back, force=True)
    assert read(source).document.projection() == read(back).document.projection()


def test_preserve_rejects_cross_format_and_unsupported_tex(tmp_path: Path) -> None:
    docx = tmp_path / "sample.docx"
    Document().save(docx)
    with pytest.raises(ConversionError, match="preserve fidelity"):
        convert(docx, tmp_path / "out.md", fidelity="preserve")
    assert not (tmp_path / "out.md").exists()

    source = tmp_path / "custom.tex"
    source.write_text(
        "\\documentclass{article}\n\\begin{document}\n\\foobar{hidden}\n\\end{document}\n",
        encoding="utf-8",
    )
    output = tmp_path / "custom.md"
    with pytest.raises(ConversionError, match="loss policy"):
        convert(source, output, fidelity="preserve")
    assert not output.exists()


def test_unchanged_docx_is_byte_identical(tmp_path: Path) -> None:
    source = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("Unchanged sentence.")
    document.save(source)
    output = tmp_path / "copy.docx"
    before = source.read_bytes()
    result = convert(source, output, fidelity="preserve", force=True)
    assert result.execution_mode == "EXACT_REUSE"
    assert output.read_bytes() == before
    assert source.read_bytes() == before


def test_write_unchanged_markdown_matches_direct_render(tmp_path: Path) -> None:
    source = tmp_path / "sample.md"
    source.write_text("Plain sentence.\n", encoding="utf-8")
    loaded = read(source)
    via_write = tmp_path / "write.docx"
    via_convert = tmp_path / "convert.docx"
    write(loaded, via_write, force=True)
    convert(source, via_convert, force=True)
    assert _docx_view(via_write) == _docx_view(via_convert)


def test_style_profile_yaml_and_restyle_keeps_text(tmp_path: Path) -> None:
    profile_path = tmp_path / "style.yaml"
    profile_path.write_text(
        "schema: docforge.style.v1\n"
        "roles:\n"
        "  body:\n"
        "    word_style: Normal\n"
        "    font:\n"
        "      latin: Times New Roman\n"
        "      east_asia: SimSun\n"
        "      size_pt: 14\n"
        "    paragraph:\n"
        "      alignment: justify\n"
        "      line_spacing: 1.15\n"
        "  caption:\n"
        "    word_style: Caption\n"
        "    font: {size_pt: 9}\n"
        "page:\n"
        "  size: A4\n"
        "  margins_mm: {top: 25, bottom: 25, left: 20, right: 20}\n"
        "policy:\n"
        "  missing_style: create\n"
        "  direct_formatting: clear_conflicting\n",
        encoding="utf-8",
    )
    profile = load_profile(profile_path)
    assert profile.roles["body"].font.size_pt == 14
    assert profile.roles["body"].font.east_asia == "SimSun"
    source = _fixture(tmp_path / "styled.docx")
    output = tmp_path / "restyled.docx"
    again = tmp_path / "again.docx"
    summary = restyle_docx(source, output, profile, force=True)
    assert summary["content_unchanged"] is True
    assert _texts(source) == _texts(output)
    resolved = resolved_paragraph_styles(output)
    body = next(item for item in resolved if item["text"].startswith("Body"))
    caption = next(item for item in resolved if item["text"] == "Caption text")
    assert body["size_pt"] == 14
    assert body["latin"] == "Times New Roman"
    assert body["east_asia"] == "SimSun"
    assert body["italic"] is True
    assert body["subscript"] is True
    assert _page_size(output) == ("11906", "16838")
    assert caption["size_pt"] == 9
    assert caption["text"] == "Caption text"
    restyle_docx(output, again, profile, force=True)
    assert resolved_paragraph_styles(output) == resolved_paragraph_styles(again)
    assert cli_main(["restyle", str(source), str(tmp_path / "cli.docx"), "--profile", str(profile_path), "--force"]) == 0


def test_convert_style_profile_beats_historical_run_defaults(tmp_path: Path) -> None:
    source = tmp_path / "body.md"
    source.write_text("Body sentence.\n", encoding="utf-8")
    legacy = tmp_path / "legacy.docx"
    convert(source, legacy, force=True)
    legacy_runs = [row for row in _direct_run_rows(legacy) if row["text"] == "Body sentence."]
    assert legacy_runs
    assert legacy_runs[0]["sz"] == "24"
    assert legacy_runs[0]["color"] == "000000"
    styled = tmp_path / "styled.docx"
    result = convert(source, styled, style={
        "schema": "docforge.style.v1",
        "roles": {"body": {"word_style": "Normal", "font": {"size_pt": 15, "color": "0000FF", "latin": "Georgia"}}},
        "policy": {"missing_style": "error", "direct_formatting": "clear_conflicting"},
    }, force=True)
    assert result.execution_mode == "SOURCE_PATCH"
    body = next(item for item in resolved_paragraph_styles(styled) if item["text"] == "Body sentence.")
    assert body["size_pt"] == 15
    assert body["latin"] == "Georgia"
    assert body["direct_size"] is False
    styled_runs = [row for row in _direct_run_rows(styled) if row["text"] == "Body sentence."]
    assert styled_runs
    assert styled_runs[0]["sz"] is None
    assert (styled_runs[0]["color"] or "").upper() != "000000"


def _direct_run_rows(path: Path) -> list[dict[str, str | None]]:
    namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    root = etree.fromstring(zipfile.ZipFile(path).read("word/document.xml"))
    rows = []
    for run in root.findall(f".//{{{namespace}}}r"):
        properties = run.find(f"{{{namespace}}}rPr")
        size = properties.find(f"{{{namespace}}}sz") if properties is not None else None
        color = properties.find(f"{{{namespace}}}color") if properties is not None else None
        rows.append({
            "text": "".join(node.text or "" for node in run.findall(f"{{{namespace}}}t")),
            "sz": size.get(f"{{{namespace}}}val") if size is not None else None,
            "color": color.get(f"{{{namespace}}}val") if color is not None else None,
        })
    return rows


def test_direct_formatting_policies(tmp_path: Path) -> None:
    source = _fixture(tmp_path / "direct.docx")
    preserved = tmp_path / "preserved.docx"
    cleared = tmp_path / "cleared.docx"
    profile = {
        "schema": "docforge.style.v1",
        "roles": {"body": {"word_style": "Normal", "font": {"size_pt": 16}}},
        "policy": {"direct_formatting": "preserve", "missing_style": "error"},
    }
    restyle_docx(source, preserved, profile, force=True)
    body = next(item for item in resolved_paragraph_styles(preserved) if item["text"].startswith("Body"))
    assert body["direct_size"] is True
    assert body["size_pt"] == 20
    profile["policy"]["direct_formatting"] = "clear_all"
    restyle_docx(source, cleared, profile, force=True)
    body = next(item for item in resolved_paragraph_styles(cleared) if item["text"].startswith("Body"))
    assert body["direct_italic"] is False
    assert body["subscript"] is True
    assert body["size_pt"] == 16


def test_style_inheritance_and_missing_role(tmp_path: Path) -> None:
    source = _fixture(tmp_path / "inherit.docx")
    with pytest.raises(ConversionError, match="no Word style"):
        restyle_docx(source, tmp_path / "nope.docx", {
            "schema": "docforge.style.v1",
            "roles": {"missing": {"word_style": "NotAStyle", "font": {"size_pt": 12}}},
            "policy": {"missing_style": "error"},
        })
    assert not (tmp_path / "nope.docx").exists()
    created = tmp_path / "created.docx"
    summary = restyle_docx(source, created, {
        "schema": "docforge.style.v1",
        "roles": {"missing": {"word_style": "NotAStyle", "font": {"size_pt": 12}}},
        "policy": {"missing_style": "create"},
    }, force=True)
    assert created.is_file()
    assert summary["roles"] == []
    assert summary["unused"] == ["missing"]
    assert all(item["style_id"] != "NotAStyle" for item in resolved_paragraph_styles(created))


def test_created_style_is_attached_to_selected_paragraphs(tmp_path: Path) -> None:
    source = _fixture(tmp_path / "styled.docx")
    output = tmp_path / "selected.docx"
    summary = restyle_docx(source, output, {
        "schema": "docforge.style.v1",
        "roles": {
            "pullquote": {
                "word_style": "PullQuote",
                "select": "Normal",
                "font": {"size_pt": 13, "latin": "Georgia"},
            }
        },
        "policy": {"missing_style": "create", "direct_formatting": "clear_conflicting"},
    }, force=True)
    assert summary["roles"] == ["pullquote"]
    assert summary["unused"] == []
    rows = resolved_paragraph_styles(output)
    body = next(item for item in rows if str(item["text"]).startswith("Body"))
    caption = next(item for item in rows if item["text"] == "Caption text")
    assert body["style_id"] == "PullQuote"
    assert body["size_pt"] == 13
    assert body["latin"] == "Georgia"
    assert caption["style_id"] != "PullQuote"
    assert caption["text"] == "Caption text"


def test_style_template_is_the_style_source_and_fails_closed(tmp_path: Path) -> None:
    template = _house_template(tmp_path / "house.docx")
    source = tmp_path / "plain.docx"
    document = Document()
    document.add_paragraph("Keep me")
    document.save(source)
    output = tmp_path / "from-template.docx"
    summary = restyle_docx(source, output, {
        "schema": "docforge.style.v1",
        "template": str(template),
        "roles": {"body": {"word_style": "Normal"}},
        "policy": {"direct_formatting": "preserve", "missing_style": "error"},
    }, force=True)
    assert summary["template"] == str(template)
    assert "Normal" in summary["imported_styles"]
    assert _texts(source) == _texts(output) == ("Keep me",)
    body = next(item for item in resolved_paragraph_styles(output) if item["text"] == "Keep me")
    assert body["size_pt"] == 16
    assert body["latin"] == "Times New Roman"
    overlay = tmp_path / "overlay.docx"
    restyle_docx(source, overlay, {
        "schema": "docforge.style.v1",
        "template": str(template),
        "roles": {"body": {"word_style": "Normal", "font": {"size_pt": 18, "latin": "Georgia"}}},
        "policy": {"direct_formatting": "clear_conflicting", "missing_style": "error"},
    }, force=True)
    styled = next(item for item in resolved_paragraph_styles(overlay) if item["text"] == "Keep me")
    assert styled["size_pt"] == 18
    assert styled["latin"] == "Georgia"
    profile_path = tmp_path / "style.yaml"
    profile_path.write_text(
        "schema: docforge.style.v1\n"
        "template: house.docx\n"
        "roles:\n"
        "  body:\n"
        "    word_style: Normal\n"
        "policy:\n"
        "  direct_formatting: preserve\n"
        "  missing_style: error\n",
        encoding="utf-8",
    )
    via_yaml = tmp_path / "yaml.docx"
    loaded = load_profile(profile_path)
    assert loaded.template == "house.docx"
    assert loaded.template_base == profile_path.resolve().parent
    restyle_docx(source, via_yaml, profile_path, force=True)
    yaml_body = next(item for item in resolved_paragraph_styles(via_yaml) if item["text"] == "Keep me")
    assert yaml_body["size_pt"] == 16
    assert yaml_body["latin"] == "Times New Roman"
    converted = tmp_path / "converted.docx"
    result = convert(source, converted, style=profile_path, force=True)
    assert result.execution_mode == "SOURCE_PATCH"
    assert any(item.code == "style.template" for item in result.loss_report.diagnostics)
    assert _texts(converted) == ("Keep me",)
    kept = tmp_path / "kept.docx"
    kept.write_bytes(b"SENTINEL")
    missing = tmp_path / "missing.docx"
    no_styles = tmp_path / "no-styles.docx"
    _repack(template, no_styles, drop={"word/styles.xml"})
    bad_styles = tmp_path / "bad-styles.docx"
    _repack(template, bad_styles, replace={"word/styles.xml": b"<styles"})
    for profile in (
        {"schema": "docforge.style.v1", "template": "missing-template.docx", "roles": {"body": {"word_style": "Normal"}}},
        {"schema": "docforge.style.v1", "template": str(missing), "roles": {"body": {"word_style": "Normal"}}},
        {"schema": "docforge.style.v1", "template": str(_broken_template(tmp_path / "broken.docx")), "roles": {"body": {"word_style": "Normal"}}},
        {"schema": "docforge.style.v1", "template": str(no_styles), "roles": {"body": {"word_style": "Normal"}}},
        {"schema": "docforge.style.v1", "template": str(bad_styles), "roles": {"body": {"word_style": "Normal"}}},
    ):
        with pytest.raises(ConversionError, match="template") as caught:
            restyle_docx(source, kept, profile, force=True)
        assert caught.value.code == "STYLE_TEMPLATE"
        assert kept.read_bytes() == b"SENTINEL"


def test_docx_projection_keeps_body_order(tmp_path: Path) -> None:
    png = _png(tmp_path / "dot.png")
    document = Document()
    document.add_paragraph("Before")
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.text = "Cell"
    nested = cell.add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested"
    paragraph = document.add_paragraph()
    paragraph.add_run("Left")
    paragraph.add_run().add_picture(str(png))
    paragraph.add_run("Right")
    path = tmp_path / "order.docx"
    document.save(path)
    loaded = read(path)
    sequence = [(block.kind, block.plain()) for block in loaded.document.blocks]
    assert [kind for kind, _text in sequence] == [
        "paragraph", "table", "paragraph", "figure", "paragraph",
    ], sequence
    assert [text for kind, text in sequence if kind == "paragraph"] == ["Before", "Left", "Right"]
    table_block = loaded.document.blocks[1]
    cell_block = table_block.children[0].children[0]
    child_kinds = [child.kind for child in cell_block.children]
    assert child_kinds.index("paragraph") < child_kinds.index("table")
    assert "Cell" in cell_block.plain()
    assert "Nested" in cell_block.plain()
    figure = next(block for block in loaded.document.blocks if block.kind == "figure")
    resource = loaded.document.resource(figure.attr("resource"))
    assert resource is not None and resource.sha256
    assert LossReport("error", loaded.diagnostics).allows_publication()


def test_preserve_write_rejects_stale_and_edited_docx(tmp_path: Path) -> None:
    source = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("Original sentence.")
    document.save(source)
    loaded = read(source)
    document.add_paragraph("Changed on disk.")
    document.save(source)
    stale = tmp_path / "stale.docx"
    stale.write_bytes(b"SENTINEL-STALE")
    with pytest.raises(ConversionError, match="changed after"):
        write(loaded, stale, fidelity="preserve", force=True)
    assert stale.read_bytes() == b"SENTINEL-STALE"

    fresh = tmp_path / "fresh.docx"
    document = Document()
    document.add_paragraph("Original sentence.")
    document.save(fresh)
    loaded = read(fresh)
    edited = replace(
        loaded,
        document=replace(loaded.document, blocks=loaded.document.blocks + (
            BlockNode("extra", "opaque", attrs=(("element", "widget"),)),
        )),
    )
    edited_output = tmp_path / "edited.docx"
    edited_output.write_bytes(b"SENTINEL-EDITED")
    with pytest.raises(ConversionError, match="preserve fidelity") as caught:
        write(edited, edited_output, fidelity="preserve", force=True)
    assert caught.value.code == "UNSUPPORTED_PATCH"
    assert edited_output.read_bytes() == b"SENTINEL-EDITED"


def test_source_patch_confines_paragraph_edits(tmp_path: Path) -> None:
    png = _png(tmp_path / "dot.png")
    document = Document()
    document.add_paragraph("Alpha")
    document.add_paragraph("Beta")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Cell"
    document.add_paragraph("Gamma")
    document.add_picture(str(png))
    source = tmp_path / "bundle.docx"
    document.save(source)
    loaded = read(source)
    assert loaded.document.blocks[0].provenance is not None
    assert str(loaded.document.blocks[0].provenance.anchor).startswith("word/document.xml:child:")
    blocks = []
    for block in loaded.document.blocks:
        if block.plain() == "Beta":
            blocks.append(replace(block, inlines=(
                text_node("Beta "),
                InlineNode("emph", children=(text_node("edited"),)),
            )))
        else:
            blocks.append(block)
    replaced = replace(loaded, document=replace(loaded.document, blocks=tuple(blocks)))
    output = tmp_path / "patched.docx"
    result = write(replaced, output, fidelity="preserve", force=True)
    assert result.execution_mode == "SOURCE_PATCH"
    assert any(item.code == "patch.replaced" for item in result.loss_report.diagnostics)
    assert _changed_parts(source, output) == ["word/document.xml"]
    assert _paragraph_rows(output)[0] == _paragraph_rows(source)[0]
    beta_before = _paragraph_rows(source)[1]
    beta_after = _paragraph_rows(output)[1]
    assert beta_before[0] == "Beta"
    assert beta_after[0] == "Beta edited"
    assert beta_before[1] == beta_after[1]
    assert _paragraph_rows(output)[2][0] == "Gamma"
    assert _media_bytes(source) == _media_bytes(output)

    removed_blocks = tuple(block for block in blocks if block.plain() != "Gamma")
    removed = replace(loaded, document=replace(loaded.document, blocks=removed_blocks))
    deleted = tmp_path / "deleted.docx"
    deleted_result = write(removed, deleted, fidelity="preserve", force=True)
    assert deleted_result.execution_mode == "SOURCE_PATCH"
    assert [row[0] for row in _paragraph_rows(deleted) if row[0]] == ["Alpha", "Beta edited"]
    assert Document(deleted).tables[0].cell(0, 0).text == "Cell"
    assert _changed_parts(source, deleted) == ["word/document.xml"]
    assert _media_bytes(source) == _media_bytes(deleted)

    inserted_blocks = (*blocks[:1], BlockNode("added", "paragraph", inlines=(text_node("Inserted"),)), *blocks[1:])
    inserted = replace(loaded, document=replace(loaded.document, blocks=tuple(inserted_blocks)))
    added = tmp_path / "inserted.docx"
    added_result = write(inserted, added, fidelity="preserve", force=True)
    assert added_result.execution_mode == "SOURCE_PATCH"
    assert [row[0] for row in _paragraph_rows(added) if row[0]] == ["Alpha", "Inserted", "Beta edited", "Gamma"]
    assert _changed_parts(source, added) == ["word/document.xml"]

    table_block = next(block for block in loaded.document.blocks if block.kind == "table")
    cell = table_block.children[0].children[0]
    paragraph = cell.children[0]
    changed_cell = replace(cell, children=(replace(paragraph, inlines=(text_node("Changed"),)),))
    changed_row = replace(table_block.children[0], children=(changed_cell,))
    changed_table = replace(table_block, children=(changed_row,))
    refused_blocks = tuple(changed_table if block.kind == "table" else block for block in loaded.document.blocks)
    refused = replace(loaded, document=replace(loaded.document, blocks=refused_blocks))
    sentinel = tmp_path / "sentinel.docx"
    sentinel.write_bytes(b"SENTINEL-PATCH")
    with pytest.raises(ConversionError, match="preserve fidelity") as caught:
        write(refused, sentinel, fidelity="preserve", force=True)
    assert caught.value.code == "UNSUPPORTED_PATCH"
    assert sentinel.read_bytes() == b"SENTINEL-PATCH"
    plain = tmp_path / "plain.docx"
    plain_doc = Document()
    plain_doc.add_paragraph("Alpha")
    plain_doc.save(plain)
    plain_loaded = read(plain)
    plain_edited = replace(
        plain_loaded,
        document=replace(plain_loaded.document, blocks=(
            replace(plain_loaded.document.blocks[0], inlines=(text_node("Alpha edited"),)),
        )),
    )
    semantic = tmp_path / "rebuilt.docx"
    rebuilt = write(plain_edited, semantic, force=True)
    assert rebuilt.execution_mode == "REBUILD"
    preserved = tmp_path / "preserved.docx"
    patched = write(plain_edited, preserved, fidelity="preserve", force=True)
    assert patched.execution_mode == "SOURCE_PATCH"
    assert Document(preserved).paragraphs[0].text == "Alpha edited"


def test_source_patch_body_text_keeps_projected_header(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("Alpha")
    header = document.sections[0].header
    header.is_linked_to_previous = False
    header.paragraphs[0].text = "Running head"
    source = tmp_path / "headed.docx"
    document.save(source)
    loaded = read(source)
    story = [
        block for block in loaded.document.blocks
        if block.provenance is not None and block.provenance.anchor.startswith("word/header")
    ]
    assert [block.plain() for block in story] == ["Running head"]
    body = next(block for block in loaded.document.blocks if block.plain() == "Alpha")
    edited_blocks = tuple(
        replace(block, inlines=(text_node("Alpha edited"),)) if block is body else block
        for block in loaded.document.blocks
    )
    edited = replace(loaded, document=replace(loaded.document, blocks=edited_blocks))
    output = tmp_path / "patched.docx"
    result = write(edited, output, fidelity="preserve", force=True)
    assert result.execution_mode == "SOURCE_PATCH"
    assert _changed_parts(source, output) == ["word/document.xml"]
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(output) as after:
        headers = [name for name in before.namelist() if name.startswith("word/header")]
        assert headers
        assert all(before.read(name) == after.read(name) for name in headers)
        assert b"Running head" in before.read(headers[0])
    assert [row[0] for row in _paragraph_rows(output) if row[0]] == ["Alpha edited"]
    assert Document(output).sections[0].header.paragraphs[0].text == "Running head"


def _changed_parts(left: Path, right: Path) -> list[str]:
    with zipfile.ZipFile(left) as source, zipfile.ZipFile(right) as output:
        names = source.namelist()
        assert names == output.namelist()
        return [name for name in names if source.read(name) != output.read(name)]


def _paragraph_rows(path: Path) -> list[tuple[str, bytes]]:
    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    root = etree.fromstring(zipfile.ZipFile(path).read("word/document.xml"))
    rows = []
    for paragraph in root.findall("w:body/w:p", namespace):
        text = "".join(node.text or "" for node in paragraph.findall(".//w:t", namespace))
        properties = paragraph.find("w:pPr", namespace)
        rows.append((text, etree.tostring(properties) if properties is not None else b""))
    return rows


def _media_bytes(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as package:
        return {name: package.read(name) for name in package.namelist() if name.startswith("word/media/")}


def test_loss_policy_error_publishes_nothing(tmp_path: Path) -> None:
    source = _textbox_docx(tmp_path / "native.docx")
    loaded = read(source)
    assert any(block.plain() == "Native box" for block in loaded.document.walk())
    assert any(item.code == "docx.textboxes" and item.details.get("category") == "loss" for item in loaded.diagnostics)
    markdown = tmp_path / "native.md"
    markdown.write_bytes(b"SENTINEL-MD")
    with pytest.raises(ConversionError, match="loss policy"):
        convert(source, markdown, loss_policy="error", force=True)
    assert markdown.read_bytes() == b"SENTINEL-MD"

    opaque = SemanticDocument(blocks=(
        BlockNode("kept", "paragraph", inlines=(text_node("Kept"),)),
        BlockNode("hidden", "opaque", attrs=(("origin", "widget"),)),
    ))
    for name in ("opaque.md", "opaque.tex"):
        destination = tmp_path / name
        destination.write_bytes(b"SENTINEL-OPAQUE")
        with pytest.raises(ConversionError, match="loss policy"):
            write(opaque, destination, loss_policy="error", force=True)
        assert destination.read_bytes() == b"SENTINEL-OPAQUE"


def test_invalid_model_does_not_replace_output(tmp_path: Path) -> None:
    document = SemanticDocument(blocks=(
        BlockNode("dup", "paragraph", inlines=(text_node("One"),)),
        BlockNode("dup", "paragraph", inlines=(text_node("Two"),)),
    ))
    for name in ("invalid.md", "invalid.tex"):
        destination = tmp_path / name
        destination.write_bytes(b"SENTINEL-INVALID")
        with pytest.raises(ConversionError, match="validation"):
            write(document, destination, force=True)
        assert destination.read_bytes() == b"SENTINEL-INVALID"
    for name in ("absent.md", "absent.tex"):
        destination = tmp_path / "missing" / name
        with pytest.raises(ConversionError, match="validation"):
            write(document, destination, force=True)
        assert not destination.exists()
        assert not destination.parent.exists()


def test_audit_resolves_images_headers_and_relative_targets(tmp_path: Path) -> None:
    png = _png(tmp_path / "dot.png")
    document = Document()
    document.add_paragraph("See")
    document.add_picture(str(png))
    header = document.sections[0].header
    header.is_linked_to_previous = False
    header.paragraphs[0].text = "Running head"
    source = tmp_path / "picture.docx"
    document.save(source)
    assert audit(source).ok

    media = [name for name in zipfile.ZipFile(source).namelist() if name.startswith("word/media/")]
    assert media
    missing = tmp_path / "missing.docx"
    _repack(source, missing, drop=set(media))
    missing_report = audit(missing)
    assert not missing_report.ok
    assert any(item.code == "docx.missing_target" for item in missing_report.diagnostics)

    absolute = tmp_path / "absolute.docx"
    _absolutize_media_targets(source, absolute)
    assert audit(absolute).ok

    relative = tmp_path / "relative.docx"
    media_name = Path(media[0]).name
    _repack(source, relative, replace={
        "word/charts/_rels/chart1.xml.rels": (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"'
            f' Target="../media/{media_name}"/>'
            "</Relationships>"
        ).encode("utf-8"),
    })
    assert audit(relative).ok

    dangling = tmp_path / "dangling.docx"
    _retarget_embed(source, dangling, "rIdMissing")
    dangling_report = audit(dangling)
    assert not dangling_report.ok
    assert any(item.code == "docx.dangling_relationship" for item in dangling_report.diagnostics)

    headers = [name for name in zipfile.ZipFile(source).namelist() if name.startswith("word/header")]
    assert headers
    broken_header = tmp_path / "header.docx"
    _repack(source, broken_header, drop=set(headers))
    header_report = audit(broken_header)
    assert not header_report.ok
    assert any(item.code == "docx.header_footer" for item in header_report.diagnostics)


def test_audit_package_relationships_and_encoded_targets(tmp_path: Path) -> None:
    png = _png(tmp_path / "dot.png")
    document = Document()
    document.add_paragraph("See")
    document.add_picture(str(png))
    header = document.sections[0].header
    header.is_linked_to_previous = False
    header.paragraphs[0].text = "Running head"
    source = tmp_path / "picture.docx"
    document.save(source)
    media = next(name for name in zipfile.ZipFile(source).namelist() if name.startswith("word/media/"))
    media_name = Path(media).name

    package_absolute = tmp_path / "package-absolute.docx"
    _prefix_relationship_targets(source, package_absolute, "_rels/.rels")
    assert audit(package_absolute).ok

    missing_core = tmp_path / "missing-core.docx"
    _repack(source, missing_core, drop={"docProps/core.xml"})
    core_report = audit(missing_core)
    assert not core_report.ok
    assert any(item.code == "docx.missing_target" and item.details.get("owner") == "[package]" for item in core_report.diagnostics)

    encoded = tmp_path / "encoded.docx"
    _retarget_media(source, encoded, part_name="image 1.png", target="media/image%201.png")
    assert audit(encoded).ok

    swapped = tmp_path / "case.docx"
    _retarget_media(source, swapped, part_name=media_name, target=f"media/{media_name.swapcase()}")
    assert audit(swapped).ok

    dotted = tmp_path / "dotted.docx"
    _retarget_media(source, dotted, part_name=media_name, target=f"media/../media/{media_name}")
    assert audit(dotted).ok

    header_media = tmp_path / "header-media.docx"
    _repack(source, header_media, replace={
        "word/_rels/header1.xml.rels": (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"'
            f' Target="media/{media_name}"/>'
            "</Relationships>"
        ).encode("utf-8"),
    })
    assert audit(header_media).ok

    escaped = tmp_path / "escaped.docx"
    _retarget_media(source, escaped, part_name=media_name, target="../../outside.png")
    assert not audit(escaped).ok

    empty_target = tmp_path / "empty-target.docx"
    _add_relationship(source, empty_target, 'Id="rIdEmpty" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target=""')
    assert not audit(empty_target).ok

    linked = tmp_path / "linked.docx"
    _rename_blip_attribute(source, linked, "link")
    assert audit(linked).ok

    dangling_link = tmp_path / "dangling-link.docx"
    _rename_blip_attribute(source, dangling_link, "link", value="rIdMissing")
    link_report = audit(dangling_link)
    assert not link_report.ok
    assert any(item.code == "docx.dangling_relationship" for item in link_report.diagnostics)

    diagram = tmp_path / "diagram.docx"
    _add_relationship_attribute(source, diagram, "dm", "rIdMissing")
    diagram_report = audit(diagram)
    assert not diagram_report.ok
    assert any(item.code == "docx.dangling_relationship" for item in diagram_report.diagnostics)

    strict = tmp_path / "strict.docx"
    _add_relationship_attribute(
        source, strict, "embed", "rIdMissing",
        namespace="http://purl.oclc.org/ooxml/officeDocument/relationships",
    )
    assert not audit(strict).ok

    office = tmp_path / "office.docx"
    _add_relationship_attribute(
        source, office, "relid", "rIdMissing",
        namespace="urn:schemas-microsoft-com:office:office",
    )
    assert not audit(office).ok


def test_restyle_audit_failure_keeps_existing_output(tmp_path: Path) -> None:
    source = tmp_path / "picture.docx"
    document = Document()
    document.add_paragraph("Body")
    document.add_picture(str(_png(tmp_path / "dot.png")))
    document.save(source)
    broken = tmp_path / "broken.docx"
    _retarget_embed(source, broken, "rIdMissing")
    output = tmp_path / "kept.docx"
    output.write_bytes(b"SENTINEL-RESTYLE")
    with pytest.raises(ConversionError, match="audit"):
        restyle_docx(broken, output, {
            "schema": "docforge.style.v1",
            "roles": {"body": {"word_style": "Normal", "font": {"size_pt": 12}}},
        }, force=True)
    assert output.read_bytes() == b"SENTINEL-RESTYLE"


def test_markdown_latex_escapes_links_images_and_copies_media(tmp_path: Path) -> None:
    source_dir = tmp_path / "in"
    source_dir.mkdir()
    png = _png(source_dir / "plot.png")
    source = source_dir / "sample.md"
    source.write_text(
        "See [site](https://example.com) and symbols { } $ ~ ^ & % # _ \\\\ done.\n\n"
        "![Plot](plot.png)\n",
        encoding="utf-8",
    )
    output_dir = tmp_path / "out"
    output = output_dir / "sample.tex"
    convert(source, output, force=True)
    text = output.read_text(encoding="utf-8")
    assert "\\usepackage{graphicx}" in text
    assert "\\usepackage{hyperref}" in text
    assert "\\usepackage{textcomp}" in text
    assert text.index("\\usepackage{graphicx}") < text.index("\\usepackage{hyperref}")
    assert "\\href{https://example.com}{site}" in text
    assert "\\includegraphics{plot.png}" in text
    assert r"\{" in text and r"\}" in text and r"\$" in text
    assert r"\&" in text and r"\%" in text and r"\#" in text and r"\_" in text
    assert r"\textasciitilde{}" in text and r"\textasciicircum{}" in text
    assert r"\textbackslash{}" in text
    assert str(tmp_path) not in text
    assert (output_dir / "plot.png").read_bytes() == png.read_bytes()
    engine = shutil.which("pdflatex") or shutil.which("lualatex") or shutil.which("xelatex")
    if engine:
        completed = subprocess.run(
            [engine, "-interaction=nonstopmode", "-halt-on-error", output.name],
            cwd=output_dir, capture_output=True, text=True, timeout=120, check=False,
        )
        assert completed.returncode == 0, completed.stdout[-2000:] + completed.stderr[-2000:]


def test_markdown_latex_rebases_media_outside_the_source_tree(tmp_path: Path) -> None:
    shared = tmp_path / "shared"
    shared.mkdir()
    png = _png(shared / "plot.png")
    source_dir = tmp_path / "chapter"
    source_dir.mkdir()
    nested_dir = source_dir / "figs"
    nested_dir.mkdir()
    nested = _png(nested_dir / "nested.png")
    source = source_dir / "sample.md"
    source.write_text(
        "See [site](https://example.com) and symbols { } $ ~ ^.\n\n"
        "![Plot](../shared/plot.png)\n\n"
        "![Nested](figs/nested.png)\n",
        encoding="utf-8",
    )
    loaded = read(source)
    assert str(tmp_path) not in str(loaded.document.to_dict())
    output = tmp_path / "build" / "sample.tex"
    convert(source, output, force=True)
    text = output.read_text(encoding="utf-8")
    assert "\\includegraphics{plot.png}" in text
    assert "\\includegraphics{figs/nested.png}" in text
    assert ".." not in text
    assert str(tmp_path) not in text
    assert (output.parent / "plot.png").read_bytes() == png.read_bytes()
    assert (output.parent / "figs" / "nested.png").read_bytes() == nested.read_bytes()
    engine = shutil.which("pdflatex") or shutil.which("lualatex") or shutil.which("xelatex")
    if engine:
        completed = subprocess.run(
            [engine, "-interaction=nonstopmode", "-halt-on-error", output.name],
            cwd=output.parent, capture_output=True, text=True, timeout=120, check=False,
        )
        assert completed.returncode == 0, completed.stdout[-2000:] + completed.stderr[-2000:]


def test_missing_figure_does_not_publish_export(tmp_path: Path) -> None:
    source_dir = tmp_path / "in"
    source_dir.mkdir()
    png = _png(source_dir / "plot.png")
    source = source_dir / "sample.md"
    source.write_text("![Plot](plot.png)\n", encoding="utf-8")
    loaded = read(source)
    edited = replace(
        loaded,
        document=replace(loaded.document, blocks=loaded.document.blocks + (
            BlockNode("note", "paragraph", inlines=(text_node("Note"),)),
        )),
    )
    png.unlink()
    for name in ("sample.md", "sample.tex"):
        destination = tmp_path / "out" / name
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_bytes(b"SENTINEL")
        sibling = destination.parent / "plot.png"
        sibling.write_bytes(b"KEEP")
        with pytest.raises(ConversionError, match="audit"):
            write(edited, destination, force=True)
        assert destination.read_bytes() == b"SENTINEL"
        assert sibling.read_bytes() == b"KEEP"


def test_public_restyle_attaches_selected_style_and_reports_unused(tmp_path: Path) -> None:
    source = _fixture(tmp_path / "styled.docx")
    unused_output = tmp_path / "unused.docx"
    unused = restyle(source, unused_output, profile={
        "schema": "docforge.style.v1",
        "roles": {"missing": {"word_style": "NotAStyle", "font": {"size_pt": 12}}},
        "policy": {"missing_style": "create"},
    }, force=True)
    assert any(item.code == "style.unused" and item.details.get("role") == "missing" for item in unused.loss_report.diagnostics)
    assert not any(item.code == "style.applied" for item in unused.loss_report.diagnostics)
    assert all(item["style_id"] != "NotAStyle" for item in resolved_paragraph_styles(unused_output))
    assert _texts(source) == _texts(unused_output)

    selected_output = tmp_path / "selected.docx"
    selected = restyle(source, selected_output, profile={
        "schema": "docforge.style.v1",
        "roles": {
            "pullquote": {
                "word_style": "PullQuote",
                "select": "Normal",
                "font": {"size_pt": 13, "latin": "Georgia"},
            }
        },
        "policy": {"missing_style": "create", "direct_formatting": "clear_conflicting"},
    }, force=True)
    assert any(item.code == "style.applied" and item.details.get("role") == "pullquote" for item in selected.loss_report.diagnostics)
    rows = resolved_paragraph_styles(selected_output)
    body = next(item for item in rows if str(item["text"]).startswith("Body"))
    caption = next(item for item in rows if item["text"] == "Caption text")
    assert body["style_id"] == "PullQuote"
    assert body["size_pt"] == 13
    assert body["latin"] == "Georgia"
    assert caption["style_id"] != "PullQuote"
    assert _texts(source) == _texts(selected_output)

    kept = tmp_path / "kept.docx"
    kept.write_bytes(b"SENTINEL")
    with pytest.raises(ConversionError, match="template") as caught:
        restyle(source, kept, profile={
            "schema": "docforge.style.v1",
            "template": "template.docx",
            "roles": {"body": {"word_style": "Normal", "font": {"size_pt": 12}}},
        }, force=True)
    assert caught.value.code == "STYLE_TEMPLATE"
    assert kept.read_bytes() == b"SENTINEL"


def _docx_view(path: Path) -> list[tuple[str, list[tuple[str, bool, bool]]]]:
    document = Document(path)
    return [
        (paragraph.text, [(run.text, bool(run.bold), bool(run.italic)) for run in paragraph.runs])
        for paragraph in document.paragraphs
    ]


def _page_size(path: Path) -> tuple[str | None, str | None]:
    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    root = etree.fromstring(zipfile.ZipFile(path).read("word/document.xml"))
    size = root.find(".//w:sectPr/w:pgSz", namespace)
    assert size is not None
    return size.get(f"{{{namespace['w']}}}w"), size.get(f"{{{namespace['w']}}}h")


def _texts(path: Path) -> tuple[str, ...]:
    return tuple(paragraph.text for paragraph in Document(path).paragraphs)


def _house_template(path: Path) -> Path:
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(16)
    document.add_paragraph("Sample")
    document.save(path)
    return path


def _broken_template(path: Path) -> Path:
    path.write_bytes(b"this is not a word package")
    return path


def _fixture(path: Path) -> Path:
    document = Document()
    caption_style = document.styles["Caption"]
    caption_style.font.size = Pt(9)
    body = document.add_paragraph()
    body.style = document.styles["Normal"]
    lead = body.add_run("Body ")
    lead.font.size = Pt(20)
    term = body.add_run("term")
    term.italic = True
    term.font.size = Pt(20)
    subscript = body.add_run("2")
    subscript.font.subscript = True
    subscript.font.size = Pt(20)
    caption = document.add_paragraph("Caption text")
    caption.style = caption_style
    document.save(path)
    return path


_PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
_DRAWING_NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}


def _png(path: Path) -> Path:
    image = Image.new("RGB", (2, 2), (10, 20, 30))
    image.save(path, format="PNG")
    return path


def _textbox_docx(path: Path) -> Path:
    document = Document()
    document.add_paragraph("Visible")
    document.element.body.append(parse_xml(
        '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
        ' xmlns:v="urn:schemas-microsoft-com:vml">'
        "<w:r><w:pict><v:shape><v:textbox><w:txbxContent>"
        "<w:p><w:r><w:t>Native box</w:t></w:r></w:p>"
        "</w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>"
    ))
    document.save(path)
    return path


def _repack(source: Path, dest: Path, *, replace: dict[str, bytes] | None = None, drop: set[str] | None = None) -> None:
    replace = replace or {}
    drop = drop or set()
    with zipfile.ZipFile(source) as incoming, zipfile.ZipFile(dest, "w") as outgoing:
        written = set()
        for info in incoming.infolist():
            if info.filename in drop:
                continue
            outgoing.writestr(info, replace.get(info.filename, incoming.read(info.filename)))
            written.add(info.filename)
        for name, payload in replace.items():
            if name not in written and name not in drop:
                outgoing.writestr(name, payload)


def _absolutize_media_targets(source: Path, dest: Path) -> None:
    with zipfile.ZipFile(source) as package:
        root = etree.fromstring(package.read("word/_rels/document.xml.rels"))
    for item in root.findall(f"{{{_PKG_REL}}}Relationship"):
        target = item.get("Target") or ""
        if "media/" in target and not target.startswith("/"):
            item.set("Target", "/word/" + target.removeprefix("./"))
    _repack(source, dest, replace={
        "word/_rels/document.xml.rels": etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True),
    })


def _prefix_relationship_targets(source: Path, dest: Path, rels_name: str) -> None:
    with zipfile.ZipFile(source) as package:
        root = etree.fromstring(package.read(rels_name))
    for item in root.findall(f"{{{_PKG_REL}}}Relationship"):
        target = item.get("Target") or ""
        mode = (item.get("TargetMode") or "").casefold()
        if mode == "external" or not target or target.startswith(("/", "#")) or "://" in target:
            continue
        item.set("Target", "/" + target)
    _repack(source, dest, replace={
        rels_name: etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True),
    })


def _retarget_media(source: Path, dest: Path, *, part_name: str, target: str) -> None:
    with zipfile.ZipFile(source) as package:
        payloads = {name: package.read(name) for name in package.namelist()}
    media = next(name for name in payloads if name.startswith("word/media/") and not name.endswith("/"))
    renamed = str(Path(media).with_name(part_name)).replace("\\", "/")
    payloads[renamed] = payloads.pop(media)
    rels_name = "word/_rels/document.xml.rels"
    root = etree.fromstring(payloads[rels_name])
    for item in root.findall(f"{{{_PKG_REL}}}Relationship"):
        current = item.get("Target") or ""
        if "media/" in current.replace("\\", "/"):
            item.set("Target", target)
    payloads[rels_name] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    with zipfile.ZipFile(dest, "w") as outgoing:
        for name, payload in payloads.items():
            outgoing.writestr(name, payload)


def _add_relationship(source: Path, dest: Path, attributes: str) -> None:
    with zipfile.ZipFile(source) as package:
        root = etree.fromstring(package.read("word/_rels/document.xml.rels"))
    root.append(etree.fromstring(
        f'<Relationship xmlns="{_PKG_REL}" {attributes}/>'
    ))
    _repack(source, dest, replace={
        "word/_rels/document.xml.rels": etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True),
    })


def _rename_blip_attribute(source: Path, dest: Path, local_name: str, *, value: str | None = None) -> None:
    with zipfile.ZipFile(source) as package:
        root = etree.fromstring(package.read("word/document.xml"))
    for blip in root.xpath(".//a:blip", namespaces=_DRAWING_NS):
        current = blip.get(f"{{{_DRAWING_NS['r']}}}embed")
        if not current:
            continue
        del blip.attrib[f"{{{_DRAWING_NS['r']}}}embed"]
        blip.set(f"{{{_DRAWING_NS['r']}}}{local_name}", value or current)
    _repack(source, dest, replace={
        "word/document.xml": etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True),
    })


def _add_relationship_attribute(source: Path, dest: Path, local_name: str, value: str, *, namespace: str | None = None) -> None:
    namespace = namespace or _DRAWING_NS["r"]
    with zipfile.ZipFile(source) as package:
        root = etree.fromstring(package.read("word/document.xml"))
    body = root.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}body")
    assert body is not None
    marker = etree.Element("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p")
    marker.set(f"{{{namespace}}}{local_name}", value)
    body.append(marker)
    _repack(source, dest, replace={
        "word/document.xml": etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True),
    })


def _retarget_embed(source: Path, dest: Path, rel_id: str) -> None:
    with zipfile.ZipFile(source) as package:
        root = etree.fromstring(package.read("word/document.xml"))
    for blip in root.xpath(".//a:blip", namespaces=_DRAWING_NS):
        if blip.get(f"{{{_DRAWING_NS['r']}}}embed"):
            blip.set(f"{{{_DRAWING_NS['r']}}}embed", rel_id)
    _repack(source, dest, replace={
        "word/document.xml": etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True),
    })
