from __future__ import annotations

import json
import shutil
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

from docforge import _pandoc
from docforge.docxdiff import create_tracked_docx
from docforge.docxdiff.redline import visible_text
from docforge.markdown import (
    Block,
    docx_to_markdown,
    normalize_script_boundaries,
    render_blocks_to_doc,
    split_markdown_sections,
)
from docforge.tex import docx_to_tex


PANDOC_AVAILABLE = shutil.which("pandoc") is not None


def _sample_docx(path: Path, image: Path) -> None:
    document = Document()
    document.add_heading("Title", level=1)
    document.add_paragraph("Author")
    document.add_heading("Abstract", level=1)
    paragraph = document.add_paragraph("H")
    superscript = paragraph.add_run("2")
    superscript.font.superscript = True
    document.add_heading("Main", level=1)
    document.add_paragraph("Body")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "2"
    document.add_picture(str(image))
    document.add_heading("References", level=1)
    document.add_paragraph("1. Example reference")
    document.add_heading("Methods", level=1)
    document.add_paragraph("Method")
    document.add_heading("References", level=1)
    document.add_paragraph("51. Method reference")
    document.add_heading("Data availability", level=1)
    document.add_paragraph("Data")
    document.save(path)


@pytest.mark.skipif(not PANDOC_AVAILABLE, reason="Pandoc is required for integration tests")
def test_docx_to_markdown_and_split(tmp_path: Path) -> None:
    image = tmp_path / "source.png"
    Image.new("RGB", (16, 8), "white").save(image)
    source = tmp_path / "source.docx"
    _sample_docx(source, image)
    split_dir = tmp_path / "sections"
    section_map = tmp_path / "section-map.json"
    section_map.write_text(
        json.dumps(
            {
                "sections": [
                    {"file": "00_metadata.md", "kind": "metadata", "end": {"heading": "Abstract"}},
                    {"file": "01_abstract.md", "start": "Abstract", "end": "Main"},
                    {"file": "02_main.md", "start": "Main", "end": {"heading": "References", "occurrence": 1}},
                    {"file": "03_main_references.md", "start": {"heading": "References", "occurrence": 1}, "end": "Methods"},
                    {"file": "04_methods.md", "start": "Methods", "end": {"heading": "References", "occurrence": 2}},
                    {"file": "05_methods_references.md", "start": {"heading": "References", "occurrence": 2}, "end": "Data availability"},
                    {"file": "06_endmatter.md", "start": "Data availability"},
                ]
            }
        ),
        encoding="utf-8",
    )
    result = docx_to_markdown(source, split_dir=split_dir, section_map_path=section_map, force=True)
    assert result.output is None
    assert (split_dir / "02_main.md").read_text(encoding="utf-8").startswith("# Main")
    assert "| A" in (split_dir / "02_main.md").read_text(encoding="utf-8")
    assert "![" in (split_dir / "02_main.md").read_text(encoding="utf-8")
    assert (split_dir / "media").exists()
    metadata = (split_dir / "00_metadata.md").read_text(encoding="utf-8")
    assert "# TITLE" in metadata and "# AUTHOR" in metadata
    assert "C:\\" not in "\n".join(path.read_text(encoding="utf-8") for path in split_dir.glob("*.md"))


@pytest.mark.skipif(not PANDOC_AVAILABLE, reason="Pandoc is required for integration tests")
def test_docx_to_tex_is_standalone_and_uses_png_media(tmp_path: Path) -> None:
    image = tmp_path / "source.png"
    Image.new("RGB", (16, 8), "white").save(image)
    source = tmp_path / "source.docx"
    _sample_docx(source, image)
    output = tmp_path / "output.tex"
    result = docx_to_tex(source, output=output, force=True)
    text = output.read_text(encoding="utf-8")
    assert result.output == output
    assert "\\documentclass" in text
    assert "\\includegraphics" in text
    assert "media/" in text


def test_sup_sub_markup_is_rendered_as_true_scripts() -> None:
    document = render_blocks_to_doc([Block("paragraph", "H<sub>2</sub>O<sup>+−</sup>")])
    runs = document.paragraphs[0].runs
    assert any(run.text == "2" and run.font.subscript for run in runs)
    assert any(run.text == "+−" and run.font.superscript for run in runs)


def test_pandoc_subscript_run_does_not_capture_formula_delimiters() -> None:
    pandoc_markdown = r"(Et<sub>4</sub>N)<sub>2</sub>\[Cu<sub>8</sub>(N<sub>3)18\]</sub>"
    normalized = normalize_script_boundaries(pandoc_markdown)
    assert normalized == r"(Et<sub>4</sub>N)<sub>2</sub>\[Cu<sub>8</sub>(N<sub>3</sub>)<sub>18</sub>\]"

    document = render_blocks_to_doc([Block("paragraph", normalized)])
    run_modes = [
        (
            run.text,
            "subscript" if run.font.subscript else "superscript" if run.font.superscript else "normal",
        )
        for run in document.paragraphs[0].runs
    ]
    assert ("3", "subscript") in run_modes
    assert (")", "normal") in run_modes
    assert ("18", "subscript") in run_modes
    assert any(text.endswith("]") and mode == "normal" for text, mode in run_modes)


def test_split_markdown_sections_handles_duplicate_headings() -> None:
    text = "**Title**\n\n**Abstract**:\n\nA\n\n**Main**\n\nB\n\n**References**\n\n1. R\n\n**Methods**\n\nC\n\n**References**\n\n51. S\n\n**Data availability**\n\nD\n"
    sections = split_markdown_sections(
        text,
        [
            {"file": "00_metadata.md", "kind": "metadata", "end": "Abstract"},
            {"file": "01_abstract.md", "start": "Abstract", "end": "Main"},
            {"file": "02_main.md", "start": "Main", "end": {"heading": "References", "occurrence": 1}},
            {"file": "03_refs.md", "start": {"heading": "References", "occurrence": 1}, "end": "Methods"},
            {"file": "04_methods.md", "start": "Methods", "end": {"heading": "References", "occurrence": 2}},
            {"file": "05_refs.md", "start": {"heading": "References", "occurrence": 2}, "end": "Data availability"},
            {"file": "06_end.md", "start": "Data availability"},
        ],
    )
    assert [name for name, _ in sections] == [
        "00_metadata.md", "01_abstract.md", "02_main.md", "03_refs.md", "04_methods.md", "05_refs.md", "06_end.md"
    ]
    assert sections[2][1].startswith("# Main")
    assert sections[5][1].startswith("# References")


def test_run_pandoc_reports_missing_external_dependency(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    def missing(*args, **kwargs):
        raise FileNotFoundError("pandoc")

    monkeypatch.setattr(_pandoc.subprocess, "run", missing)
    with pytest.raises(RuntimeError, match="pandoc is required"):
        _pandoc.run_pandoc(tmp_path / "input.docx", tmp_path / "out.md", output_format="gfm")


def _tracked_current(path: Path) -> None:
    source = path.with_suffix(".source.docx")
    document = Document()
    document.add_paragraph("after")
    document.save(source)
    with zipfile.ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    root = __import__("lxml.etree", fromlist=["etree"]).fromstring(parts["word/document.xml"])
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    paragraph = root.find(".//w:body/w:p", ns)
    assert paragraph is not None
    run = paragraph.find("./w:r", ns)
    assert run is not None
    paragraph.remove(run)
    deleted = OxmlElement("w:del")
    deleted.set(qn("w:id"), "1")
    deleted_run = OxmlElement("w:r")
    deleted_text = OxmlElement("w:delText")
    deleted_text.text = "before"
    deleted_run.append(deleted_text)
    deleted.append(deleted_run)
    inserted = OxmlElement("w:ins")
    inserted.set(qn("w:id"), "2")
    inserted_run = OxmlElement("w:r")
    inserted_text = OxmlElement("w:t")
    inserted_text.text = "after"
    inserted_run.append(inserted_text)
    inserted.append(inserted_run)
    paragraph.append(deleted)
    paragraph.append(inserted)
    parts["word/document.xml"] = __import__("lxml.etree", fromlist=["etree"]).tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, payload in parts.items():
            archive.writestr(name, payload)
    source.unlink()


def test_redline_accepts_current_revisions_before_diff(tmp_path: Path) -> None:
    base = tmp_path / "base.docx"
    base_document = Document()
    base_document.add_paragraph("before")
    base_document.save(base)
    current = tmp_path / "current.docx"
    _tracked_current(current)
    output = tmp_path / "tracked.docx"
    summary = create_tracked_docx(base, current, output, overwrite=True)
    assert summary["changed"] == 1
    with zipfile.ZipFile(output) as archive:
        from lxml import etree

        root = etree.fromstring(archive.read("word/document.xml"))
    assert visible_text(root, "final").strip() == "after"
    assert visible_text(root, "original").strip() == "before"
    assert not root.xpath(".//w:rPrChange", namespaces={"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"})
