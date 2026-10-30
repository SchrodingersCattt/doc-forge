"""Regression coverage for generic template formatting and Markdown layout."""
from __future__ import annotations

import json
import os
from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from lxml import etree
from PIL import Image

from docforge.cli import main
from docforge.markdown.launcher import convert_unicode_scripts_in_docx
from docforge.markdown.omml import parse_math_omml
from docforge.markdown.template import assemble_markdown_template
from docforge.math.pandoc import latex_to_omml
from docforge.style.engine import restyle_docx

NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
      "m": "http://schemas.openxmlformats.org/officeDocument/2006/math"}


@pytest.fixture
def fixture(tmp_path):
    template = tmp_path / "template.docx"
    doc = Document()
    for name in ("Title1", "Authors", "Address", "Abstract", "P1", "H1", "H2",
                 "FigureCaption", "References", "TableCaption", "TableHead", "TableBody",
                 "TableFoot", "Keywords", "BackMatterHeading", "BackMatter", "TableOfContentText", "Algorithm"):
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Georgia"
        style.font.size = Pt(8.5)
        style.font.color.rgb = RGBColor.from_string("123456")
        doc.add_paragraph("Example " + name, style=name)
    doc.styles["References"].paragraph_format.left_indent = Pt(14)
    doc.styles["References"].paragraph_format.first_line_indent = Pt(-14)
    # Deliberately conflicting sample formatting must not override References.
    sample = next(p for p in doc.paragraphs if p.style.name == "References")
    sample.paragraph_format.left_indent = Pt(0)
    sample.paragraph_format.space_after = Pt(6)
    sample.runs[0].font.size = Pt(22)
    doc.save(template)
    metadata = tmp_path / "metadata.md"
    metadata.write_text("# TITLE\nA title\n# AUTHORS\nA. Writer\n# ACKNOWLEDGMENTS\nThanks to helpers.\n", encoding="utf-8")
    source = tmp_path / "source.md"
    output = tmp_path / "output.docx"
    def build(text, **options):
        source.write_text(text, encoding="utf-8")
        result = assemble_markdown_template([source], template_path=template, metadata_path=metadata,
                                            output=output, force=True, **options)
        return Document(output), result
    return tmp_path, template, metadata, source, output, build


def test_template_inherits_typography_and_reference_layout(fixture):
    root, template, metadata, source, output, build = fixture
    bibliography = root / "refs.json"
    bibliography.write_text(json.dumps({"entry": "Author. Work. 2026."}), encoding="utf-8")
    doc, result = build("# Abstract\nAbstract words.\nKeywords: one · two\n\n# Introduction\nBody **bold** and *italic* \\citep{entry}.\n", bibliography_path=bibliography)
    for paragraph in doc.paragraphs:
        assert not paragraph._p.xpath(".//w:rPr/w:sz | .//w:rPr/w:szCs | .//w:rPr/w:color | .//w:rPr/w:rFonts")
    reference = next(p for p in doc.paragraphs if p.style.name == "References")
    assert not reference._p.xpath("./w:pPr/w:ind | ./w:pPr/w:spacing")
    assert reference.style.paragraph_format.first_line_indent.pt == -14
    abstract = next(p for p in doc.paragraphs if p.style.name == "Abstract")
    assert "Keywords:" not in abstract.text
    assert next(p for p in doc.paragraphs if p.style.name == "Keywords").text == "Keywords: one · two"
    assert next(p for p in doc.paragraphs if p.text == "ACKNOWLEDGMENTS").style.name == "BackMatterHeading"
    assert next(p for p in doc.paragraphs if p.text == "Thanks to helpers.").style.name == "BackMatter"
    assert any(r.bold for p in doc.paragraphs for r in p.runs if r.text == "bold")


def test_table_roles_safe_merges_and_keep_next(fixture):
    _, _, _, _, _, build = fixture
    doc, result = build("# Results\n\n## Values\n\nTable: Measurements.\n\n| Group | Sample | Y |\n|---|---|---|\n| a | s1 | -22.8 |\n| a | s2 | -22.8 |\n\n[a] A table note.\n\nTable note: Another note.\n")
    caption = next(p for p in doc.paragraphs if p.text.startswith("Table 1."))
    assert caption.style.name == "TableCaption"
    assert caption.paragraph_format.keep_with_next is True
    for heading in ("Results", "Values"):
        assert next(p for p in doc.paragraphs if p.text == heading).paragraph_format.keep_with_next is True
    for index, row in enumerate(doc.tables[0].rows):
        for cell in row.cells:
            # Continuation grouping cells resolve to their origin in python-docx.
            assert cell.paragraphs[0].style.name == ("TableHead" if index == 0 else "TableBody")
    xml_rows = doc.tables[0]._tbl.findall(qn("w:tr"))
    assert xml_rows[1].findall(qn("w:tc"))[0].find(qn("w:tcPr") + "/" + qn("w:vMerge")) is not None
    for row in xml_rows[1:]:
        data = row.findall(qn("w:tc"))[2]
        assert not data.findall(".//" + qn("w:vMerge"))
        assert "".join(data.itertext()).count("-22.8") >= 1
    assert [p.text for p in doc.paragraphs if p.style.name == "TableFoot"] == ["[a] A table note.", "Another note."]


@pytest.mark.parametrize("placement", ["source", "end"])
def test_end_matter_placement(fixture, placement):
    _, _, _, _, _, build = fixture
    doc, result = build("# Graphical Abstract\nA summary.\n\n# Introduction\nBody.\n\n[1] Reference.\n", end_matter_placement=placement)
    texts = [p.text for p in doc.paragraphs]
    index = texts.index("Graphical Abstract")
    assert (index > texts.index("[1] Reference.")) == (placement == "end")
    assert next(p for p in doc.paragraphs if p.text == "A summary.").style.name == "TableOfContentText"
    if placement == "end":
        assert doc.paragraphs[index - 1]._p.xpath('.//w:br[@w:type="page"]')
    assert result.end_matter_placement == placement


@pytest.mark.parametrize("policy", ["clear_all", "clear_conflicting", "preserve"])
def test_restyle_template_inheritance_and_semantic_emphasis(tmp_path, policy):
    template = tmp_path / "template.docx"
    doc = Document()
    base = doc.styles.add_style("BaseBody", WD_STYLE_TYPE.PARAGRAPH)
    base.font.size = Pt(8.5)
    base.font.name = "Georgia"
    base.font.color.rgb = RGBColor.from_string("123456")
    base.font.bold = False
    base.paragraph_format.left_indent = Pt(15)
    base.paragraph_format.space_after = Pt(3)
    target = doc.styles.add_style("CustomBody", WD_STYLE_TYPE.PARAGRAPH)
    target.base_style = base
    doc.save(template)
    source = tmp_path / "input.docx"
    paragraph = doc.add_paragraph(style="CustomBody")
    paragraph.paragraph_format.left_indent = Pt(0)
    paragraph.paragraph_format.space_after = Pt(30)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run("emphasis")
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor.from_string("000000")
    run.font.name = "Arial"
    run.bold = run.italic = run.font.superscript = True
    link = etree.SubElement(paragraph._p, qn("w:hyperlink"))
    nested = etree.SubElement(link, qn("w:r"))
    props = etree.SubElement(nested, qn("w:rPr"))
    etree.SubElement(props, qn("w:sz")).set(qn("w:val"), "40")
    etree.SubElement(nested, qn("w:t")).text = "nested"
    doc.save(source)
    output = tmp_path / "output.docx"
    profile = {"schema": "docforge.style.v1", "template": str(template),
               "roles": {"body": {"word_style": "CustomBody", "select": "CustomBody"}},
               "policy": {"direct_formatting": policy}}
    restyle_docx(source, output, profile)
    paragraph = Document(output).paragraphs[0]
    assert paragraph.runs[0].bold and paragraph.runs[0].italic and paragraph.runs[0].font.superscript
    assert paragraph.paragraph_format.keep_with_next
    if policy == "preserve":
        assert paragraph.runs[0].font.size.pt == 20
        assert paragraph.paragraph_format.space_after.pt == 30
    else:
        assert not paragraph._p.xpath(".//w:rPr/w:sz | .//w:rPr/w:color | .//w:rPr/w:rFonts")
        assert not paragraph._p.xpath("./w:pPr/w:ind | ./w:pPr/w:spacing")


def test_cli_v1_profile_and_audit_sidecars(fixture):
    root, template, metadata, source, output, _ = fixture
    source.write_text("Body.\n", encoding="utf-8")
    profile = root / "profile.json"
    profile.write_text(json.dumps({"schema": "docforge.style.v1", "template": template.name,
                                  "roles": {"body": {"word_style": "P1", "font": {"size_pt": 12}}}}), encoding="utf-8")
    assert main(["md2docx", str(source), "-o", str(output), "--template", str(template),
                 "--metadata", str(metadata), "--style", str(profile)]) == 0
    assert output.with_suffix(".sha256").is_file()
    assert output.with_suffix(".manifest.json").is_file()
    doc = Document(output)
    body = next(p for p in doc.paragraphs if p.text == "Body.")
    assert body.style.font.size.pt == 12


@pytest.mark.parametrize("slot", [False, True])
def test_tall_figures_fit_text_area_with_caption(fixture, slot):
    root, template, _, _, _, build = fixture
    image = root / "portrait.png"
    Image.new("RGB", (200, 1200), "white").save(image)
    if slot:
        doc = Document(template)
        doc.add_paragraph().add_run().add_picture(str(image), width=1000000)
        doc.add_paragraph("Figure 1. Example", style="FigureCaption")
        doc.save(template)
    doc, _ = build("# Results\n\n![A caption with several words.](portrait.png)\n")
    picture = next(p for p in doc.paragraphs if p._p.xpath(".//w:drawing"))
    extent = picture._p.xpath(".//wp:extent")[0]
    width, height = int(extent.get("cx")), int(extent.get("cy"))
    section = doc.sections[-1]
    text_height = section.page_height - section.top_margin - section.bottom_margin
    assert height < text_height - Pt(24)
    assert height / width == pytest.approx(6, rel=1e-4)
    assert picture.paragraph_format.keep_with_next is True


def test_native_inline_math_and_indented_algorithm(fixture):
    _, _, _, _, _, build = fixture
    doc, result = build(r"""# Methods

Body $m^\star \gets \arg\min_{m \in \text{cand}} \min_{r \in R} d(m,r)$.

```pseudocode
**Require:** $\mathbf{u} - \mathbf{1}_{\text{x.y}}$
    while $i < n$
        return $\frac{1}{|A|}$
```
""", inline_math="omml")
    body = next(p for p in doc.paragraphs if p.text.startswith("Body "))
    assert body._p.xpath("./m:oMath")
    assert body._p.xpath(".//m:sSup")
    algorithms = [p for p in doc.paragraphs if p.style.name == "Algorithm"]
    assert len(algorithms) == 3
    assert [p.paragraph_format.left_indent.pt for p in algorithms] == [0, 17, 34]
    assert algorithms[0]._p.xpath(".//m:sSub")
    assert algorithms[-1]._p.xpath('.//m:d/m:dPr/m:begChr[@m:val="|"]')
    assert result.inline_math == "omml"


def test_bars_in_nested_math_and_unpaired_operator():
    native = latex_to_omml(r"\frac{1}{|A|} + \frac{|B|}{|C|}")
    assert len(native.xpath('.//m:d/m:dPr/m:begChr[@m:val="|"]', namespaces=NS)) == 3
    assert not native.xpath('.//m:dPr/m:sepChr[@m:val=""]', namespaces=NS)
    small = etree.Element("root")
    small.extend(parse_math_omml(r"\frac{1}{|A|} + p(x|y)"))
    assert len(small.xpath(".//m:d", namespaces=NS)) == 1
    assert small.xpath('.//m:r[m:t="|"]', namespaces=NS)


def test_native_toc_has_cached_headings_and_update_request(fixture):
    _, _, _, _, output, build = fixture
    doc, _ = build("# Methods\nBody.\n\n## Details\nMore.\n", native_toc=True)
    assert doc.settings.element.find(qn("w:updateFields")).get(qn("w:val")) == "true"
    paragraph = next(p for p in doc.paragraphs if p._p.xpath(".//w:instrText"))
    assert "Methods" in paragraph.text and "Details" in paragraph.text
    assert "Right-click" not in paragraph.text
    assert paragraph._p.xpath('.//w:fldChar[@w:fldCharType="end"]')


def test_temp_handles_close_before_package_writes_and_replace(fixture, monkeypatch):
    _, _, _, _, output, build = fixture
    build("Unicode x₂.\n")
    import docforge.style.engine as engine
    import zipfile
    real_mkstemp, real_zip, real_close = engine.tempfile.mkstemp, zipfile.ZipFile, os.close
    descriptors = set()
    def mkstemp(*args, **kwargs):
        fd, name = real_mkstemp(*args, **kwargs)
        descriptors.add(fd)
        return fd, name
    def close(fd):
        real_close(fd)
        descriptors.discard(fd)
    class CheckedZip(real_zip):
        def __init__(self, file, mode="r", *args, **kwargs):
            if mode == "w":
                assert not descriptors
            super().__init__(file, mode, *args, **kwargs)
    monkeypatch.setattr(engine.tempfile, "mkstemp", mkstemp)
    monkeypatch.setattr(engine.os, "close", close)
    monkeypatch.setattr(zipfile, "ZipFile", CheckedZip)
    target = output.with_name("restyled.docx")
    restyle_docx(output, target, {"roles": {"body": {"word_style": "P1"}}})
    # Existing script conversion closes its named temporary file before replace.
    convert_unicode_scripts_in_docx(target)
    assert target.is_file()


@pytest.mark.parametrize("label", ["Table of Contents", "Graphical Abstract"])
def test_only_end_section_with_page_figure_preserves_audit_and_geometry(fixture, label):
    root, template, _, _, _, build = fixture
    image = root / "figure.png"
    Image.new("RGB", (200, 1200), "white").save(image)
    doc = Document(template)
    columns = doc.sections[-1]._sectPr.find(qn("w:cols"))
    columns.set(qn("w:num"), "2")
    doc.save(template)
    doc, result = build(f"# {label}\nSummary.\n\n![figure.png|span=page](figure.png)\n\nFigure 1. End figure.\n",
                        end_matter_placement="end", native_toc=True)
    assert len(result.figures) == 1
    assert result.figures[0]["span"] == "page"
    assert result.figures[0]["caption"] == "Figure 1. End figure."
    assert result.figures[0]["sha256"]
    assert result.template_sections_after == 3
    assert len(doc.inline_shapes) == 1
    headings = [p.text for p in doc.paragraphs]
    assert headings.index("CONTENTS") > headings.index("ACKNOWLEDGMENTS")


def test_algorithm_tabs_nbsp_blank_lines_and_default_font(fixture):
    _, template, _, _, _, build = fixture
    doc = Document(template)
    element = doc.styles["Algorithm"].element
    element.getparent().remove(element)
    doc.save(template)
    doc, _ = build("# Methods\n\n```algorithm\nstart\n\tstep\n\n\u00a0\u00a0\u00a0\u00a0return\n```\n")
    lines = [p for p in doc.paragraphs if p.text in {"start", "step", "return", ""} and p.paragraph_format.left_indent is not None]
    assert [p.text for p in lines] == ["start", "step", "", "return"]
    assert [p.paragraph_format.left_indent.pt for p in lines] == [0, 17, 0, 17]
    assert all(run.font.name == "Consolas" for p in lines for run in p.runs if run.text)


def test_clear_conflicting_preserves_unset_font_and_spacing_slots(tmp_path):
    source = tmp_path / "input.docx"
    doc = Document()
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    for child in list(defaults):
        defaults.remove(child)
    style = doc.styles["Normal"]
    style.font.name = "Georgia"
    style.paragraph_format.space_before = Pt(4)
    p = doc.add_paragraph("Body")
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(30)
    run = p.runs[0]
    run.font.name = "Arial"
    run._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Example East Asian")
    run.font.size = Pt(20)
    doc.save(source)
    output = tmp_path / "output.docx"
    restyle_docx(source, output, {"roles": {"body": {"word_style": "Normal"}}})
    p = Document(output).paragraphs[0]
    assert p.paragraph_format.space_before is None
    assert p.paragraph_format.space_after.pt == 30
    assert p.runs[0].font.name is None
    assert p.runs[0]._r.rPr.rFonts.get(qn("w:eastAsia")) == "Example East Asian"
    assert p.runs[0].font.size.pt == 20


def test_absolute_value_delimiters_render_in_libreoffice(fixture):
    from docforge.docxmd.verify import _soffice
    from docforge.gates.render import layout_gate
    if _soffice() is None:
        pytest.skip("LibreOffice unavailable; this skip is not a render pass")
    fitz = pytest.importorskip("fitz")
    root, _, _, _, output, build = fixture
    build(r"Body $\frac{1}{|A|}$.", inline_math="omml")
    report = layout_gate(output, root / "render", "libreoffice")
    assert report["status"] == "passed"
    with fitz.open(root / "render" / "output.pdf") as pdf:
        text = "".join(page.get_text() for page in pdf)
    assert "|A|" in "".join(text.split())
