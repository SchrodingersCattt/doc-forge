from __future__ import annotations

import json
import hashlib
import shutil
from pathlib import Path

import pytest
from docx import Document

from docforge.roundtrip import roundtrip_docx


PANDOC_AVAILABLE = shutil.which("pandoc") is not None


@pytest.mark.skipif(not PANDOC_AVAILABLE, reason="Pandoc is required for roundtrip export")
def test_roundtrip_reuses_exact_source_then_rebuilds_changed_markdown(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    document = Document()
    document.add_heading("Title", level=1)
    document.add_paragraph("Body")
    document.save(source)

    bundle = tmp_path / "bundle"
    exact = roundtrip_docx(source, workdir=bundle, output=tmp_path / "exact.docx", force=True)
    assert exact.mode == "exact-reuse"
    assert exact.output_sha256 == exact.source_sha256
    assert (tmp_path / "exact.docx").read_bytes() == source.read_bytes()

    manifest = json.loads((bundle / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["schema"] == "docforge.roundtrip.v1"
    assert manifest["source_copy"] == "roundtrip/source.docx"
    assert manifest["full"] == "full.md"

    full = bundle / manifest["full"]
    full.write_text(full.read_text(encoding="utf-8") + "\nEdited.\n", encoding="utf-8")
    rebuilt = roundtrip_docx(source, workdir=bundle, output=tmp_path / "rebuilt.docx", force=True)
    assert rebuilt.mode == "rebuild"
    assert (tmp_path / "rebuilt.docx").read_bytes() != source.read_bytes()


def test_roundtrip_rejects_snapshot_path_outside_bundle(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    Document().save(source)
    bundle = tmp_path / "bundle"
    bundle.mkdir()
    outside = tmp_path / "outside.docx"
    outside.write_bytes(source.read_bytes())
    payload = {
        "schema": "docforge.roundtrip.v1",
        "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "source_copy": "../outside.docx",
        "source_copy_sha256": "0" * 64,
        "full": "full.md",
        "full_sha256": "0" * 64,
        "section_sha256": {},
        "media_sha256": {},
    }
    (bundle / "manifest.json").write_text(json.dumps(payload), encoding="utf-8")
    with pytest.raises(ValueError, match="escapes"):
        roundtrip_docx(source, workdir=bundle, output=tmp_path / "output.docx", force=True)
