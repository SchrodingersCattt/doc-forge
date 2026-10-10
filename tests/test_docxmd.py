from __future__ import annotations

import io
import json
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.shared import Pt, RGBColor
from lxml import etree

from docforge.docxdiff.redline import create_tracked_docx
from docforge.docxmd import ExportOptions, build_docx, compare_docx, export_docx, roundtrip
from docforge.docxmd.inline import parse, write
from docforge.docxmd.ooxml import W, build, diff, flatten, merge, parse_tokens, format_tokens


INLINE_SAMPLES = [
    "plain text with spaces",
    "**bold** and *italic* and ***both***",
    "H<sub>2</sub>O and x<sup>2</sup> and <u>under</u> and ~~gone~~ and ==marked==",
    "[blue words]{color=1F4E79 sz=20} next",
    "{++inserted++}{author=\"A. B.\" date=2026-01-01T00:00:00Z} and {--deleted--}{author=gmy date=2026-01-01T00:00:00Z}",
    "tab<tab/>break<br/>end",
    "cite here.\\citep{smith2020,lee2021}",
    "escaped \\* star and \\{brace\\} and a \\| pipe",
    "[link text](https://example.org/a)",
]


@pytest.mark.parametrize("text", INLINE_SAMPLES)
def test_inline_round_trip(text):
    items = parse(text)
    assert parse(write(items)) == items


def test_tokens_round_trip_and_diff():
    xml = (
        f'<w:pPr xmlns:w="{W}"><w:pStyle w:val="Body"/><w:spacing w:after="240" w:line="360"/>'
        '<w:ind w:firstLine="480"/><w:jc w:val="both"/><w:rPr><w:b/><w:lang w:eastAsia="zh-CN"/></w:rPr></w:pPr>'
    )
    element = etree.fromstring(xml)
    tokens = flatten(element)
    assert flatten(build(element.tag, tokens)) == tokens
    assert parse_tokens(format_tokens(tokens)) == tokens
    target = [token for token in tokens if token[0] != "jc@val"] + [("keepNext", None)]
    assert dict(merge(tokens, diff(tokens, target))) == dict(target)


def _png() -> bytes:
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (40, 30), (200, 30, 30)).save(buffer, format="PNG")
    return buffer.getvalue()


def _sample_docx(path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.left_margin = Pt(60)
    title = doc.add_paragraph(style="Title")
    title.add_run("A Reversible Manuscript")
    doc.add_paragraph("Ada Lovelace, Alan Turing")
    heading = doc.add_heading("Introduction", level=1)
    for bookmark_id, name in ((1, "_RefIntro"), (2, "OLE_LINK1")):
        start = etree.SubElement(heading._p, f"{{{W}}}bookmarkStart")
        start.set(f"{{{W}}}id", str(bookmark_id))
        start.set(f"{{{W}}}name", name)
        heading._p.insert(1, start)
        end = etree.SubElement(heading._p, f"{{{W}}}bookmarkEnd")
        end.set(f"{{{W}}}id", str(bookmark_id))
    body = doc.add_paragraph()
    body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body.paragraph_format.first_line_indent = Pt(24)
    body.add_run("Water is H")
    body.add_run("2").font.subscript = True
    body.add_run("O, energy scales as x")
    body.add_run("2").font.superscript = True
    body.add_run(". Some ")
    body.add_run("bold").bold = True
    body.add_run(", ")
    body.add_run("italic").italic = True
    body.add_run(", ")
    body.add_run("underlined").underline = True
    body.add_run(", ")
    body.add_run("struck").font.strike = True
    body.add_run(", ")
    marked = body.add_run("highlighted")
    marked.font.highlight_color = WD_COLOR_INDEX.YELLOW
    body.add_run(" and ")
    red = body.add_run("red")
    red.font.color.rgb = RGBColor(0xC0, 0, 0)
    body.add_run(" text, as reported.")
    body.add_run("1,2").font.superscript = True
    center = doc.add_paragraph("A centered line in another font.")
    center.alignment = WD_ALIGN_PARAGRAPH.CENTER
    center.runs[0].font.name = "Arial"
    doc.add_heading("Results", level=2)
    doc.add_picture(io.BytesIO(_png()))
    doc.add_paragraph("Figure 1. A red square.")
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    for row, values in enumerate((("Property", "Value"), ("Density", "1.9"))):
        for col, value in enumerate(values):
            table.cell(row, col).text = value
    doc.add_heading("References", level=1)
    doc.add_paragraph("[1]\tSmith, J. A study. J. Chem. 2020, 1, 1.")
    doc.add_paragraph("[2]\tLee, K. Another study. J. Chem. 2021, 2, 2.")
    doc.save(path)


def test_docx_round_trip_and_minimal_redline(tmp_path):
    source = tmp_path / "source.docx"
    _sample_docx(source)
    split = [{"file": "00.front.md", "start": None}, {"file": "01.body.md", "start": "^Introduction$"}]
    report = roundtrip(source, tmp_path / "md", "main", split=split, bibliography="refs.json")
    assert report["issues"] == []
    assert report["ok"], report

    bundle = tmp_path / "md"
    body = (bundle / "01.body.md").read_text(encoding="utf-8")
    assert "H<sub>2</sub>O" in body
    assert "**bold**" in body and "*italic*" in body and "<u>underlined</u>" in body
    assert "~~struck~~" in body and "==highlighted==" in body
    assert "\\citep{smith2020,lee2021}" in body
    assert "![](media/main/" in body
    assert "::: table" in body and "::: references" in body
    assert "<bookmark-start name=\"_RefIntro\"/>" in body and "OLE_LINK" not in body
    refs = json.loads((bundle / "refs.json").read_text(encoding="utf-8"))
    assert list(key for key in refs if not key.startswith("_")) == ["smith2020", "lee2021"]
    manifest = json.loads((bundle / "bundle.json").read_text(encoding="utf-8"))
    assert manifest["documents"][0]["metadata"]["title"] == "A Reversible Manuscript"

    (bundle / "01.body.md").write_text(body.replace("Some **bold**", "Some **strong**"), encoding="utf-8")
    edited = tmp_path / "edited.docx"
    build_docx(bundle, "main", edited)
    issues = compare_docx(source, edited)
    assert len(issues) == 1 and "bold" in issues[0]
    summary = create_tracked_docx(source, edited, tmp_path / "tracked.docx", author="tester", workers=1)
    assert summary["changed"] == 1 and summary["inserted"] == summary["deleted"] == 0
    with zipfile.ZipFile(tmp_path / "tracked.docx") as package:
        document = package.read("word/document.xml").decode("utf-8")
    assert document.count("<w:ins ") == 1 and document.count("<w:del ") == 1
    assert "strong" in document


def test_export_is_stable(tmp_path):
    source = tmp_path / "source.docx"
    _sample_docx(source)
    bundle = tmp_path / "md"
    export_docx(source, bundle, ExportOptions("main"))
    first = (bundle / "main.md").read_text(encoding="utf-8")
    rebuilt = tmp_path / "rebuilt.docx"
    build_docx(bundle, "main", rebuilt)
    again = tmp_path / "again"
    export_docx(rebuilt, again, ExportOptions("main"))
    assert (again / "main.md").read_text(encoding="utf-8") == first


def test_comment_ids_follow_markdown(tmp_path):
    source = tmp_path / "source.docx"
    _sample_docx(source)
    bundle = tmp_path / "md"
    export_docx(source, bundle, ExportOptions("main"))
    text = (bundle / "main.md").read_text(encoding="utf-8")
    text = text.replace("Some **bold**", 'Some <comment-start id="c7"/>**bold**<comment-end id="c7"/> '
                        '<comment-start id="note"/>text<comment-end id="note"/>', 1)
    text += ('\n::: comment {: id=c7 author=a date=2026-01-01T00:00:00Z}\nKept\n:::\n'
             '\n::: comment {: id=note author=b date=2026-01-01T00:00:00Z}\nNew\n:::\n')
    (bundle / "main.md").write_text(text, encoding="utf-8")
    built = tmp_path / "built.docx"
    build_docx(bundle, "main", built)
    with zipfile.ZipFile(built) as package:
        comments = etree.fromstring(package.read("word/comments.xml"))
        document = etree.fromstring(package.read("word/document.xml"))
    ids = [node.get(f"{{{W}}}id") for node in comments.iter(f"{{{W}}}comment")]
    assert ids == ["7", "8"]
    refs = {node.get(f"{{{W}}}id") for node in document.iter(f"{{{W}}}commentReference")}
    assert refs == {"7", "8"}


def test_redline_follows_moved_comment(tmp_path):
    source = tmp_path / "source.docx"
    _sample_docx(source)
    bundle = tmp_path / "md"
    export_docx(source, bundle, ExportOptions("main"))
    original = (bundle / "main.md").read_text(encoding="utf-8")
    note = '\n::: comment {: id=c3 author=a date=2026-01-01T00:00:00Z}\nCheck\n:::\n'
    anchored = '<comment-start id="c3"/>**bold**<comment-end id="c3"/>'
    (bundle / "main.md").write_text(original.replace("**bold**", anchored, 1) + note, encoding="utf-8")
    base = tmp_path / "base.docx"
    build_docx(bundle, "main", base)
    moved = original.replace("Some **bold**", "Some **strong**", 1)
    moved = moved.replace("in another font.", 'in <comment-start id="c3"/>a new font<comment-end id="c3"/>.', 1)
    (bundle / "main.md").write_text(moved + note, encoding="utf-8")
    current = tmp_path / "current.docx"
    build_docx(bundle, "main", current)
    create_tracked_docx(base, current, tmp_path / "tracked.docx", author="tester", workers=1)
    with zipfile.ZipFile(tmp_path / "tracked.docx") as package:
        document = etree.fromstring(package.read("word/document.xml"))
    for name in ("commentRangeStart", "commentRangeEnd", "commentReference"):
        (node,) = document.iter(f"{{{W}}}{name}")
        paragraph = next(node.iterancestors(f"{{{W}}}p"))
        assert "a new font" in "".join(paragraph.itertext())
