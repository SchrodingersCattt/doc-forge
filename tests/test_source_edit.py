from __future__ import annotations

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from docforge.docxdiff import create_tracked_docx
from docforge.docxdiff.redline import visible_text
from docforge.markdown.source_edit import apply_markdown_delta
from docforge.markdown.source_edit import _copy_paragraph_properties, _list_headings, _run_properties
from docforge.markdown.source_edit import _Block, _map_blocks


def test_markdown_delta_keeps_unedited_paragraphs(tmp_path):
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("Electronic Supplementary Information")
    document.add_paragraph("Keep this sentence exactly, including its place in the document.")
    document.add_paragraph(
        "EAP-4 and EAP-8 show nearly identical onset temperatures but very different decomposition behaviors."
    )
    document.add_paragraph("Sensitivity Tests")
    document.add_paragraph("Friction sensitivity measurements were performed in accordance with STANAG 4487.")
    document.save(source)

    baseline = tmp_path / "baseline.md"
    edited = tmp_path / "edited.md"
    baseline.write_text(
        "\n\n".join(
            [
                "Electronic Supplementary Information",
                "Keep this sentence exactly, including its place in the document.",
                "EAP-4 and EAP-8 show nearly identical onset temperatures but very different decomposition behaviors.",
                "4.  **Sensitivity Tests**",
                "Friction sensitivity measurements were performed in accordance with STANAG 4487.",
            ]
        ),
        encoding="utf-8",
    )
    edited.write_text(
        "\n\n".join(
            [
                "Electronic Supplementary Information",
                "Keep this sentence exactly, including its place in the document.",
                "EAP-4 and **EAP-8** show nearly identical onset temperatures but markedly different decomposition kinetics.",
                "The 1666 K trajectories follow nitrogen provenance.",
                "## Reactive molecular dynamics",
                "The trajectories used a four-member committee.",
                "4.  **Sensitivity Tests**",
                "Friction sensitivity measurements were performed in accordance with STANAG 4487.",
            ]
        ),
        encoding="utf-8",
    )
    current = tmp_path / "current.docx"
    summary = apply_markdown_delta(source, baseline, edited, current, overwrite=True)
    assert summary["replaced"] == 1
    assert summary["inserted"] == 0

    current_doc = Document(current)
    texts = [paragraph.text for paragraph in current_doc.paragraphs]
    assert texts[0] == "Electronic Supplementary Information"
    assert texts[1] == "Keep this sentence exactly, including its place in the document."
    assert "markedly different" in texts[2]
    assert "1666 K" in texts[3]
    assert texts[4] == "Reactive molecular dynamics"
    assert "four-member" in texts[5]
    assert texts[6] == "Sensitivity Tests"

    tracked = tmp_path / "tracked.docx"
    redline = create_tracked_docx(source, current, tracked, author="gmy", overwrite=True)
    assert redline["deleted"] == 0
    accepted = []
    from lxml import etree
    from zipfile import ZipFile

    root = etree.fromstring(ZipFile(tracked).read("word/document.xml"))
    for paragraph in root.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"):
        text = visible_text(paragraph).strip()
        if text:
            accepted.append(text)
    assert "Keep this sentence exactly, including its place in the document." in accepted
    assert "Electronic Supplementary Information" in accepted


def test_insert_keeps_anchor_size_and_unescapes_markdown(tmp_path):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    source = tmp_path / "source.docx"
    document = Document()
    title = document.add_paragraph("Electronic Supplementary Information")
    title.runs[0].font.size = Pt(18)
    title.runs[0].bold = True
    body = document.add_paragraph("Keep this sentence exactly, including its place in the document.")
    body.runs[0].font.size = Pt(12)
    heading = document.add_paragraph("Sensitivity Tests")
    heading.runs[0].bold = True
    heading.runs[0].font.size = Pt(14)
    p_pr = heading._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "1")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "2")
    num_pr.append(ilvl)
    num_pr.append(num_id)
    p_pr.append(num_pr)
    document.save(source)

    baseline = tmp_path / "baseline.md"
    edited = tmp_path / "edited.md"
    baseline.write_text(
        "\n\n".join(
            [
                "Electronic Supplementary Information",
                "Keep this sentence exactly, including its place in the document.",
                "4.  **Sensitivity Tests**",
            ]
        ),
        encoding="utf-8",
    )
    edited.write_text(
        "\n\n".join(
            [
                "Electronic Supplementary Information",
                "Keep this sentence exactly, including its place in the document.",
                "## Reactive molecular dynamics",
                r"\[11\] Y. Zhang, H. Wang, example citation for the delta.",
                "4.  **Sensitivity Tests**",
            ]
        ),
        encoding="utf-8",
    )
    current = tmp_path / "current.docx"
    apply_markdown_delta(source, baseline, edited, current, overwrite=True)
    current_doc = Document(current)
    texts = [paragraph.text for paragraph in current_doc.paragraphs]
    assert texts[2] == "Reactive molecular dynamics"
    assert texts[3].startswith("[11] Y. Zhang")
    assert "\\[" not in texts[3]
    heading_run = current_doc.paragraphs[2].runs[0]
    citation_run = current_doc.paragraphs[3].runs[0]
    assert heading_run.bold is True
    assert heading_run.font.size == Pt(14)
    assert citation_run.font.size == Pt(12)
    assert current_doc.paragraphs[2]._p.find(qn("w:pPr")).find(qn("w:numPr")) is None


def test_insertions_follow_the_next_mapped_paragraph(tmp_path):
    source = tmp_path / "source.docx"
    document = Document()
    original = [
        "P stays byte-identical.",
        "Q is the paragraph whose wording changes.",
        "R is the next mapped paragraph.",
        "S stays byte-identical too.",
    ]
    for text in original:
        document.add_paragraph(text)
    document.save(source)

    baseline = tmp_path / "baseline.md"
    baseline.write_text("\n\n".join(original), encoding="utf-8")
    edited = tmp_path / "edited.md"
    edited.write_text(
        "\n\n".join(
            [
                original[0],
                "Q has revised wording.",
                "Inserted before R (one).",
                "Inserted before R (two).",
                original[2],
                original[3],
            ]
        ),
        encoding="utf-8",
    )
    output = tmp_path / "output.docx"
    summary = apply_markdown_delta(source, baseline, edited, output, overwrite=True)
    assert summary == {"replaced": 1, "inserted": 0}
    assert [paragraph.text for paragraph in Document(output).paragraphs] == [
        original[0],
        "Q has revised wording.",
        "Inserted before R (one).",
        "Inserted before R (two).",
        original[2],
        original[3],
    ]


def test_mapping_pairs_duplicate_sentences_in_source_order():
    blocks = [_Block("text", "Repeated sentence."), _Block("text", "Repeated sentence.")]
    assert _map_blocks(blocks, [(3, "Repeated sentence."), (8, "Repeated sentence.")]) == {
        0: 3,
        1: 8,
    }


def test_mapping_solves_exact_and_fuzzy_duplicates_in_one_ordered_pass():
    # The exact second block must not greedily consume source[0] before the
    # earlier fuzzy block has had a chance to pair with it.
    blocks = [_Block("text", "A repeated sentence with one edit."), _Block("text", "A repeated sentence.")]
    mapping = _map_blocks(
        blocks,
        [
            (0, "A repeated sentence with one edit!"),
            (1, "A repeated sentence."),
        ],
    )
    assert mapping == {0: 0, 1: 1}


def test_mapping_keeps_short_fuzzy_duplicate_before_exact_duplicate():
    # A one-character typo in a short key must not let the later exact
    # duplicate consume the first source paragraph.
    assert _map_blocks(
        [_Block("text", "fooo"), _Block("text", "foo")],
        [(0, "foo"), (1, "foo")],
    ) == {0: 0, 1: 1}


def test_body_insertion_before_heading_uses_body_formatting(tmp_path):
    source = tmp_path / "source.docx"
    document = Document()
    heading = document.add_paragraph("Heading")
    heading.runs[0].bold = True
    heading.runs[0].font.size = Pt(18)
    body = document.add_paragraph("Body anchor")
    body.runs[0].font.size = Pt(11)
    document.save(source)

    baseline = tmp_path / "baseline.md"
    baseline.write_text("Heading\n\nBody anchor", encoding="utf-8")
    edited = tmp_path / "edited.md"
    edited.write_text("Inserted body\n\nHeading\n\nBody anchor", encoding="utf-8")
    output = tmp_path / "output.docx"
    apply_markdown_delta(source, baseline, edited, output, overwrite=True)

    inserted = Document(output).paragraphs[0]
    assert inserted.text == "Inserted body"
    assert inserted.runs[0].font.size == Pt(11)
    assert inserted.runs[0].bold is not True


def test_body_insertion_does_not_borrow_unstyled_title_metrics(tmp_path):
    from docx.shared import Pt

    source = tmp_path / "source.docx"
    document = Document()
    title = document.add_paragraph("A very long unstyled title that must not be used as body prototype")
    title.runs[0].font.size = Pt(18)
    body = document.add_paragraph("Body anchor with enough words for matching.")
    body.runs[0].font.size = Pt(11)
    heading = document.add_paragraph("Heading")
    heading.runs[0].bold = True
    document.save(source)
    baseline = tmp_path / "baseline.md"
    baseline.write_text(
        "A very long unstyled title that must not be used as body prototype\n\n"
        "Body anchor with enough words for matching.\n\nHeading",
        encoding="utf-8",
    )
    edited = tmp_path / "edited.md"
    edited.write_text(
        "A very long unstyled title that must not be used as body prototype\n\n"
        "Body anchor with enough words for matching.\n\nInserted body text.\n\nHeading",
        encoding="utf-8",
    )
    output = tmp_path / "output.docx"
    apply_markdown_delta(source, baseline, edited, output, overwrite=True)
    inserted = Document(output).paragraphs[2]
    assert inserted.text == "Inserted body text."
    assert inserted.runs[0].font.size == Pt(11)


def test_mapping_allows_short_headings_and_does_not_jump_to_distant_fuzzy_match():
    assert _map_blocks([_Block("text", "# R")], [(2, "R")]) == {0: 2}
    block = _Block("text", "A sentence with a local edit.")
    # The distant paragraph is a strong fuzzy match, but the nearest source
    # neighbour is intentionally unrelated. Ordered matching must leave the
    # block unmapped instead of reaching across that neighbour.
    mapping = _map_blocks(
        [block],
        [
            (0, "An unrelated paragraph."),
            (1, "A sentence with a local edit!"),
        ],
    )
    assert mapping == {}

    # A future exact anchor is also the next neighbour for the margin check;
    # a close fuzzy candidate must not steal that anchor.
    assert _map_blocks(
        [
            _Block("text", "A sentence with a local edit."),
            _Block("text", "A sentence with a local edit?"),
        ],
        [(0, "A sentence with a local edit!"), (1, "A sentence with a local edit?")],
    ) == {1: 1}


def test_insertion_before_short_heading_skips_unmapped_equation(tmp_path):
    from lxml import etree

    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("Opening paragraph.")
    equation = document.add_paragraph()
    math = etree.SubElement(equation._p, "{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath")
    etree.SubElement(math, "{http://schemas.openxmlformats.org/officeDocument/2006/math}t").text = "E=mc2"
    document.add_paragraph("R")
    document.add_paragraph("Closing paragraph.")
    document.save(source)

    baseline = tmp_path / "baseline.md"
    baseline.write_text("\n\n".join(["Opening paragraph.", r"\[E=mc^2\]", "# R", "Closing paragraph."]), encoding="utf-8")
    edited = tmp_path / "edited.md"
    edited.write_text(
        "\n\n".join(
            [
                "Opening paragraph.",
                r"\[E=mc^2\]",
                "Inserted after the equation.",
                "# R",
                "Closing paragraph.",
            ]
        ),
        encoding="utf-8",
    )
    output = tmp_path / "output.docx"
    apply_markdown_delta(source, baseline, edited, output, overwrite=True)
    assert [paragraph.text for paragraph in Document(output).paragraphs] == [
        "Opening paragraph.",
        "",
        "Inserted after the equation.",
        "R",
        "Closing paragraph.",
    ]
def test_insert_formatting_drops_list_numbering_and_placeholder_highlight(tmp_path):
    source = tmp_path / "source.docx"
    document = Document()
    anchor = document.add_paragraph("Anchor paragraph")
    ppr = anchor._p.get_or_add_pPr()
    style = OxmlElement("w:pStyle")
    style.set(qn("w:val"), "ListParagraph")
    ppr.append(style)
    indentation = OxmlElement("w:ind")
    indentation.set(qn("w:left"), "720")
    indentation.set(qn("w:hanging"), "360")
    ppr.append(indentation)
    numbering = OxmlElement("w:numPr")
    numbering.append(OxmlElement("w:ilvl"))
    ppr.append(numbering)
    run = anchor.runs[0]
    run.font.size = Pt(13)
    highlight = OxmlElement("w:highlight")
    highlight.set(qn("w:val"), "yellow")
    run._r.get_or_add_rPr().append(highlight)
    document.save(source)

    baseline = tmp_path / "baseline.md"
    edited = tmp_path / "edited.md"
    baseline.write_text("Anchor paragraph", encoding="utf-8")
    edited.write_text("Anchor paragraph\n\nInserted paragraph", encoding="utf-8")
    output = tmp_path / "output.docx"
    apply_markdown_delta(source, baseline, edited, output, overwrite=True)
    inserted = Document(output).paragraphs[1]
    properties = inserted._p.pPr
    assert properties.find(qn("w:numPr")) is None
    assert properties.find(qn("w:pStyle")) is None
    assert inserted.runs[0].font.size == Pt(13)
    assert inserted.runs[0]._r.find(".//" + qn("w:highlight")) is None


def test_heading_list_prototype_uses_first_clean_bold_sample_per_level():
    def numbered(text, level, *, highlight=False, tab=False, field=None):
        paragraph = OxmlElement("w:p")
        ppr = OxmlElement("w:pPr")
        num = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), str(level))
        num.append(ilvl)
        ppr.append(num)
        paragraph.append(ppr)
        run = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        rpr.append(OxmlElement("w:b"))
        if highlight:
            rpr.append(OxmlElement("w:highlight"))
        run.append(rpr)
        if tab:
            run.append(OxmlElement("w:tab"))
        if field:
            instruction = OxmlElement("w:instrText")
            instruction.text = field
            run.append(instruction)
        text_node = OxmlElement("w:t")
        text_node.text = text
        run.append(text_node)
        paragraph.append(run)
        return paragraph

    ignored = numbered("ignored", 1, highlight=True)
    ignored_tab = numbered("ignored", 1, tab=True)
    ignored_field = numbered("ignored", 1, field=" PAGEREF _Toc1 ")
    clean = numbered("clean", 1)
    subsection = numbered("subsection", 2)
    prototypes = _list_headings([ignored, ignored_tab, ignored_field, clean, subsection])
    assert prototypes[1] is clean
    assert prototypes[2] is subsection
    copied = _copy_paragraph_properties(clean)
    assert copied.find(qn("w:numPr")) is None
    assert _run_properties(clean).find(qn("w:b")) is not None


def test_body_insertion_does_not_borrow_unstyled_title_metrics(tmp_path):
    source = tmp_path / "source.docx"
    document = Document()
    title = document.add_paragraph("A very long unstyled title that must not be used as body prototype")
    title.runs[0].font.size = Pt(18)
    body = document.add_paragraph("Body anchor with enough words for matching.")
    body.runs[0].font.size = Pt(11)
    heading = document.add_paragraph("Heading")
    heading.runs[0].bold = True
    document.save(source)
    baseline = tmp_path / "baseline.md"
    baseline.write_text(
        "A very long unstyled title that must not be used as body prototype\n\n"
        "Body anchor with enough words for matching.\n\nHeading",
        encoding="utf-8",
    )
    edited = tmp_path / "edited.md"
    edited.write_text(
        "A very long unstyled title that must not be used as body prototype\n\n"
        "Body anchor with enough words for matching.\n\nInserted body text.\n\nHeading",
        encoding="utf-8",
    )
    output = tmp_path / "output.docx"
    apply_markdown_delta(source, baseline, edited, output, overwrite=True)
    inserted = Document(output).paragraphs[2]
    assert inserted.text == "Inserted body text."
    assert inserted.runs[0].font.size == Pt(11)
