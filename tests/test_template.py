from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
from docforge.cli import build_parser

from docforge.markdown import (
    _format_bibliography_record,
    assemble_markdown_template,
    discover_template_styles,
    parse_metadata,
    verify_template_output,
    write_assembly_sidecars,
)
from docforge.markdown.template import _word_compatible_image_bytes
from docforge.output import validate_output_path


def test_output_path_rejects_final_token(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="forbidden token 'final'"):
        validate_output_path(tmp_path / "main-final.docx")
    validate_output_path(tmp_path / "main-paph2.docx")


def _template(path: Path) -> None:
    document = Document()
    for name in (
        "BB_Author_Name",
        "FA_Corresponding_Author_Footnote",
        "BD_Abstract_Title",
        "BD_Abstract",
        "TA_Main_Text",
        "TF_References_Section",
        "EndNote Bibliography",
    ):
        document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("[TITLE] [placeholder text]").style = document.styles["BB_Author_Name"]
    document.add_paragraph("[ABSTRACT] [placeholder text]").style = document.styles["BD_Abstract"]
    document.add_section(WD_SECTION_START.CONTINUOUS)
    document.add_paragraph("[SECTION] [placeholder text]")
    document.save(path)


def _one_section_template(path: Path) -> None:
    document = Document()
    document.add_paragraph("[TITLE] [placeholder text]")
    document.save(path)


def _figure_template(path: Path, image: Path) -> None:
    document = Document()
    for name in ("BB_Author_Name", "BD_Abstract", "TA_Main_Text", "Caption", "EndNote Bibliography"):
        if name not in [style.name for style in document.styles]:
            document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("[TITLE]").style = document.styles["BB_Author_Name"]
    document.add_paragraph("[ABSTRACT]").style = document.styles["BD_Abstract"]
    document.add_section(WD_SECTION_START.CONTINUOUS)
    document.add_paragraph("Body").style = document.styles["TA_Main_Text"]
    document.add_section(WD_SECTION_START.CONTINUOUS)
    document.add_paragraph().add_run().add_picture(str(image), width=1000000)
    document.add_paragraph("[FIGURE CAPTION] [placeholder text]").style = document.styles["Caption"]
    document.save(path)


def test_metadata_and_style_discovery(tmp_path: Path) -> None:
    metadata = tmp_path / "metadata.md"
    metadata.write_text(
        "# TITLE\n\nA title\n# AUTHOR\n\nA. Author\n# AFFILIATION:\n\nA Lab\n# EMAILS\n\n*one@example.com\n**two@example.com\n",
        encoding="utf-8",
    )
    parsed = parse_metadata(metadata)
    assert parsed.title == "A title"
    assert parsed.authors == "A. Author"
    assert parsed.affiliations == "A Lab"
    assert parsed.contacts == ("*one@example.com", "**two@example.com")
    template = tmp_path / "template.docx"
    _template(template)
    styles = discover_template_styles(Document(template))
    assert styles["title"] == "BB_Author_Name"
    assert styles["body"] == "TA_Main_Text"
    assert styles["reference"] == "EndNoteBibliography"


def test_metadata_parses_author_contributions(tmp_path: Path) -> None:
    metadata = tmp_path / "metadata.md"
    metadata.write_text("# TITLE\n\nA title\n# AUTHOR CONTRIBUTIONS\n\nA. Author contributed.\n", encoding="utf-8")
    assert parse_metadata(metadata).author_contributions == "A. Author contributed."


def test_correspondence_stars_render_literally(tmp_path: Path) -> None:
    template = tmp_path / "template.docx"
    _template(template)
    metadata = tmp_path / "metadata.md"
    metadata.write_text(
        "# TITLE\n\nA title\n# AUTHOR\n\nA. One*, B. Two*\n# EMAILS\n\n*one@example.com\n*two@example.com\n",
        encoding="utf-8",
    )
    source = tmp_path / "source.md"
    source.write_text("Body.\n", encoding="utf-8")
    output = tmp_path / "output.docx"
    assemble_markdown_template([source], template_path=template, output=output, metadata_path=metadata)
    document = Document(output)
    text = "\n".join(p.text for p in document.paragraphs)
    assert "A. One*, B. Two*" in text
    assert "*one@example.com" in text and "*two@example.com" in text
    author = next(p for p in document.paragraphs if "A. One" in p.text)
    assert all(run.font.italic is not True for run in author.runs)


def test_structured_bibliography_record_matches_formatter_contract() -> None:
    formatted = _format_bibliography_record(
        {
            "authors": ["Author, One", "Author, Two"],
            "title": "A title",
            "year": 2025,
            "journal": "Journal",
            "volume": "12",
            "issue": "3",
            "locator": "101-110",
            "doi": "https://doi.org/10.1000/example",
        },
        "example",
    )
    assert formatted == "Author, One, and Author, Two. A title. *Journal*, **2025**, *12* (3), 101-110. DOI: 10.1000/example"

def test_template_assembly_replaces_placeholders_and_numbers_citations(tmp_path: Path) -> None:
    template = tmp_path / "template.docx"
    _template(template)
    metadata = tmp_path / "metadata.md"
    metadata.write_text("# TITLE\n\nA title\n", encoding="utf-8")
    abstract = tmp_path / "abstract.md"
    abstract.write_text("# Abstract\n\nAbstract cites \\citep{ref_b}.\n", encoding="utf-8")
    body = tmp_path / "body.md"
    body.write_text("# Introduction\n\nBody cites \\cite{ref_a}.\n", encoding="utf-8")
    bibliography = tmp_path / "refs.json"
    bibliography.write_text(
        json.dumps({"ref_a": "Author A. *Journal*, **2020**.", "ref_b": "Author B. *Journal*, **2021**."}),
        encoding="utf-8",
    )
    output = tmp_path / "output.docx"

    result = assemble_markdown_template(
        [abstract, body],
        template_path=template,
        output=output,
        metadata_path=metadata,
        bibliography_path=bibliography,
    )
    document = Document(output)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert result.used_citations == ("ref_b", "ref_a")
    assert "Abstract cites [1]." in text
    assert "Body cites [2]." in text
    assert "1.\tAuthor B." in text
    assert "2.\tAuthor A." in text
    assert "placeholder" not in text.lower()
    assert len(document.sections) == 2
    verify_template_output(
        output,
        expected_sections=2,
        expected_geometry=result.section_geometry,
    )
    manifest, checksum = write_assembly_sidecars(
        result,
        inputs=[abstract, body],
        template_path=template,
        metadata_path=metadata,
        bibliography_path=bibliography,
        command=["docforge", "md2docx"],
    )
    assert manifest.exists()
    assert checksum.exists()
    payload = json.loads(manifest.read_text(encoding="utf-8"))
    assert payload["rendering"]["figure_span"] == "column"


def test_unknown_citation_fails_before_writing(tmp_path: Path) -> None:
    template = tmp_path / "template.docx"
    _template(template)
    metadata = tmp_path / "metadata.md"
    metadata.write_text("# TITLE\n\nA title\n", encoding="utf-8")
    source = tmp_path / "source.md"
    source.write_text("# Intro\n\n\\citep{missing}\n", encoding="utf-8")
    bibliography = tmp_path / "refs.json"
    bibliography.write_text("{}", encoding="utf-8")
    output = tmp_path / "output.docx"
    with pytest.raises(ValueError, match="Unknown citation key"):
        assemble_markdown_template(
            [source],
            template_path=template,
            output=output,
            metadata_path=metadata,
            bibliography_path=bibliography,
        )
    assert not output.exists()


def test_one_section_template_uses_terminal_section_properties(tmp_path: Path) -> None:
    template = tmp_path / "template.docx"
    _one_section_template(template)
    metadata = tmp_path / "metadata.md"
    metadata.write_text("# TITLE\n\nA title\n", encoding="utf-8")
    source = tmp_path / "source.md"
    source.write_text("# Introduction\n\nBody.\n", encoding="utf-8")
    output = tmp_path / "output.docx"
    result = assemble_markdown_template(
        [source], template_path=template, output=output, metadata_path=metadata
    )
    assert result.template_sections_before == 1
    assert result.template_sections_after == 1
    assert "placeholder" not in "\n".join(p.text for p in Document(output).paragraphs).lower()


def _front_and_body_same_style_template(path: Path) -> None:
    document = Document()
    title = document.add_paragraph("Supporting Information")
    title.alignment = 1
    title.runs[0].italic = True
    document.add_paragraph("A long body prototype with ordinary left-aligned formatting.")
    document.add_paragraph("Section", style="Heading 1")
    document.add_paragraph("A second long body prototype with ordinary left-aligned formatting.")
    document.save(path)


def _mixed_figure_template(path: Path, image: Path) -> None:
    document = Document()
    document.add_paragraph("[TITLE]")
    document.add_paragraph("Body")
    for index in (1, 2):
        document.add_paragraph().add_run().add_picture(str(image), width=1000000)
        document.add_paragraph(f"Figure {index}. Placeholder caption.")
    document.save(path)


def test_body_prototype_skips_front_matter_when_style_is_shared(tmp_path: Path) -> None:
    template = tmp_path / "template.docx"
    _front_and_body_same_style_template(template)
    metadata = tmp_path / "metadata.md"
    metadata.write_text("# TITLE\n\nA title\n", encoding="utf-8")
    source = tmp_path / "source.md"
    source.write_text("# Intro\n\nGenerated body text.\n", encoding="utf-8")
    output = tmp_path / "output.docx"
    assemble_markdown_template(
        [source], template_path=template, output=output, metadata_path=metadata
    )
    body = next(p for p in Document(output).paragraphs if "Generated body" in p.text)
    assert body.alignment != 1
    assert all(run.font.italic is not True for run in body.runs)


def test_template_uses_all_mixed_region_figure_slots(tmp_path: Path) -> None:
    source_image = tmp_path / "source.png"
    Image.new("RGB", (80, 40), "white").save(source_image)
    template = tmp_path / "template.docx"
    _mixed_figure_template(template, source_image)
    metadata = tmp_path / "metadata.md"
    metadata.write_text("# TITLE\n\nA title\n", encoding="utf-8")
    source = tmp_path / "source.md"
    source.write_text(
        "Body.\n\n![Figure 1](source.png)\n\nFigure 1. One.\n\n"
        "![Figure 2](source.png)\n\nFigure 2. Two.\n",
        encoding="utf-8",
    )
    output = tmp_path / "output.docx"
    result = assemble_markdown_template(
        [source], template_path=template, output=output, metadata_path=metadata
    )
    assert len(result.figures) == 2
    assert all(figure["span"] == "column" for figure in result.figures)
    assert len(Document(output).inline_shapes) == 2


def test_template_assembly_keeps_inline_figure_and_caption_pair(tmp_path: Path) -> None:
    source_image = tmp_path / "source.png"
    Image.new("RGB", (80, 40), "white").save(source_image)
    template = tmp_path / "template.docx"
    _figure_template(template, source_image)
    metadata = tmp_path / "metadata.md"
    metadata.write_text("# TITLE\n\nA title\n", encoding="utf-8")
    source = tmp_path / "source.md"
    source.write_text(
        "# Introduction\n\nBody.\n\n![Figure 1|span=page](source.png)\n\nFigure 1. Example inline figure.\n",
        encoding="utf-8",
    )
    output = tmp_path / "output.docx"
    result = assemble_markdown_template([source], template_path=template, output=output, metadata_path=metadata)
    document = Document(output)
    assert result.figures and result.figures[0]["caption"].startswith("Figure 1.")
    assert result.figures[0]["span"] == "page"
    assert result.figures[0]["columns"] == 1
    assert result.figures[0]["orientation"] == "portrait"
    assert len(document.inline_shapes) == 1
    assert any(p.text.startswith("Figure 1.") for p in document.paragraphs)
    verify_template_output(output, expected_sections=4)


def test_replacement_figure_drops_template_crop_and_fills_section_width(tmp_path: Path) -> None:
    template_image = tmp_path / "template.png"
    source_image = tmp_path / "source.png"
    Image.new("RGB", (200, 80), "white").save(template_image)
    Image.new("RGB", (400, 300), "white").save(source_image)
    template = tmp_path / "template.docx"
    _figure_template(template, template_image)
    document = Document(template)
    drawing = document.inline_shapes[0]._inline
    source_rect = OxmlElement("a:srcRect")
    source_rect.set("t", "10000")
    source_rect.set("b", "5000")
    blip_fill = drawing.find(".//" + qn("pic:blipFill"))
    blip_fill.insert(1, source_rect)
    document.save(template)
    metadata = tmp_path / "metadata.md"
    metadata.write_text("# TITLE\n\nA title\n", encoding="utf-8")
    source = tmp_path / "source.md"
    source.write_text(
        "![Figure 1](source.png)\n\nFigure 1. Replacement.\n",
        encoding="utf-8",
    )
    output = tmp_path / "output.docx"
    assemble_markdown_template(
        [source], template_path=template, output=output, metadata_path=metadata
    )
    rendered = Document(output)
    shape = rendered.inline_shapes[0]
    section = rendered.sections[-1]
    expected_width = section.page_width - section.left_margin - section.right_margin
    assert shape.width == expected_width
    assert abs((shape.width / shape.height) - (4 / 3)) < 1e-6
    source_rects = shape._inline.findall(".//" + qn("a:srcRect"))
    assert all(not node.attrib for node in source_rects)


def test_word_compatible_image_bytes_normalizes_large_rgba_png(tmp_path: Path) -> None:
    source = tmp_path / "large-rgba.png"
    Image.new("RGBA", (5000, 3000), (255, 255, 255, 255)).save(source)
    payload = _word_compatible_image_bytes(source)
    with Image.open(BytesIO(payload)) as image:
        assert image.format == "PNG"
        assert image.mode == "RGB"
        assert image.size == (4096, 2458)


def test_figure_span_cli_default_and_override() -> None:
    parser = build_parser()
    default = parser.parse_args(["md2docx", "source.md", "-o", "out.docx"])
    override = parser.parse_args(
        ["md2docx", "source.md", "-o", "out.docx", "--figure-span", "page"]
    )
    assert default.figure_span == "column"
    assert override.figure_span == "page"


def test_portrait_page_span_uses_continuous_two_one_two_sections(tmp_path: Path) -> None:
    source_image = tmp_path / "source.png"
    Image.new("RGB", (160, 80), "white").save(source_image)
    template = tmp_path / "template.docx"
    _figure_template(template, source_image)
    document = Document(template)
    body_columns = document.sections[1]._sectPr.find(qn("w:cols"))
    if body_columns is None:
        body_columns = OxmlElement("w:cols")
        document.sections[1]._sectPr.append(body_columns)
    body_columns.set(qn("w:num"), "2")
    body_columns.set(qn("w:space"), "475")
    document.save(template)
    metadata = tmp_path / "metadata.md"
    metadata.write_text("# TITLE\n\nA title\n", encoding="utf-8")
    source = tmp_path / "source.md"
    source.write_text(
        "Body before.\n\n![Figure 1|span=page](source.png)\n\n"
        "Figure 1. Wide portrait figure.\n\nBody after.\n",
        encoding="utf-8",
    )
    output = tmp_path / "output.docx"
    result = assemble_markdown_template(
        [source], template_path=template, output=output, metadata_path=metadata,
        line_numbers="on", figure_span="column",
    )
    rendered = Document(output)
    columns = []
    section_types = []
    for section in rendered.sections:
        cols = section._sectPr.find(qn("w:cols"))
        columns.append(int(cols.get(qn("w:num"), "1")) if cols is not None else 1)
        section_type = section._sectPr.find(qn("w:type"))
        section_types.append(section_type.get(qn("w:val")) if section_type is not None else None)
        assert section._sectPr.find(qn("w:lnNumType")) is not None
    assert columns[-3:] == [2, 1, 2]
    assert all(value in {None, "continuous"} for value in section_types)
    assert '<w:br w:type="page"' not in rendered._element.xml
    assert result.figure_span == "column"
    assert result.figures[0]["span"] == "page"
    assert result.figures[0]["columns"] == 1


def test_invalid_and_conflicting_image_span_markers_fail(tmp_path: Path) -> None:
    source_image = tmp_path / "source.png"
    Image.new("RGB", (80, 40), "white").save(source_image)
    template = tmp_path / "template.docx"
    _figure_template(template, source_image)
    metadata = tmp_path / "metadata.md"
    metadata.write_text("# TITLE\n\nA title\n", encoding="utf-8")
    for marker, message in (
        ("span=wide", "span marker"),
        ("span=page|columns=single", "cannot be combined"),
    ):
        source = tmp_path / f"{marker.replace('|', '-')}.md"
        source.write_text(f"![Figure 1|{marker}](source.png)\n", encoding="utf-8")
        with pytest.raises(ValueError, match=message):
            assemble_markdown_template(
                [source], template_path=template, output=tmp_path / "out.docx",
                metadata_path=metadata,
            )


def test_terminal_headings_share_h1_style_without_list_numbering(tmp_path: Path) -> None:
    template = tmp_path / "template.docx"
    _template(template)
    document = Document(template)
    heading = document.add_paragraph("Numbered prototype", style="Heading 1")
    numbering = OxmlElement("w:numPr")
    numbering.append(OxmlElement("w:ilvl"))
    numbering.append(OxmlElement("w:numId"))
    heading._p.get_or_add_pPr().append(numbering)
    document.save(template)
    metadata = tmp_path / "metadata.md"
    metadata.write_text(
        "# TITLE\n\nA title\n# ACKNOWLEDGMENTS\n\nThanks.\n"
        "# AUTHOR CONTRIBUTIONS\n\nA. Author contributed.\n"
        "# CODE AVAILABILITY\n\nCode is available.\n",
        encoding="utf-8",
    )
    source = tmp_path / "source.md"
    source.write_text("Body \\citep{ref}.\n", encoding="utf-8")
    bibliography = tmp_path / "refs.json"
    bibliography.write_text('{"ref": "Author. Journal. 2026."}', encoding="utf-8")
    output = tmp_path / "output.docx"
    result = assemble_markdown_template(
        [source], template_path=template, output=output, metadata_path=metadata,
        bibliography_path=bibliography, line_numbers="on",
    )
    rendered = Document(output)
    terminal_texts = {
        "ACKNOWLEDGMENTS", "AUTHOR CONTRIBUTIONS", "CODE AVAILABILITY", "REFERENCES"
    }
    headings = [p for p in rendered.paragraphs if p.text in terminal_texts]
    assert {p.text for p in headings} == terminal_texts
    assert {p.style.style_id for p in headings} == {result.style_map["heading_1"]}
    assert all(p._p.find(".//" + qn("w:numPr")) is None for p in headings)
    assert all(p._p.find(".//" + qn("w:outlineLvl")) is not None for p in headings)
    assert all(section._sectPr.find(qn("w:lnNumType")) is not None for section in rendered.sections)


def test_table_fits_explicit_one_column_section(tmp_path: Path) -> None:
    template = tmp_path / "template.docx"
    _template(template)
    document = Document(template)
    columns = document.sections[1]._sectPr.find(qn("w:cols"))
    if columns is None:
        columns = OxmlElement("w:cols")
        document.sections[1]._sectPr.append(columns)
    columns.set(qn("w:num"), "2")
    document.save(template)
    metadata = tmp_path / "metadata.md"
    metadata.write_text("# TITLE\n\nA title\n", encoding="utf-8")
    source = tmp_path / "source.md"
    source.write_text("# Intro\n\n| One | Two |\n|---|---|\n| Alpha | Beta |\n", encoding="utf-8")
    output = tmp_path / "output.docx"
    assemble_markdown_template([source], template_path=template, output=output, metadata_path=metadata, columns="one")
    result = Document(output)
    table = result.tables[0]._tbl
    width = sum(int(node.get(qn("w:w"))) for node in table.find(qn("w:tblGrid")))
    section = result.sections[1]
    expected = int((section.page_width - section.left_margin - section.right_margin) / 635)
    assert width == expected


def test_template_inherits_superscript_citations_without_private_markers(tmp_path: Path) -> None:
    template = tmp_path / "template.docx"
    _template(template)
    document = Document(template)
    document.add_paragraph().add_run("1").font.superscript = True
    document.save(template)
    metadata = tmp_path / "metadata.md"
    metadata.write_text("# TITLE\n\nA title\n", encoding="utf-8")
    source = tmp_path / "source.md"
    source.write_text("# Intro\n\nFirst \\citep{ref}.\n\nSecond \\cite{ref}.\n\nThird \\citep{ref}.\n", encoding="utf-8")
    bibliography = tmp_path / "references.json"
    bibliography.write_text('{"ref": "Author. Journal. 2026."}', encoding="utf-8")
    output = tmp_path / "output.docx"
    result = assemble_markdown_template([source], template_path=template, output=output, metadata_path=metadata, bibliography_path=bibliography)
    document = Document(output)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "\ue000" not in text and "\ue001" not in text
    assert result.verification["citation_count"] == 3
    assert sum(run.font.superscript is True for paragraph in document.paragraphs for run in paragraph.runs) == 3
    first = next(p for p in document.paragraphs if p.text.startswith("First"))
    assert first.text == "First1."
    assert first.runs[0].text == "First"
    assert first.runs[1].text == "1" and first.runs[1].font.superscript is True
    assert all(
        (section._sectPr.find(qn("w:type")) is None
         or section._sectPr.find(qn("w:type")).get(qn("w:val")) == "continuous")
        for section in document.sections
    )


def test_si_reuses_main_citation_numbers_and_prefixes_si_only_references(tmp_path: Path) -> None:
    template = tmp_path / "template.docx"
    _template(template)
    metadata = tmp_path / "metadata.md"
    metadata.write_text("# TITLE\n\nA title\n", encoding="utf-8")
    bibliography = tmp_path / "references.json"
    bibliography.write_text(
        '{"shared": "Shared. Journal. 2020.", "si_only": "SI only. Journal. 2021."}',
        encoding="utf-8",
    )
    main_source = tmp_path / "main.md"
    main_source.write_text("# Introduction\n\nMain \\citep{shared}.\n", encoding="utf-8")
    main_output = tmp_path / "main.docx"
    main_result = assemble_markdown_template(
        [main_source], template_path=template, output=main_output,
        metadata_path=metadata, bibliography_path=bibliography,
    )
    main_manifest, _ = write_assembly_sidecars(
        main_result, inputs=[main_source], template_path=template,
        metadata_path=metadata, bibliography_path=bibliography,
        command=["docforge", "md2docx"],
    )
    si_source = tmp_path / "si.md"
    si_source.write_text("# Supporting Information\n\nSI \\citep{shared,si_only}.\n", encoding="utf-8")
    si_output = tmp_path / "si.docx"
    si_result = assemble_markdown_template(
        [si_source], template_path=template, output=si_output,
        metadata_path=metadata, bibliography_path=bibliography,
        citation_base_path=main_manifest,
    )
    text = "\n".join(paragraph.text for paragraph in Document(si_output).paragraphs)
    assert si_result.citation_map == {"shared": 1, "si_only": "S1"}
    assert "\ue000" not in text and "[1,S1]" in text
    assert "S1.\tSI only." in text


def test_template_assembly_preserves_inline_runs_and_numbers_display_equation(tmp_path: Path) -> None:
    template = tmp_path / "template.docx"
    _template(template)
    metadata = tmp_path / "metadata.md"
    metadata.write_text("# TITLE\n\nA title\n", encoding="utf-8")
    source = tmp_path / "source.md"
    source.write_text(
        "# Introduction\n\nInline $t_{\\mathrm{chem}}$ uses runs.\n\n"
        "$$\nE = mc^2\n$$\n",
        encoding="utf-8",
    )
    output = tmp_path / "output.docx"
    assemble_markdown_template([source], template_path=template, output=output, metadata_path=metadata)
    document = Document(output)
    inline = next(p for p in document.paragraphs if "Inline" in p.text)
    equation = next(p for p in document.paragraphs if "(1)" in p.text)
    assert "oMath" not in inline._p.xml
    assert any(run.text == "chem" and run.font.subscript for run in inline.runs)
    assert "oMathPara" in equation._p.xml
    assert "(1)" in equation.text
    tab = equation._p.find(".//" + qn("w:tab"))
    assert tab is not None and int(tab.get(qn("w:pos"))) > 0


def test_native_toc_and_heading_number_reset(tmp_path: Path) -> None:
    template = tmp_path / "template.docx"
    _template(template)
    metadata = tmp_path / "metadata.md"
    metadata.write_text("# TITLE\n\nA title\n", encoding="utf-8")
    source = tmp_path / "si.md"
    source.write_text(
        "# Methods\n\n## One\n\nBody.\n\n## Two\n\nBody.\n\n"
        "# Supplementary Tables\n\n## Table Section\n\n### Detail\n\nBody.\n\n"
        "# Supplementary Figures\n\n## Figure Section\n\nBody.\n",
        encoding="utf-8",
    )
    output = tmp_path / "si.docx"
    result = assemble_markdown_template(
        [source], template_path=template, output=output, metadata_path=metadata,
        native_toc=True, restart_heading_numbering=True, body_first_line_chars=2,
        page_break_before_h1=True,
    )
    document = Document(output)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "CONTENTS" in text
    contents = next(p for p in document.paragraphs if p.text == "CONTENTS")
    assert contents.style.name == "TOC Heading"
    assert all(value in text for value in ("METHODS", "SUPPLEMENTARY TABLES", "SUPPLEMENTARY FIGURES"))
    headings = [p for p in document.paragraphs if p._p.find(".//" + qn("w:outlineLvl")) is not None]
    numbered = [p.text for p in headings if p.text[:1].isdigit()]
    assert numbered == ["1. One", "2. Two", "1. Table Section", "1. Figure Section"]
    assert "Detail" in [p.text for p in headings]
    detail = next(p for p in headings if p.text == "Detail")
    assert not detail.text[:1].isdigit()
    for paragraph in headings:
        indentation = paragraph._p.find(".//" + qn("w:ind"))
        assert indentation is not None
        assert indentation.get(qn("w:left")) == "0"
        assert indentation.get(qn("w:firstLine")) == "0"
    xml = document._element.xml
    assert 'TOC \\o "1-3" \\h \\z \\u' in xml
    assert result.native_toc is True
    assert result.restart_heading_numbering is True
    body = next(p for p in document.paragraphs if p.text == "Body.")
    body_indent = body._p.find(".//" + qn("w:ind"))
    assert body_indent is not None
    assert body_indent.get(qn("w:firstLineChars")) == "200"
    assert result.body_first_line_chars == 2
    for paragraph in headings:
        if paragraph.text in {"METHODS", "SUPPLEMENTARY TABLES", "SUPPLEMENTARY FIGURES"}:
            assert paragraph._p.find(".//" + qn("w:pageBreakBefore")) is not None
    assert result.page_break_before_h1 is True


def test_template_body_and_caption_typography(tmp_path: Path) -> None:
    source_image = tmp_path / "source.png"
    Image.new("RGB", (80, 40), "white").save(source_image)
    template = tmp_path / "template.docx"
    _figure_template(template, source_image)
    metadata = tmp_path / "metadata.md"
    metadata.write_text("# TITLE\n\nA title\n", encoding="utf-8")
    source = tmp_path / "source.md"
    source.write_text(
        "# Figures\n\nBody text.\n\nTable: Example values.\n\n"
        "| A | B |\n|---|---|\n| 1 | 2 |\n\n"
        "![Figure 1](source.png)\n\nFigure 1. Caption body.\n",
        encoding="utf-8",
    )
    output = tmp_path / "output.docx"
    result = assemble_markdown_template(
        [source], template_path=template, output=output, metadata_path=metadata,
        numbering_prefix="S",
    )
    document = Document(output)
    body = next(p for p in document.paragraphs if p.text == "Body text.")
    assert all(run._r.rPr.sz.get(qn("w:val")) == "22" for run in body.runs if run.text)
    for prefix in ("Table S1.", "Figure S1."):
        caption = next(p for p in document.paragraphs if p.text.startswith(prefix))
        assert all(run._r.rPr.sz.get(qn("w:val")) == "20" for run in caption.runs if run.text)
        assert caption.runs[0].bold is True
        assert all(run.bold is not True and run.italic is not True for run in caption.runs[1:] if run.text.strip())
    assert result.figures[0]["orientation"] == "portrait"


def test_caption_preserves_inline_math_italics_and_superscripts(tmp_path: Path) -> None:
    source_image = tmp_path / "source.png"
    Image.new("RGB", (80, 40), "white").save(source_image)
    template = tmp_path / "template.docx"
    _figure_template(template, source_image)
    metadata = tmp_path / "metadata.md"
    metadata.write_text("# TITLE\n\nA title\n", encoding="utf-8")
    source = tmp_path / "source.md"
    source.write_text(
        "![Figure 1](source.png)\n\nFigure 1. The $f^{+}$ response.\n",
        encoding="utf-8",
    )
    output = tmp_path / "output.docx"
    assemble_markdown_template(
        [source], template_path=template, output=output, metadata_path=metadata
    )
    caption = next(p for p in Document(output).paragraphs if p.text.startswith("Figure 1."))
    f_run = next(run for run in caption.runs if run.text == "f")
    sign_run = next(run for run in caption.runs if run.text == "+")
    assert f_run.italic is True
    assert sign_run.font.superscript is True


def test_no_title_strips_level_one_and_generated_reference_heading(tmp_path: Path) -> None:
    template = tmp_path / "template.docx"
    _template(template)
    metadata = tmp_path / "metadata.md"
    metadata.write_text("# TITLE\n\nA title\n", encoding="utf-8")
    source = tmp_path / "source.md"
    source.write_text(
        "# Introduction\n\nBody \\citep{ref}.\n\n## Detail\n\nMore.\n",
        encoding="utf-8",
    )
    bibliography = tmp_path / "references.json"
    bibliography.write_text('{"ref": "Author. Journal. 2026."}', encoding="utf-8")
    output = tmp_path / "output.docx"
    result = assemble_markdown_template(
        [source],
        template_path=template,
        output=output,
        metadata_path=metadata,
        bibliography_path=bibliography,
        strip_level_one_headings=True,
    )
    text = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
    assert result.strip_level_one_headings is True
    assert "A title" in text
    assert "Body [1]." in text
    assert "Detail" in text
    assert "Introduction" not in text
    assert "REFERENCES" in text
    assert "1.\tAuthor." in text


def test_si_only_reference_scope_omits_shared_entries_and_metadata_back_matter(tmp_path: Path) -> None:
    template = tmp_path / "template.docx"
    _template(template)
    metadata = tmp_path / "metadata.md"
    metadata.write_text(
        "# TITLE\n\nA title\n# ACKNOWLEDGMENTS\n\nThanks.\n# AUTHOR CONTRIBUTIONS\n\nContributed.\n",
        encoding="utf-8",
    )
    bibliography = tmp_path / "references.json"
    bibliography.write_text('{"shared": "Shared. Journal. 2020.", "si_only": "SI only. Journal. 2021."}', encoding="utf-8")
    main_source = tmp_path / "main.md"
    main_source.write_text("# Intro\n\nMain \\citep{shared}.\n", encoding="utf-8")
    main_output = tmp_path / "main.docx"
    main_result = assemble_markdown_template(
        [main_source], template_path=template, output=main_output,
        metadata_path=metadata, bibliography_path=bibliography,
    )
    main_manifest, _ = write_assembly_sidecars(
        main_result, inputs=[main_source], template_path=template,
        metadata_path=metadata, bibliography_path=bibliography,
        command=["docforge", "md2docx"],
    )
    si_source = tmp_path / "si.md"
    si_source.write_text("# SI\n\nSI \\citep{shared,si_only}.\n", encoding="utf-8")
    si_output = tmp_path / "si.docx"
    si_result = assemble_markdown_template(
        [si_source], template_path=template, output=si_output,
        metadata_path=metadata, bibliography_path=bibliography,
        citation_base_path=main_manifest, bibliography_scope="new-only",
        include_metadata_back_matter=False,
    )
    text = "\n".join(p.text for p in Document(si_output).paragraphs)
    assert si_result.reference_keys == ("si_only",)
    assert "[1,S1]" in text
    assert "S1.\tSI only." in text
    assert "Shared. Journal" not in text
    assert "ACKNOWLEDGMENTS" not in text
    assert "AUTHOR CONTRIBUTIONS" not in text
    assert "REFERENCES" in text


def test_unfilled_optional_metadata_is_omitted(tmp_path: Path) -> None:
    template = tmp_path / "template.docx"
    _template(template)
    metadata = tmp_path / "metadata.md"
    metadata.write_text("# TITLE\n\nA title\n\n# Acknowledgement\n\nTODO\n", encoding="utf-8")
    source = tmp_path / "source.md"
    source.write_text("# Introduction\n\nBody.\n", encoding="utf-8")
    output = tmp_path / "output.docx"
    assemble_markdown_template(
        [source], template_path=template, output=output, metadata_path=metadata
    )
    text = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
    assert "TODO" not in text
    assert "ACKNOWLEDGMENT" not in text

def test_template_rendering_args_override_font_style_and_line_numbers(tmp_path: Path) -> None:
    template = tmp_path / "template.docx"
    _template(template)
    metadata = tmp_path / "metadata.md"
    metadata.write_text("# TITLE\n\nA title\n", encoding="utf-8")
    style = tmp_path / "styles.json"
    style.write_text('{"body": "TA_Main_Text"}', encoding="utf-8")
    source = tmp_path / "source.md"
    source.write_text("# Introduction\n\nBody.\n", encoding="utf-8")
    output = tmp_path / "output.docx"
    result = assemble_markdown_template(
        [source],
        template_path=template,
        output=output,
        metadata_path=metadata,
        font_family="Arial",
        east_asia_font="SimSun",
        style_profile=str(style),
        line_numbers="on",
        strip_level_one_headings=True,
        heading_before=6,
    )
    document = Document(output)
    assert result.style_profile == str(style)
    assert result.font_family == "Arial"
    assert result.east_asia_font == "SimSun"
    assert result.line_numbers == "on"
    assert result.include_title is True
    assert result.strip_level_one_headings is True
    assert result.heading_before == 6
    assert all(section._sectPr.find(qn("w:lnNumType")) is not None for section in document.sections)
    text = "\n".join(p.text for p in document.paragraphs)
    assert "A title" in text
    assert "Introduction" not in text
    text_runs = [run for paragraph in document.paragraphs for run in paragraph.runs if run.text]
    assert text_runs
    for run in text_runs:
        fonts = run._r.rPr.rFonts
        assert fonts.get(qn("w:ascii")) == "Arial"
        assert fonts.get(qn("w:hAnsi")) == "Arial"
        assert fonts.get(qn("w:eastAsia")) == "SimSun"


def test_si_numbering_prefixes_figures_tables_and_equations(tmp_path: Path) -> None:
    source_image = tmp_path / "source.png"
    Image.new("RGB", (80, 40), "white").save(source_image)
    template = tmp_path / "template.docx"
    _figure_template(template, source_image)
    metadata = tmp_path / "metadata.md"
    metadata.write_text("# TITLE\n\nA title\n", encoding="utf-8")
    source = tmp_path / "si.md"
    source.write_text(
        "# Supporting Information\n\n"
        "![Figure 9](source.png)\n\n"
        "Figure 9. Example figure.\n\n"
        "Table: Example values.\n\n"
        "| A | B |\n|---|---|\n| 1 | 2 |\n\n"
        "$$\nE = mc^2\n$$\n",
        encoding="utf-8",
    )
    output = tmp_path / "si.docx"
    result = assemble_markdown_template(
        [source],
        template_path=template,
        output=output,
        metadata_path=metadata,
        numbering_prefix="S",
    )
    document = Document(output)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Figure S1." in text
    assert "Table S1." in text
    assert "(S1)" in text
    assert result.figures[0]["label"] == "Figure S1"
    assert result.tables[0]["label"] == "Table S1"
    assert result.numbering_prefix == "S"


def test_main_numbering_prefix_remains_numeric_by_default(tmp_path: Path) -> None:
    source = tmp_path / "source.md"
    source.write_text("$$\na = b\n$$\n", encoding="utf-8")
    from docforge.markdown import parse_markdown, render_blocks_to_doc
    document = render_blocks_to_doc(parse_markdown(source))
    assert "(1)" in "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "(S1)" not in document._element.xml
